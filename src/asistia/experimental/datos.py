"""Lectura nativa DRE, controles y publicación atómica de perfiles institucionales."""

from __future__ import annotations

import calendar
import copy
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from psycopg.types.json import Jsonb

from asistia.cierre.ingesta import MESES, normal
from asistia.db import sha256_de

VERSION = "PERFILES_DRE_1"
CAMPOS = (
    "dias_justificados",
    "dias_injustificados",
    "horas_no_justificadas",
    "horas_justificadas",
)
FEATURES = (
    "dias_justificados_por_registro_conocido",
    "dias_injustificados_por_registro_conocido",
    "horas_no_justificadas_por_registro_conocido",
    "fraccion_registros_con_injustificadas",
    "fraccion_conteos_desconocidos",
)


def huella(valor):
    return hashlib.sha256(
        json.dumps(
            valor, sort_keys=True, ensure_ascii=False, default=str, allow_nan=False
        ).encode()
    ).hexdigest()


def mes(valor):
    if not re.fullmatch(r"\d{4}-\d{2}", str(valor)):
        raise ValueError("Usa un período válido en formato AAAA-MM.")
    try:
        return date.fromisoformat(valor + "-01")
    except ValueError:
        raise ValueError("Usa un período válido en formato AAAA-MM.") from None


def mapa_periodos(origenes, inicio=None):
    """Trasladar un intervalo sin juntar, duplicar o cambiar el orden de los meses."""
    origenes = sorted({str(x)[:7] for x in origenes})
    if not origenes:
        raise ValueError("No hay períodos de origen disponibles.")
    primero = mes(origenes[0])
    destino = mes(inicio) if inicio is not None else primero
    desplazamiento = (destino.year - primero.year) * 12 + destino.month - primero.month
    resultado = {}
    for valor in origenes:
        fuente = mes(valor)
        indice = fuente.year * 12 + fuente.month - 1 + desplazamiento
        try:
            resultado[valor] = date(indice // 12, indice % 12 + 1, 1).strftime("%Y-%m")
        except ValueError:
            raise ValueError(
                "El intervalo de destino queda fuera del calendario permitido."
            ) from None
    return resultado


def conteo(valor, campo, periodo):
    """Un guion es cero supuesto solo aquí. Vacío/ilegible conservan ausencia."""
    if valor is None or str(valor).strip() == "":
        return {"recibido": valor, "observado": None, "usado": None, "estado": "VACIO"}
    if str(valor).strip() in {"-", "--", "—", "–"}:
        return {
            "recibido": valor,
            "observado": None,
            "usado": 0.0,
            "estado": "SUPUESTO_GUION_CERO",
        }
    try:
        if isinstance(valor, bool):
            raise TypeError
        numero = float(str(valor).strip().replace(",", "."))
        if not math.isfinite(numero):
            raise ValueError
    except (ValueError, TypeError):
        return {
            "recibido": str(valor),
            "observado": None,
            "usado": None,
            "estado": "NO_INTERPRETABLE",
        }
    limite = (
        calendar.monthrange(periodo.year, periodo.month)[1]
        if campo.startswith("dias_")
        else 24 * calendar.monthrange(periodo.year, periodo.month)[1]
    )
    valido = 0 <= numero <= limite
    return {
        "recibido": valor,
        "observado": numero,
        "usado": numero if valido else None,
        "estado": "OBSERVADO" if valido else "FUERA_DE_RANGO",
    }


def cabecera(ws):
    """Contrato de las tablas DRE actuales; no adivinar columnas por posición sola."""
    encabezado = next(
        (
            row[0].row
            for row in ws.iter_rows(max_row=min(40, ws.max_row))
            if any("APELLIDOS Y NOMBRES" in normal(c.value) for c in row)
        ),
        None,
    )
    if not encabezado:
        raise ValueError("DRE sin encabezado nominal reconocible.")
    columnas = {}
    for c in ws[encabezado + 1]:
        if normal(c.value) == "JUSTIFICADAS" and "DIAS" in normal(
            ws.cell(encabezado + 2, c.column).value
        ):
            columnas["dias_justificados"] = c.column
        if normal(c.value) == "INJUSTIFICADAS" and "DIAS" in normal(
            ws.cell(encabezado + 2, c.column).value
        ):
            columnas["dias_injustificados"] = c.column
    for c in ws[encabezado]:
        if normal(c.value) == "HORAS NO JUSTIFICADAS":
            columnas["horas_no_justificadas"] = c.column
        if normal(c.value) == "HORAS JUSTIFICADAS":
            columnas["horas_justificadas"] = c.column
    # Secundaria agrupa HORAS debajo de INJUSTIFICADAS, tras la columna de fecha.
    j = columnas.get("dias_injustificados")
    if j and normal(ws.cell(encabezado + 2, j + 2).value) == "HORAS":
        columnas["horas_no_justificadas"] = j + 2
    if set(columnas) != set(CAMPOS):
        raise ValueError(
            "La tabla DRE cambió: faltan encabezados de conteos. Requiere adaptar el lector."
        )
    candidatos = []
    for row in ws.iter_rows(max_row=encabezado - 1):
        for c in row:
            t = normal(c.value)
            for nombre, numero in MESES.items():
                encontrado = re.search(
                    r"MES DE\s+" + nombre + r"\s*[-–]?\s*(20\d{2})\b", t
                )
                if encontrado:
                    candidatos.append(
                        {
                            "celda": c.coordinate,
                            "periodo": f"{encontrado[1]}-{numero:02}",
                            "texto": str(c.value),
                        }
                    )
    if not candidatos:
        raise ValueError("DRE sin mes/año asociado a la tabla.")
    # Encabezado más cercano a la tabla, preservando las etiquetas antiguas.
    return columnas, mes(candidatos[-1]["periodo"]), candidatos


def resumir(filas):
    componentes = {}
    for campo in CAMPOS:
        celdas = [r["conteos"][campo] for r in filas]
        conocidos = [c["usado"] for c in celdas if c["usado"] is not None]
        observados = [c["observado"] for c in celdas if c["estado"] == "OBSERVADO"]
        componentes[campo] = {
            "suma_observada_valida": sum(observados) if observados else None,
            "celdas_observadas_validas": len(observados),
            "suma_utilizada": sum(conocidos) if conocidos else None,
            "registros_conocidos": len(conocidos),
            "registros_desconocidos": len(filas) - len(conocidos),
            "estados": dict(Counter(c["estado"] for c in celdas)),
        }
    features = {}
    for campo in CAMPOS[:3]:
        datos = componentes[campo]
        features[campo + "_por_registro_conocido"] = (
            datos["suma_utilizada"] / datos["registros_conocidos"]
            if datos["registros_conocidos"]
            else None
        )
    injustificados = [
        r["conteos"]["dias_injustificados"]["usado"]
        for r in filas
        if r["conteos"]["dias_injustificados"]["usado"] is not None
    ]
    features["fraccion_registros_con_injustificadas"] = (
        sum(x > 0 for x in injustificados) / len(injustificados)
        if injustificados
        else None
    )
    features["fraccion_conteos_desconocidos"] = sum(
        componentes[c]["registros_desconocidos"] for c in CAMPOS[:3]
    ) / (3 * len(filas))
    return {
        "registros_laborales": len(filas),
        "componentes": componentes,
        "features": features,
        "apto_modelo": any(componentes[c]["registros_conocidos"] for c in CAMPOS[:2]),
        "naturaleza": "SINTETICO_DERIVADO_DRE",
        "remuneracion": "NO_APLICA_EXPERIMENTO",
    }


def leer_corpus(conn, carpeta):
    """Usa originales y la identidad ya resuelta, sin copiar nombres al experimento."""
    archivos = sorted(Path(carpeta).glob("*.xlsx"))
    archivos = [p for p in archivos if not p.name.startswith("~$")]
    if not archivos:
        raise ValueError("La carpeta no contiene reportes DRE .xlsx.")
    fuentes, filas, excluidas = [], [], []
    for ruta in archivos:
        sha = sha256_de(ruta)
        doc = conn.execute(
            "SELECT cierre_documento_id FROM cierre_documento WHERE tipo='DRE' AND sha256=%s",
            (sha,),
        ).fetchone()
        if not doc:
            raise ValueError(
                f"Primero importa el padrón DRE desde Cierre: {ruta.name}."
            )
        registros = conn.execute(
            "SELECT * FROM padron_evidencia WHERE cierre_documento_id=%s ORDER BY hoja,fila",
            (doc["cierre_documento_id"],),
        ).fetchall()
        if not registros:
            raise ValueError(
                f"El DRE no tiene filas nominales importadas: {ruta.name}."
            )
        libro = openpyxl.load_workbook(ruta, data_only=True)
        try:
            cabeceras = {ws.title: cabecera(ws) for ws in libro}
            fuentes.append(
                {
                    "archivo": ruta.name,
                    "sha256": sha,
                    "cierre_documento_id": str(doc["cierre_documento_id"]),
                    "filas_padron": len(registros),
                    "hojas": {
                        k: {"periodo": str(v[1]), "encabezados": v[2]}
                        for k, v in cabeceras.items()
                    },
                }
            )
            for r in registros:
                cols, periodo, _ = cabeceras[r["hoja"]]
                loc = {
                    "archivo_sha256": sha,
                    "cierre_documento_id": str(doc["cierre_documento_id"]),
                    "hoja": r["hoja"],
                    "fila": r["fila"],
                    "periodo_fuente": str(periodo),
                    "padron_evidencia_id": str(r["padron_evidencia_id"]),
                }
                if r["institucion_educativa_id"] is None:
                    excluidas.append({**loc, "motivo": "INSTITUCION_NO_RESUELTA"})
                    continue
                celdas = {}
                ws = libro[r["hoja"]]
                for campo, columna in cols.items():
                    celda = ws.cell(r["fila"], columna)
                    anterior = r["datos_raw"]["valores_fila"][columna - 1]
                    if celda.value != anterior:
                        raise ValueError(
                            f"La fila {r['fila']} no coincide con el padrón conservado en {ruta.name}."
                        )
                    celdas[campo] = {
                        **conteo(celda.value, campo, periodo),
                        "celda": celda.coordinate,
                    }
                filas.append(
                    {
                        **loc,
                        "institucion_educativa_id": r["institucion_educativa_id"],
                        "identidad_fila": huella(
                            [normal(r["nombres_raw"]), normal(r["cargo_raw"])]
                        ),
                        "conteos": celdas,
                    }
                )
        finally:
            libro.close()
        if sha256_de(ruta) != sha:
            raise ValueError(
                "Un original cambió durante la lectura; vuelve a intentar con la versión fijada."
            )
    grupos = defaultdict(list)
    for r in filas:
        grupos[
            (r["institucion_educativa_id"], r["periodo_fuente"], r["identidad_fila"])
        ].append(r)
    aceptadas = []
    for grupo in grupos.values():
        valores = {
            huella({k: v["recibido"] for k, v in r["conteos"].items()}) for r in grupo
        }
        if len(valores) != 1:
            excluidas.extend(
                {**r, "motivo": "CONFLICTO_MISMA_IDENTIDAD_Y_ROL"} for r in grupo
            )
        else:
            aceptadas.append(grupo[0])
            excluidas.extend(
                {**r, "motivo": "DUPLICADO_EXACTO_MISMA_IDENTIDAD_Y_ROL"}
                for r in grupo[1:]
            )
    grupos_mes = defaultdict(list)
    for r in aceptadas:
        grupos_mes[(r["institucion_educativa_id"], r["periodo_fuente"])].append(r)
    reportes = [
        {
            "institucion_educativa_id": ie,
            "periodo_fuente": periodo,
            "datos": resumir(rows),
            "fuentes": rows,
        }
        for (ie, periodo), rows in sorted(grupos_mes.items())
    ]
    total = sum(f["filas_padron"] for f in fuentes)
    assert total == len(aceptadas) + len(excluidas)
    return reportes, {
        "version": VERSION,
        "fuentes": fuentes,
        "filas_leidas": total,
        "filas_aceptadas": len(aceptadas),
        "filas_excluidas": len(excluidas),
        "exclusiones_por_motivo": dict(Counter(r["motivo"] for r in excluidas)),
        "exclusiones": excluidas,
        "supuestos": [
            "Guion se usa como cero exclusivamente en este lote sintético.",
            "Vacíos y conteos fuera de rango permanecen desconocidos.",
            "No se reconstruyen fechas diarias ni pago.",
            "Una fila laboral no equivale necesariamente a una persona única.",
        ],
    }


def publicar(
    conn,
    reportes,
    manifiesto,
    *,
    inicio=None,
    autor="AGENTE_TECNICO",
    motivo,
    padre_id=None,
):
    motivo = motivo.strip()
    if not motivo or len(motivo) > 4000:
        raise ValueError("Escribe un motivo de hasta 4000 caracteres.")
    mapa = mapa_periodos([r["periodo_fuente"] for r in reportes], inicio)
    config = {"version": VERSION, "mapa_periodos": mapa}
    hash_lote = huella(
        {
            "config": config,
            "reportes": reportes,
            "manifiesto": manifiesto,
            "padre": str(padre_id) if padre_id else None,
        }
    )
    with conn.transaction():
        nuevo = conn.execute(
            """INSERT INTO ml_lote(huella,padre_id,configuracion,manifiesto,autor,motivo)
            VALUES(%s,%s,%s,%s,%s,%s) ON CONFLICT(huella) DO NOTHING RETURNING ml_lote_id""",
            (hash_lote, padre_id, Jsonb(config), Jsonb(manifiesto), autor, motivo),
        ).fetchone()
        if not nuevo:
            return conn.execute(
                "SELECT ml_lote_id FROM ml_lote WHERE huella=%s", (hash_lote,)
            ).fetchone()["ml_lote_id"], False
        lid = nuevo["ml_lote_id"]
        for r in reportes:
            conn.execute(
                """INSERT INTO ml_reporte_mensual(ml_lote_id,institucion_educativa_id,periodo_fuente,periodo,datos,fuentes)
                VALUES(%s,%s,%s,%s,%s,%s)""",
                (
                    lid,
                    r["institucion_educativa_id"],
                    r["periodo_fuente"],
                    mes(mapa[str(r["periodo_fuente"])[:7]]),
                    Jsonb(r["datos"]),
                    Jsonb(r["fuentes"]),
                ),
            )
    return lid, True


def redefinir(conn, lote_id, inicio, *, autor, motivo):
    lote = conn.execute(
        "SELECT * FROM ml_lote WHERE ml_lote_id=%s", (lote_id,)
    ).fetchone()
    if not lote:
        raise ValueError("El lote solicitado no existe.")
    filas = conn.execute(
        "SELECT institucion_educativa_id,periodo_fuente,datos,fuentes FROM ml_reporte_mensual WHERE ml_lote_id=%s ORDER BY institucion_educativa_id,periodo_fuente",
        (lote_id,),
    ).fetchall()
    filas = copy.deepcopy(filas)
    for r in filas:
        r["periodo_fuente"] = str(r["periodo_fuente"])
    return publicar(
        conn,
        filas,
        lote["manifiesto"],
        inicio=inicio,
        autor=autor,
        motivo=motivo,
        padre_id=lote_id,
    )

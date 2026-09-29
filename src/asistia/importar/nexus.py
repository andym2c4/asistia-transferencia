"""Importador de cortes NEXUS (.xlsx moderno y .xls legado).

Identifica la fila de encabezados por contenido (busca "CODMOD I.E."), no por posicion fija -- los
dos formatos reales vistos (nexus 2026-06-01.xlsx con encabezado en fila 1; NEXUS LUYA AL
01-04-2026.xls con encabezado en fila 5) lo exigen. Ver docs/casos/TACTAMAL_JULIO_2026.md §3 y
docs/casos/IE_258_MEMBRILLO_JUNIO_2026.md §3 para la evidencia real de ambos formatos.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import psycopg
import xlrd

from asistia.db import registrar_documento

_ANCLA_ENCABEZADO = "CODMOD I.E."

_COLUMNAS_REQUERIDAS = [
    "CODMOD I.E.",
    "CODIGO DE PLAZA",
    "TIPO DE TRABAJADOR",
    "SITUACION LABORAL",
    "TIPO DE REGISTRO",
    "DOCUMENTO DE IDENTIDAD",
    "APELLIDO PATERNO",
    "APELLIDO MATERNO",
    "NOMBRES",
    "FECHA DE INICIO",
    "FECHA DE TERMINO",
    "CODIGO LOCAL",
    "DISTRITO",
    "NIVEL EDUCATIVO",
    "NOMBRE DE LA INSTITUCION EDUCATIVA",
]


@dataclass
class ResultadoImportacionNexus:
    nexus_carga_id: str
    total_filas: int = 0
    resueltas: int = 0
    vacantes: int = 0
    documento_invalido: int = 0
    errores: list[str] = field(default_factory=list)


def _texto_numerico(raw: Any) -> str:
    """xlrd devuelve celdas numericas como float de Python (p. ej. 1129998.0 para una celda que en
    Excel se ve como texto '1129998'). Convertir con str() directo deja el sufijo '.0', que luego
    corrompe cualquier recorte de digitos por posicion. Se normaliza a entero primero cuando el
    valor es un float sin parte decimal real."""
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    return str(raw)


def _normalizar_identificador(raw: Any, ancho: int) -> str | None:
    """Identificador numerico de ancho fijo, tolerante a ceros perdidos pero NO a valores largos.

    Los cortes reales entregan el mismo identificador en dos formas: el `.xls` conserva los ceros a
    la izquierda ('008628') y el `.xlsx` los pierde al pasar por float ('8628.0'). Rellenar con
    ceros recupera esos casos y es lo que mantiene una sola identidad para la misma persona o
    institucion.

    Lo que NO se hace es recortar. La version anterior terminaba en `zfill(ancho)[-ancho:]`, de modo
    que un valor con mas digitos significativos de los esperados se quedaba con los ultimos y
    producia **una identidad distinta de la real** -- un carne de extranjeria de 9 digitos se
    convertiria en el DNI de otra persona, en silencio y sin alerta (DT05). Ante un valor asi se
    devuelve None: la fila queda como identidad no resuelta, camino que ya existe, es visible en la
    bandeja de RRHH y no desaparece del reporte.

    Un valor que solo tiene ceros tampoco es un identificador: devuelve None en vez de fabricar
    '00000000', mismo criterio que ya aplicaba `_normalizar_codigo_local`.

    Caracterizado contra los tres cortes NEXUS reales (2026-09-14): de 3174 DNI, 3164 traen 8
    digitos y 9 traen 7 (ceros perdidos, que el relleno recupera); el unico de 9 digitos
    ('005960341') lo es por ceros a la izquierda y sigue normalizando a '05960341' igual que antes.
    De 3677 cod_mod, 636 traen 6 digitos y 3041 traen 7; ninguno excede. Es decir: este cambio no
    altera el resultado de ninguna fila del corpus disponible, solo cierra el camino de
    fabricacion.
    """
    if raw is None:
        return None
    digitos = re.sub(r"\D", "", _texto_numerico(raw))
    significativos = digitos.lstrip("0")
    if not significativos:
        return None
    if len(significativos) > ancho:
        return None
    return significativos.zfill(ancho)


def _normalizar_cod_mod(raw: Any) -> str | None:
    return _normalizar_identificador(raw, 7)


def _normalizar_dni(raw: Any) -> str | None:
    return _normalizar_identificador(raw, 8)


def _normalizar_codigo_local(raw: Any) -> str | None:
    """Codigo de local escolar (ESCALE), 6 digitos con ceros a la izquierda -- mismo problema que
    cod_mod/DNI: un corte real (.xlsx) lo entrega como float ('8628.0'), y str() directo pierde los
    ceros a la izquierda que el otro corte (.xls, '008628') sí conserva. Sin esta normalizacion, el
    mismo local fisico (ej. el IEPySM de Tactamal, primaria y secundaria) queda registrado como 2
    filas distintas de `local_educativo` solo por el formato del corte usado -- confirmado con
    datos reales: 316 codigos del corte .xls ya vienen con 6 digitos, 150 del corte .xlsx pierden
    1-2 ceros a la izquierda por el redondeo de float. Un codigo que normaliza a "000000" (celda
    vacia o literal 0) no es un local real -- se devuelve None en vez de inventar una identidad
    compartida falsa. Desde 2026-09-14 comparte regla con cod_mod/DNI (DT05): un codigo con mas de 6
    digitos significativos tampoco se recorta, se descarta."""
    return _normalizar_identificador(raw, 6)


def _normalizar_tipo_registro(raw: str | None) -> str | None:
    if not raw or not str(raw).strip():
        return None
    return str(raw).strip().upper().replace(" ", "_")


def _normalizar_situacion_laboral(raw: str | None) -> str | None:
    if not raw or not str(raw).strip():
        return None
    valor = str(raw).strip().upper()
    return (
        valor
        if valor
        in {
            "NOMBRADO",
            "CONTRATADO",
            "DESIGNADO",
            "VACANTE",
            "ENCARGADO",
            "DESTACADO",
        }
        else None
    )


_EXCEL_EPOCH = datetime(1899, 12, 30)  # noqa: DTZ001 - fecha civil de Excel, sin zona.


def _parsear_fecha(valor: Any, xls_datemode: int | None = None) -> date | None:
    """Acepta datetime (openpyxl), serial numerico (xlrd/openpyxl crudo) o texto 'M/D/YYYY'.
    Cualquier otra cosa (blanco, '-', texto no reconocible) es None -- nunca se adivina."""
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        if xls_datemode is not None:
            return xlrd.xldate.xldate_as_datetime(valor, xls_datemode).date()
        return (_EXCEL_EPOCH + __import__("datetime").timedelta(days=valor)).date()
    texto = str(valor).strip()
    if not texto or texto in {"-", "-   -", "  -   -"}:
        return None
    for formato in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato).date()  # noqa: DTZ007 - solo fecha civil.
        except ValueError:
            continue
    return None


def _filas_xlsx(ruta: Path) -> Iterator[dict[str, Any]]:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb[wb.sheetnames[0]]
    fila_encabezado = None
    encabezados: dict[str, int] = {}
    for fila in ws.iter_rows(min_row=1, max_row=min(20, ws.max_row)):
        valores = [c.value for c in fila]
        if _ANCLA_ENCABEZADO in valores:
            fila_encabezado = fila[0].row
            encabezados = {str(v).strip(): i for i, v in enumerate(valores) if v}
            break
    if fila_encabezado is None:
        raise ValueError(
            f"{ruta}: no se encontro la fila de encabezado (ancla {_ANCLA_ENCABEZADO!r})"
        )

    for row_idx in range(fila_encabezado + 1, ws.max_row + 1):
        valores = [c.value for c in ws[row_idx]]
        if all(v in (None, "") for v in valores):
            continue
        yield {
            nombre: (valores[idx] if idx < len(valores) else None)
            for nombre, idx in encabezados.items()
        }


def _filas_xls(ruta: Path) -> Iterator[dict[str, Any]]:
    wb = xlrd.open_workbook(str(ruta))
    sh = wb.sheet_by_index(0)
    fila_encabezado = None
    encabezados: dict[str, int] = {}
    for r in range(min(20, sh.nrows)):
        valores = [sh.cell_value(r, c) for c in range(sh.ncols)]
        if _ANCLA_ENCABEZADO in valores:
            fila_encabezado = r
            encabezados = {str(v).strip(): i for i, v in enumerate(valores) if v}
            break
    if fila_encabezado is None:
        raise ValueError(
            f"{ruta}: no se encontro la fila de encabezado (ancla {_ANCLA_ENCABEZADO!r})"
        )

    for r in range(fila_encabezado + 1, sh.nrows):
        valores = [sh.cell_value(r, c) for c in range(sh.ncols)]
        if all(v in (None, "") for v in valores):
            continue
        fila = {nombre: valores[idx] for nombre, idx in encabezados.items()}
        fila["__xls_datemode__"] = wb.datemode
        yield fila


def _leer_filas(ruta: Path) -> Iterator[dict[str, Any]]:
    if ruta.suffix.lower() == ".xlsx":
        yield from _filas_xlsx(ruta)
    elif ruta.suffix.lower() == ".xls":
        yield from _filas_xls(ruta)
    else:
        raise ValueError(f"formato no soportado para NEXUS: {ruta.suffix}")


def _obtener_o_crear_local(
    cur: psycopg.Cursor, codigo_local: str | None, distrito: str
) -> int:
    if codigo_local:
        cur.execute(
            "SELECT local_educativo_id FROM local_educativo WHERE codlocal_escale = %s",
            (codigo_local,),
        )
        fila = cur.fetchone()
        if fila:
            return fila["local_educativo_id"]
    cur.execute(
        """
        INSERT INTO local_educativo (codlocal_escale, distrito)
        VALUES (%s, %s) RETURNING local_educativo_id
        """,
        (codigo_local, distrito),
    )
    return cur.fetchone()["local_educativo_id"]


def _obtener_o_crear_institucion(
    cur: psycopg.Cursor,
    cod_mod: str,
    nombre_raw: str,
    nivel: str,
    distrito: str,
    local_educativo_id: int,
) -> int:
    cur.execute(
        "SELECT institucion_educativa_id FROM institucion_educativa WHERE cod_mod = %s AND anexo = '0'",
        (cod_mod,),
    )
    fila = cur.fetchone()
    if fila:
        return fila["institucion_educativa_id"]
    cur.execute(
        """
        INSERT INTO institucion_educativa (local_educativo_id, cod_mod, nombre_ie, nivel_modalidad, distrito)
        VALUES (%s, %s, %s, %s, %s)
        RETURNING institucion_educativa_id
        """,
        (
            local_educativo_id,
            cod_mod,
            nombre_raw or cod_mod,
            nivel or "SIN ESPECIFICAR",
            distrito,
        ),
    )
    return cur.fetchone()["institucion_educativa_id"]


def _obtener_o_crear_plaza(
    cur: psycopg.Cursor,
    codigo_plaza: str,
    institucion_educativa_id: int,
    cargo_raw: str | None,
) -> int:
    cur.execute("SELECT plaza_id FROM plaza WHERE codigo_plaza = %s", (codigo_plaza,))
    fila = cur.fetchone()
    if fila:
        return fila["plaza_id"]
    cur.execute(
        """
        INSERT INTO plaza (institucion_educativa_id, codigo_plaza, cargo_raw)
        VALUES (%s, %s, %s) RETURNING plaza_id
        """,
        (institucion_educativa_id, codigo_plaza, cargo_raw),
    )
    return cur.fetchone()["plaza_id"]


def _obtener_o_crear_trabajador(
    cur: psycopg.Cursor,
    dni: str,
    apellido_paterno: str,
    apellido_materno: str,
    nombres: str,
) -> int:
    cur.execute("SELECT trabajador_id FROM trabajador WHERE dni = %s", (dni,))
    fila = cur.fetchone()
    if fila:
        return fila["trabajador_id"]
    cur.execute(
        """
        INSERT INTO trabajador (dni, apellido_paterno, apellido_materno, nombres)
        VALUES (%s, %s, %s, %s) RETURNING trabajador_id
        """,
        (dni, apellido_paterno or "-", apellido_materno or "-", nombres or "-"),
    )
    return cur.fetchone()["trabajador_id"]


def _obtener_o_crear_vinculo(
    cur: psycopg.Cursor,
    trabajador_id: int | None,
    plaza_id: int,
    institucion_educativa_id: int,
    rol_laboral_id: int,
    situacion_laboral: str,
    tipo_registro: str,
    estado_raw: str | None,
    motivo_vacante_raw: str | None,
    fecha_inicio: date | None,
    fecha_fin: date | None,
) -> int | None:
    """Devuelve None cuando el rango de fechas declarado choca con otro vinculo ya existente
    para la misma plaza y tipo_registro (dos cortes NEXUS en desacuerdo) -- el llamador debe
    tratarlo como PENDIENTE de revision, nunca elegir una version en silencio."""
    # Idempotencia: mismo (plaza, tipo_registro, trabajador, rango de fechas) ya importado -> reusar.
    cur.execute(
        """
        SELECT vinculo_trabajador_ie_id FROM vinculo_trabajador_ie
        WHERE plaza_id = %s AND tipo_registro = %s
          AND trabajador_id IS NOT DISTINCT FROM %s
          AND fecha_inicio IS NOT DISTINCT FROM %s
          AND fecha_fin IS NOT DISTINCT FROM %s
        """,
        (plaza_id, tipo_registro, trabajador_id, fecha_inicio, fecha_fin),
    )
    fila = cur.fetchone()
    if fila:
        return fila["vinculo_trabajador_ie_id"]

    # Dos cortes NEXUS distintos pueden declarar rangos de fecha que se solapan para la misma
    # (plaza, tipo_registro) sin ser la misma fila (p. ej. una correccion administrativa entre
    # cortes). No se decide aqui cual version es correcta -- se deja pendiente de revision en vez
    # de forzar un INSERT que el EXCLUDE de la tabla rechazaria.
    cur.execute(
        """
        SELECT 1 FROM vinculo_trabajador_ie
        WHERE plaza_id = %s AND tipo_registro = %s
          AND daterange(COALESCE(fecha_inicio, '-infinity'::date), COALESCE(fecha_fin, 'infinity'::date), '[]')
              && daterange(COALESCE(%s::date, '-infinity'::date), COALESCE(%s::date, 'infinity'::date), '[]')
        """,
        (plaza_id, tipo_registro, fecha_inicio, fecha_fin),
    )
    if cur.fetchone():
        return None

    cur.execute(
        """
        INSERT INTO vinculo_trabajador_ie
            (trabajador_id, plaza_id, institucion_educativa_id, rol_laboral_id, situacion_laboral,
             tipo_registro, estado_raw, motivo_vacante_raw, fecha_inicio, fecha_fin)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING vinculo_trabajador_ie_id
        """,
        (
            trabajador_id,
            plaza_id,
            institucion_educativa_id,
            rol_laboral_id,
            situacion_laboral,
            tipo_registro,
            estado_raw,
            motivo_vacante_raw,
            fecha_inicio,
            fecha_fin,
        ),
    )
    return cur.fetchone()["vinculo_trabajador_ie_id"]


def importar_nexus(
    conn: psycopg.Connection, ruta: Path, fecha_corte: date
) -> ResultadoImportacionNexus:
    mime = (
        "application/vnd.ms-excel"
        if ruta.suffix.lower() == ".xls"
        else ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    )
    with conn.cursor() as cur:
        documento_recibido_id = registrar_documento(cur, ruta, mime)

        filas = list(_leer_filas(ruta))
        cur.execute(
            """
            INSERT INTO nexus_carga (documento_recibido_id, fecha_corte, total_filas)
            VALUES (%s, %s, %s) RETURNING nexus_carga_id
            """,
            (documento_recibido_id, fecha_corte, len(filas)),
        )
        nexus_carga_id = cur.fetchone()["nexus_carga_id"]
    conn.commit()

    resultado = ResultadoImportacionNexus(
        nexus_carga_id=nexus_carga_id, total_filas=len(filas)
    )

    for i, fila in enumerate(filas, start=1):
        with conn.cursor() as cur:
            try:
                _importar_fila(cur, nexus_carga_id, i, fila, resultado)
                conn.commit()
            except Exception as exc:  # noqa: BLE001 - se registra y sigue con la fila siguiente
                conn.rollback()
                resultado.errores.append(f"fila {i}: {exc}")

    with conn.cursor() as cur:
        cur.execute(
            """UPDATE nexus_carga SET
               filas_registradas=(SELECT count(*) FROM nexus_registro WHERE nexus_carga_id=%s),
               total_personas_validas=%s,total_vacantes=%s,total_documento_invalido=%s
               WHERE nexus_carga_id=%s""",
            (
                nexus_carga_id,
                resultado.resueltas,
                resultado.vacantes,
                resultado.documento_invalido,
                nexus_carga_id,
            ),
        )
        cur.execute(
            """UPDATE nexus_carga SET estado_integridad=CASE
               WHEN filas_registradas=total_filas THEN 'COMPLETA'
               WHEN filas_registradas=0 THEN 'FALLIDA' ELSE 'PARCIAL' END,
               estado=CASE WHEN filas_registradas=total_filas THEN 'CERRADA' ELSE 'PROCESANDO' END
               WHERE nexus_carga_id=%s""",
            (nexus_carga_id,),
        )
    conn.commit()
    return resultado


def _importar_fila(
    cur: psycopg.Cursor,
    nexus_carga_id: str,
    fila_origen: int,
    fila: dict[str, Any],
    resultado: ResultadoImportacionNexus,
) -> None:
    cod_mod_raw = fila.get("CODMOD I.E.")
    cod_mod = _normalizar_cod_mod(cod_mod_raw)
    dni = _normalizar_dni(fila.get("DOCUMENTO DE IDENTIDAD"))
    codigo_plaza = (
        _texto_numerico(fila.get("CODIGO DE PLAZA")).strip()
        if fila.get("CODIGO DE PLAZA") not in (None, "")
        else ""
    )
    situacion_laboral = _normalizar_situacion_laboral(fila.get("SITUACION LABORAL"))
    tipo_registro = _normalizar_tipo_registro(fila.get("TIPO DE REGISTRO"))
    tipo_trabajador_raw = str(fila.get("TIPO DE TRABAJADOR") or "").strip()
    xls_datemode = fila.get("__xls_datemode__")
    fecha_inicio = _parsear_fecha(fila.get("FECHA DE INICIO"), xls_datemode)
    fecha_fin = _parsear_fecha(fila.get("FECHA DE TERMINO"), xls_datemode)

    estado_resolucion = "PENDIENTE"
    institucion_educativa_id = None
    plaza_id = None
    trabajador_id = None
    vinculo_trabajador_ie_id = None

    if not cod_mod:
        estado_resolucion = "DOCUMENTO_INVALIDO"
        resultado.documento_invalido += 1
    else:
        local_educativo_id = _obtener_o_crear_local(
            cur,
            _normalizar_codigo_local(fila.get("CODIGO LOCAL")),
            str(fila.get("DISTRITO") or "").strip(),
        )
        institucion_educativa_id = _obtener_o_crear_institucion(
            cur,
            cod_mod,
            _texto_numerico(fila.get("NOMBRE DE LA INSTITUCION EDUCATIVA") or ""),
            str(fila.get("NIVEL EDUCATIVO") or ""),
            str(fila.get("DISTRITO") or ""),
            local_educativo_id,
        )

        if situacion_laboral == "VACANTE" or not dni:
            estado_resolucion = "VACANTE"
            resultado.vacantes += 1
            if codigo_plaza:
                plaza_id = _obtener_o_crear_plaza(
                    cur,
                    codigo_plaza,
                    institucion_educativa_id,
                    str(fila.get("CARGO") or ""),
                )
        else:
            cur.execute(
                "SELECT rol_laboral_id FROM catalogo_rol_laboral WHERE codigo = %s",
                (tipo_trabajador_raw,),
            )
            rol_fila = cur.fetchone()
            # Fechas invertidas: mismo bug real que v1 documento en TECH_DEBT.md (NEXUS, filas
            # POR_REEMPLAZO con ciertos motivos de vacante traen fecha_inicio > fecha_fin). Se
            # marca PENDIENTE para revision manual en vez de inventar cual fecha esta mal.
            fechas_invertidas = fecha_inicio and fecha_fin and fecha_inicio > fecha_fin
            if (
                not rol_fila
                or not codigo_plaza
                or not situacion_laboral
                or not tipo_registro
                or fechas_invertidas
            ):
                estado_resolucion = "PENDIENTE"
            else:
                plaza_id = _obtener_o_crear_plaza(
                    cur,
                    codigo_plaza,
                    institucion_educativa_id,
                    str(fila.get("CARGO") or ""),
                )
                trabajador_id = _obtener_o_crear_trabajador(
                    cur,
                    dni,
                    str(fila.get("APELLIDO PATERNO") or ""),
                    str(fila.get("APELLIDO MATERNO") or ""),
                    str(fila.get("NOMBRES") or ""),
                )
                vinculo_trabajador_ie_id = _obtener_o_crear_vinculo(
                    cur,
                    trabajador_id,
                    plaza_id,
                    institucion_educativa_id,
                    rol_fila["rol_laboral_id"],
                    situacion_laboral,
                    tipo_registro,
                    str(fila.get("ESTADO") or "") or None,
                    str(fila.get("MOTIVO DE VACANTE") or "") or None,
                    fecha_inicio,
                    fecha_fin,
                )
                if vinculo_trabajador_ie_id is None:
                    estado_resolucion = "PENDIENTE"
                else:
                    estado_resolucion = "RESUELTO"
                    resultado.resueltas += 1

    cur.execute(
        """
        INSERT INTO nexus_registro
            (nexus_carga_id, fila_origen, cod_mod_ie_raw, institucion_educativa_id, codigo_plaza_raw,
             plaza_id, documento_identidad_raw, trabajador_id, vinculo_trabajador_ie_id,
             tipo_trabajador_raw, situacion_laboral_raw, tipo_registro_raw, fecha_inicio_raw,
             fecha_termino_raw, estado_resolucion)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            nexus_carga_id,
            fila_origen,
            # `_raw` conserva lo recibido sin recortar (migracion 0011, DT05): si el valor
            # no normaliza, es justo cuando hace falta saber que traia el archivo.
            cod_mod or (str(cod_mod_raw) if cod_mod_raw is not None else None),
            institucion_educativa_id,
            codigo_plaza or None,
            plaza_id,
            str(fila.get("DOCUMENTO DE IDENTIDAD") or ""),
            trabajador_id,
            vinculo_trabajador_ie_id,
            tipo_trabajador_raw or None,
            str(fila.get("SITUACION LABORAL") or "") or None,
            str(fila.get("TIPO DE REGISTRO") or "") or None,
            str(fila.get("FECHA DE INICIO") or "") or None,
            str(fila.get("FECHA DE TERMINO") or "") or None,
            estado_resolucion,
        ),
    )

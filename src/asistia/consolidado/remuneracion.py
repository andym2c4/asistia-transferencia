"""Cruce diario verificable; clasificar días no ejecuta pagos ni descuentos."""

import calendar
from collections import Counter
from datetime import date

from psycopg.types.json import Jsonb

from asistia.ajustes_asistencia import ajuste_del_dia, marca_efectiva
from asistia.leyendas import categoria_dia

from .universo_esperado import calendario_aplicable

VERSION = "CRUCE_REMUNERACION_2"


def cruzar_dia(calendario, asistencia, *, falta=None, impedimento=None):
    if impedimento:
        return {
            "resultado": "PENDIENTE",
            "es_remunerado": None,
            "es_falta": None,
            "motivo": impedimento,
        }
    if type(calendario) is not bool or type(asistencia) is not bool:
        return {
            "resultado": "PENDIENTE",
            "es_remunerado": None,
            "es_falta": None,
            "motivo": "Falta clasificar la remuneración o completar la fuente diaria.",
        }
    if not calendario and asistencia:
        return {
            "resultado": "OBSERVACION",
            "es_remunerado": None,
            "es_falta": False,
            "motivo": "Asistencia remunerada en un día no remunerado del calendario; revisar el cruce.",
        }
    pagado = calendario and asistencia
    return {
        "resultado": "REMUNERADO" if pagado else "NO_REMUNERADO",
        "es_remunerado": pagado,
        "es_falta": (falta if calendario and not asistencia else False),
        "motivo": "Ambas categorías son remuneradas."
        if pagado
        else (
            "Ausencia en día remunerado."
            if calendario and falta is True
            else "La categoría de asistencia no es remunerada."
            if calendario
            else "Ambas categorías son no remuneradas."
        ),
    }


def cruce_persona(conn, tid, *, calendario_id=None, dia_propuesto=None):
    p = conn.execute(
        """SELECT t.*,r.periodo,r.institucion_educativa_id,r.clasificacion_codigos,r.estado AS estado_reporte
        FROM trabajador_en_reporte t JOIN reporte_asistencia r USING(reporte_asistencia_id)
        WHERE trabajador_en_reporte_id=%s""",
        (tid,),
    ).fetchone()
    if not p:
        raise ValueError("La fila de asistencia no existe.")
    if calendario_id is None:
        with conn.cursor() as cur:
            calendario_id, _ = calendario_aplicable(
                cur, p["institucion_educativa_id"], p["periodo"].year
            )
    cv = (
        conn.execute(
            "SELECT * FROM calendarizacion_version WHERE calendarizacion_version_id=%s",
            (calendario_id,),
        ).fetchone()
        if calendario_id
        else None
    )
    calendario_dias = {
        d["fecha"]: d
        for d in conn.execute(
            "SELECT * FROM dia_calendarizacion WHERE calendarizacion_version_id=%s AND date_trunc('month',fecha)=%s",
            (calendario_id, p["periodo"]),
        ).fetchall()
    }
    asistencia_dias = {
        d["fecha"]: d
        for d in conn.execute(
            "SELECT * FROM asistencia_dia WHERE trabajador_en_reporte_id=%s", (tid,)
        ).fetchall()
    }
    if dia_propuesto is not None:
        anterior = asistencia_dias.get(dia_propuesto["fecha"], {})
        if anterior.get("asistencia_dia_id") != dia_propuesto["asistencia_dia_id"]:
            raise ValueError("El día propuesto no pertenece a esta fila de asistencia.")
        asistencia_dias[dia_propuesto["fecha"]] = dia_propuesto
    total = calendar.monthrange(p["periodo"].year, p["periodo"].month)[1]
    vigencias = (
        {
            d["fecha"]: d["esperado"]
            for d in conn.execute(
                """SELECT %s::date+n AS fecha,
        fn_vinculo_presencia_esperada(%s,%s::date+n) AS esperado FROM generate_series(0,%s) n""",
                (p["periodo"], p["vinculo_trabajador_ie_id"], p["periodo"], total - 1),
            ).fetchall()
        }
        if p["vinculo_trabajador_ie_id"]
        else {}
    )
    dias = []
    for numero in range(1, total + 1):
        fecha = date(p["periodo"].year, p["periodo"].month, numero)
        cd, ad = calendario_dias.get(fecha, {}), asistencia_dias.get(fecha, {})
        ajuste = ajuste_del_dia(p, ad)
        ad = marca_efectiva(p, ad)
        cc, ca = categoria_dia(cv or {}, cd), categoria_dia(p, ad)
        from asistia.calidad_leyendas import pendiente

        no_aplica_confirmado = ad.get("estado_captura") == "NO_APLICA" and (
            ad.get("validado") or bool(ajuste and ajuste["vigente"])
        )
        impedimento = None
        if ajuste and not ajuste["vigente"]:
            impedimento = (
                "La declaración cambió después del ajuste de RRHH; vuelve a cotejarlo."
            )
        elif not cv:
            impedimento = "No hay calendario para el año."
        elif cv["procedencia_extraccion"].get("tipo") == "DERIVADO_2025":
            impedimento = (
                "Calendario basado en 2025: no acredita la remuneración de 2026."
            )
        elif not p["vinculo_trabajador_ie_id"]:
            impedimento = "Falta resolver el vínculo del trabajador."
        elif not vigencias.get(fecha):
            impedimento = "El vínculo no tiene presencia esperada en esta fecha; revisar su aplicabilidad remunerativa."
        elif not no_aplica_confirmado and (
            ad.get("evidencia_interpretacion", {}).get("dato_no_determinado")
            or not (ad.get("codigo_interpretado") or ad.get("codigo_reportado_raw"))
        ):
            impedimento = "La asistencia diaria está vacía o no determinada."
        elif not cd or not (
            cd.get("codigo_interpretado") or cd.get("codigo_reportado_raw")
        ):
            impedimento = "Falta el código de calendario del día."
        elif (not no_aplica_confirmado and pendiente(ca)) or pendiente(cc):
            impedimento = (
                "Posible desfase de leyenda: interpretación pendiente de revisión."
            )
        resultado = cruzar_dia(
            cc.get("es_remunerado"),
            ca.get("es_remunerado"),
            falta=ca.get("es_falta"),
            impedimento=impedimento,
        )
        if no_aplica_confirmado and not impedimento:
            from asistia.monitoreo.reglas import actividad

            actividad_calendario = actividad(cc)
            sin_actividad = actividad_calendario is False
            resultado = {
                "resultado": "NO_APLICA"
                if sin_actividad
                else "OBSERVACION"
                if actividad_calendario is True
                else "PENDIENTE",
                "es_remunerado": None,
                "es_falta": False if sin_actividad else None,
                "motivo": "No correspondía asistir; no cuenta como falta ni determina remuneración."
                if sin_actividad
                else "No correspondía asistir frente a actividad programada; revisar."
                if actividad_calendario is True
                else "Falta clasificar la actividad del calendario para cotejar si correspondía asistir.",
            }
        dias.append(
            {
                "fecha": str(fecha),
                **({"ajuste_rrhh": ajuste} if ajuste else {}),
                "calendario": {
                    "codigo": cd.get("codigo_interpretado")
                    or cd.get("codigo_reportado_raw"),
                    **cc,
                },
                "asistencia": {
                    "codigo": None
                    if ad.get("estado_captura") == "NO_APLICA"
                    else (
                        ad.get("codigo_interpretado") or ad.get("codigo_reportado_raw")
                    ),
                    **ca,
                },
                "codigo_recibido": ad.get("codigo_reportado_raw"),
                "fuente_asistencia": ad.get("celda_origen"),
                "fuente_calendario": cd.get("celda_origen"),
                "evidencia_interpretacion": ad.get("evidencia_interpretacion", {}),
                **resultado,
            }
        )
    conteo = Counter(d["resultado"] for d in dias)
    return {
        "regla": VERSION,
        "periodo": str(p["periodo"]),
        "unidad": "día por vínculo y fila seleccionada",
        "calendario_version_id": str(calendario_id) if calendario_id else None,
        "estado_calendario": cv["estado"] if cv else None,
        "estado_reporte": p["estado_reporte"],
        "reporte_id": str(p["reporte_asistencia_id"]),
        "vinculo_id": p["vinculo_trabajador_ie_id"],
        "dias": dias,
        "resumen": {
            k: conteo[k]
            for k in (
                "REMUNERADO",
                "NO_REMUNERADO",
                "OBSERVACION",
                "PENDIENTE",
                "NO_APLICA",
            )
        },
        "faltas_en_dia_remunerado": sum(d["es_falta"] is True for d in dias),
        "total_dias": total,
        "alcance": "Resultado técnico según las categorías de las fuentes; no es una orden de pago ni un descuento aprobado.",
    }


def congelar_cruces(conn, consolidado_id):
    filas = conn.execute(
        """SELECT d.consolidado_dre_detalle_id,f.trabajador_en_reporte_id FROM consolidado_dre_detalle d
        JOIN consolidado_dre_detalle_fuente f USING(consolidado_dre_detalle_id) WHERE d.consolidado_dre_id=%s""",
        (consolidado_id,),
    ).fetchall()
    for f in filas:
        cruce = cruce_persona(conn, f["trabajador_en_reporte_id"])
        justificadas, injustificadas, sin_clasificar = [], [], []
        for d in cruce["dias"]:
            if d["es_falta"] is not True:
                continue
            codigo = d["asistencia"].get("estado_asistencia_codigo")
            if codigo in {"I", "3T", "H"}:
                injustificadas.append(d["fecha"])
            elif codigo == "J":
                justificadas.append(d["fecha"])
            else:
                sin_clasificar.append(d["fecha"])
        pendientes = (
            sum(
                d["es_falta"] is None and d["resultado"] == "NO_REMUNERADO"
                for d in cruce["dias"]
            )
            + cruce["resumen"]["PENDIENTE"]
            + cruce["resumen"]["OBSERVACION"]
            + len(sin_clasificar)
        )
        clasificacion = {
            "justificadas": justificadas,
            "injustificadas": injustificadas,
            "sin_clasificar": sin_clasificar,
            "dias_pendientes": pendientes,
            "estado": "CALCULABLE"
            if not pendientes
            else "PARCIAL"
            if justificadas or injustificadas
            else "NO_CALCULABLE",
            "regla": VERSION,
        }
        conn.execute(
            "UPDATE consolidado_dre_detalle SET fuente_calculo=fuente_calculo||%s WHERE consolidado_dre_detalle_id=%s",
            (
                Jsonb(
                    {"cruce_remuneracion": cruce, "clasificacion_faltas": clasificacion}
                ),
                f["consolidado_dre_detalle_id"],
            ),
        )


def exportar_cruces(wb, personas):
    hoja = wb.create_sheet("Cruce de remuneración")
    hoja.append(
        [
            "Resultado técnico; revisar antes de decidir pago o descuento. Los vacíos y observaciones no equivalen a cero."
        ]
    )
    hoja.append(
        [
            "Institución",
            "Persona",
            "Fila de fuente",
            "Fecha",
            "Tipo de día: calendario",
            "Remuneración: calendario",
            "Tipo de día: asistencia",
            "Remuneración: asistencia",
            "Resultado del cruce",
            "Falta en día remunerado",
            "Motivo",
            "Versión calendario",
            "Reporte",
            "Fuente diaria",
        ]
    )

    def etiqueta(v):
        return (
            "Remunerado"
            if v is True
            else "No remunerado"
            if v is False
            else "Pendiente"
        )

    for p in personas:
        cruce = p["fuente_calculo"].get("cruce_remuneracion", {})
        for d in cruce.get("dias", []):
            hoja.append(
                [
                    p["nombre_ie"],
                    " ".join(
                        p.get(k) or ""
                        for k in ("apellido_paterno", "apellido_materno", "nombres")
                    ),
                    str(p["trabajador_en_reporte_id"]),
                    d["fecha"],
                    d["calendario"].get("tipo_dia", "Sin leyenda"),
                    etiqueta(d["calendario"].get("es_remunerado")),
                    d["asistencia"].get("tipo_dia", "Sin leyenda"),
                    etiqueta(d["asistencia"].get("es_remunerado")),
                    d["resultado"],
                    "Sí"
                    if d["es_falta"] is True
                    else "No"
                    if d["es_falta"] is False
                    else "Pendiente",
                    d["motivo"],
                    cruce["calendario_version_id"],
                    cruce["reporte_id"],
                    d["fuente_asistencia"],
                ]
            )
    hoja.freeze_panes = "E3"
    hoja.auto_filter.ref = f"A2:N{hoja.max_row}"
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    for c in hoja[2]:
        c.font = Font(bold=True)
    for i in range(1, 15):
        hoja.column_dimensions[get_column_letter(i)].width = 30 if i != 11 else 65
    for fila in hoja:
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            # Texto extraído de documentos nunca se convierte en fórmula.
            if celda.data_type == "f":
                celda.data_type = "s"

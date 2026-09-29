"""Control estructural de leyendas: reconocer significados no valida su alineación."""

from __future__ import annotations

import copy
import re

from asistia.leyendas import descripcion_de_categoria, normal

VERSION = "ALINEACION_LEYENDA_EXCEL_1"
REGLA_ALERTA = "LEYENDA_POSIBLE_DESFASE"
MENSAJE = (
    "Posible desfase de leyenda: hay descripciones y códigos sin pareja en los "
    "extremos. Revisa las celdas del original antes de interpretar la asistencia."
)


class HojaLeyendaNoDeterminada(ValueError):
    """No hay un localizador inequívoco para aplicar el control estructural."""


def referencias(conn, dominio="asistencia"):
    from asistia.categorias import catalogo, equivalencias

    return {normal(c["nombre"]) for c in catalogo(conn, dominio)} | {
        normal(e["descripcion"]) for e in equivalencias(conn, dominio)
    }


def detectar_desfases(ws, conocidas=()):
    """Desplazamiento de una fila, >=3 códigos contiguos y >=3 descripciones.

    No usa una correspondencia universal letra/significado. Ambos extremos deben
    quedar sin pareja en la lectura por fila y alinearse con el mismo desplazamiento.
    Celdas combinadas que cruzan filas no permiten asegurar esa geometría.
    """
    conocidas = {normal(t) for t in conocidas}
    filas = list(
        ws.iter_rows(
            max_row=min(ws.max_row or 1500, 1500),
            max_col=min(ws.max_column or 120, 120),
        )
    )
    valores = {
        (c.row, c.column): c for row in filas for c in row if c.value not in (None, "")
    }
    codigos = {}
    for (f, col), c in valores.items():
        if re.fullmatch(r"[A-Za-z0-9./+-]{1,4}", str(c.value).strip()):
            codigos.setdefault(col, []).append(f)
    hallazgos = []
    for col, posiciones in codigos.items():
        bloques = []
        for f in sorted(posiciones):
            if bloques and f == bloques[-1][-1] + 1:
                bloques[-1].append(f)
            else:
                bloques.append([f])
        for bloque in bloques:
            if len(bloque) < 3:
                continue
            inicio, fin = bloque[0], bloque[-1]
            for dc in range(col + 1, min(col + 5, len(filas[0])) + 1):
                for desplazamiento in (-1, 1):
                    fs = [f + desplazamiento for f in bloque]
                    textos = [valores.get((f, dc)) for f in fs]
                    if not all(textos):
                        continue
                    reconocidas = [
                        c
                        for c in textos
                        if normal(c.value) in conocidas
                        or descripcion_de_categoria(c.value)
                    ]
                    if len(reconocidas) < 3 or len(reconocidas) / len(textos) < 0.75:
                        continue
                    extremo_codigo = fin if desplazamiento == -1 else inicio
                    extremo_texto = inicio - 1 if desplazamiento == -1 else fin + 1
                    celda_extremo = valores.get((extremo_texto, col))
                    if (extremo_codigo, dc) in valores or (
                        celda_extremo and "LEYENDA" not in normal(celda_extremo.value)
                    ):
                        continue
                    # Evitar saltar una columna intermedia con otra tabla/código.
                    if any(
                        (f, intermedia) in valores
                        for f in bloque
                        for intermedia in range(col + 1, dc)
                    ):
                        continue
                    if any(
                        r.min_row != r.max_row
                        and r.min_row <= max(fs + bloque)
                        and r.max_row >= min(fs + bloque)
                        and r.min_col <= dc
                        and r.max_col >= col
                        for r in ws.merged_cells.ranges
                    ):
                        continue
                    pares = []
                    for f, texto in zip(bloque, textos):
                        celda = valores[(f, col)]
                        recibida = valores.get((f, dc))
                        pares.append(
                            {
                                "codigo": str(celda.value).strip(),
                                "celda_codigo": celda.coordinate,
                                "descripcion_misma_fila": str(recibida.value).strip()
                                if recibida
                                else None,
                                "celda_misma_fila": ws.cell(f, dc).coordinate,
                                "descripcion_desplazada": str(texto.value).strip(),
                                "celda_desplazada": texto.coordinate,
                            }
                        )
                    hallazgos.append(
                        {
                            "regla": VERSION,
                            "tipo": "POSIBLE_DESFASE_VERTICAL",
                            "hoja": ws.title,
                            "rango": f"{ws.cell(min(fs + bloque), col).coordinate}:{ws.cell(max(fs + bloque), dc).coordinate}",
                            "desplazamiento_descripcion": desplazamiento,
                            "mensaje": MENSAJE,
                            "pares": pares,
                            "descripciones_reconocidas": len(reconocidas),
                            "total_descripciones": len(textos),
                        }
                    )
                    break
                if (
                    hallazgos
                    and hallazgos[-1]["pares"][0]["celda_codigo"]
                    == ws.cell(inicio, col).coordinate
                ):
                    break
    return hallazgos


def marcar_entradas(entradas, hallazgos):
    salida = copy.deepcopy(entradas)
    for h in hallazgos:
        for par in h["pares"]:
            existentes = [
                e
                for e in salida
                if e["codigo"] == par["codigo"] and e.get("hoja") == h["hoja"]
            ]
            if not existentes:
                existentes = [
                    {
                        "codigo": par["codigo"],
                        "descripcion": par["descripcion_misma_fila"]
                        or "Significado sin alinear",
                        "hoja": h["hoja"],
                        "celda": par["celda_codigo"] + ":" + par["celda_misma_fila"],
                    }
                ]
                salida.extend(existentes)
            for e in existentes:
                e["revision_leyenda"] = copy.deepcopy(h)
    return salida


def pendiente(categoria):
    return (
        bool(categoria.get("revision_leyenda"))
        and categoria.get("fuente") != "REVISION_WEB"
    )


def suspender(categoria, hallazgo):
    c = copy.deepcopy(categoria)
    if c.get("fuente") == "REVISION_WEB":
        return c
    if not pendiente(c):
        c["clasificacion_previa"] = copy.deepcopy(categoria)
    c.update(
        fuente="LEYENDA_POR_REVISAR",
        revision_leyenda=copy.deepcopy(hallazgo),
        es_remunerado=None,
        es_falta=None,
        grupo_actividad=None,
        estado_asistencia_codigo=None,
    )
    return c


def registrar_alerta(conn, rid, categorias):
    from psycopg.types.json import Jsonb

    afectados = {
        codigo: c["revision_leyenda"]
        for codigo, c in categorias.items()
        if pendiente(c)
    }
    if afectados:
        conn.execute(
            """INSERT INTO validacion_reporte(reporte_asistencia_id,codigo_regla,severidad,mensaje,evidencia)
            SELECT %s,%s,'ERROR',%s,%s WHERE NOT EXISTS (
              SELECT 1 FROM validacion_reporte WHERE reporte_asistencia_id=%s AND codigo_regla=%s AND estado='PENDIENTE')""",
            (rid, REGLA_ALERTA, MENSAJE, Jsonb(afectados), rid, REGLA_ALERTA),
        )


def cotejar_excel(ruta, hoja, categorias, conocidas=()):
    """No confundir un número de página con un índice de hoja.

    Usar nombre explícito, hoja única de la evidencia o un único ANEXO 3.
    """
    import openpyxl

    from asistia.leyendas import leer_leyenda_xlsx, normalizar_leyenda

    libro = openpyxl.load_workbook(ruta, data_only=True)
    try:
        if hoja not in libro.sheetnames:
            hojas = {
                c.get("evidencia", {}).get("hoja")
                for c in categorias.values()
                if c.get("evidencia", {}).get("hoja") in libro.sheetnames
            }
            if len(hojas) != 1:
                hojas = {
                    h
                    for h in libro.sheetnames
                    if re.fullmatch(r"ANEXO\s*[-_ ]?\s*3", normal(h))
                }
            if len(hojas) != 1:
                raise HojaLeyendaNoDeterminada(
                    "No se pudo localizar de forma única la hoja de asistencia del Excel. Revisa su fuente."
                )
            hoja = next(iter(hojas))
        extraidas = normalizar_leyenda(
            leer_leyenda_xlsx(libro[hoja], conocidas=conocidas), "asistencia"
        )
    finally:
        libro.close()
    despues = copy.deepcopy(categorias)
    for codigo, c in extraidas.items():
        if pendiente(c):
            despues[codigo] = suspender(
                categorias.get(codigo, c), c["revision_leyenda"]
            )
    return despues

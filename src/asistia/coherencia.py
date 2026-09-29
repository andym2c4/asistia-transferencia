"""Cruce operativo por versiones guardadas. Leer no altera documentos ni decisiones."""

from collections import Counter, defaultdict

from asistia.experimental.datos import huella
from asistia.monitoreo.datos import cargar_fuentes, construir
from asistia.monitoreo.reglas import REGLAS, VERSION

CONTRATO = "REVISION_COHERENCIA_1"
MOTIVOS = {
    "AJUSTE_RRHH_OBSOLETO": "La declaración cambió después del ajuste de RRHH; vuelve a cotejarlo.",
    "SIN_CALENDARIO": "Falta la calendarización de esta institución y año.",
    "CALENDARIO_DERIVADO_2025": "El calendario es una proyección; falta la fuente del año.",
    "CALENDARIO_SIN_CONFIRMAR": "Falta confirmar la calendarización del año.",
    "CATEGORIA_CALENDARIO_PENDIENTE": "Hay fechas sin actividad determinada en el calendario.",
    "CATEGORIA_ASISTENCIA_PENDIENTE": "Hay fechas sin declaración interpretable en el reporte.",
    "LEYENDA_POR_REVISAR": "Hay significados de símbolos pendientes de revisión.",
    "FUENTES_DIARIAS_CONTRADICTORIAS": "Las fuentes declaran datos diferentes para una misma persona y fecha.",
    "IDENTIDAD_O_ROL_PENDIENTE": "Falta resolver la identidad o el rol del personal.",
    "VIGENCIA_PENDIENTE": "Falta resolver la vigencia del vínculo.",
    "SIN_VIGENCIA_ESPERADA": "Fechas excluidas por vigencia del vínculo.",
}
DECISIONES = {
    "VERIFICADO": "Caso verificado por RRHH: conservar el valor actual",
    "PENDIENTE": "Requiere aclaración: mantener pendiente",
}


def revisar_reportes(conn, reportes):
    """Cruce operativo con ajustes RRHH; conserva la señal documental usada por IF."""
    por_anio = defaultdict(list)
    for r in reportes:
        por_anio[r["periodo"].year].append(str(r["reporte_asistencia_id"]))
    salida = {}
    for anio, ids in por_anio.items():
        fuentes = cargar_fuentes(conn, anio, reporte_ids=ids)
        perfiles = construir(fuentes, reglas_aplicadas=True, exigir_vigente=True)
        declarados = {
            (p["institucion_educativa_id"], p["periodo"], d["clave"]): d
            for p in perfiles
            for d in p["datos"]["dias"]
        }
        if any(
            d.get("evidencia_interpretacion", {}).get("ajuste_rrhh")
            for d in fuentes["dias"]
        ):
            perfiles = construir(
                fuentes, reglas_aplicadas=True, exigir_vigente=True, ajustes_rrhh=True
            )
        fuentes_ids = [r["reporte_asistencia_id"] for r in fuentes["reportes"]]
        eventos = conn.execute(
            """SELECT e.posterior,e.usuario_id,e.creado_en,e.web_revision_evento_id,u.nombre AS autor
            FROM web_revision_evento e JOIN usuario u USING(usuario_id)
            WHERE e.reporte_asistencia_id=ANY(%s) AND e.accion='REVISAR_COHERENCIA'
            ORDER BY e.creado_en,e.web_revision_evento_id""",
            (fuentes_ids,),
        ).fetchall()
        decisiones = {}
        for e in eventos:
            for clave in e["posterior"].get("casos", []):
                decisiones[clave] = {
                    "estado": e["posterior"]["decision"],
                    "autor": e["autor"],
                    "fecha": e["creado_en"],
                    "evento": str(e["web_revision_evento_id"]),
                }
        for p in perfiles:
            cv = p["fuentes"]["calendario"]
            casos = []
            for d in p["datos"]["dias"]:
                declarado = declarados[
                    (p["institucion_educativa_id"], p["periodo"], d["clave"])
                ]
                ajustes = [
                    f["ajuste_rrhh"] for f in d["fuentes"] if f.get("ajuste_rrhh")
                ]
                if not d["alertas"] and not declarado["alertas"] and not ajustes:
                    continue
                clave = huella(
                    {
                        "contrato": CONTRATO,
                        "reglas": VERSION,
                        "dia": d,
                        "calendario": str(cv["calendarizacion_version_id"])
                        if cv
                        else None,
                    }
                )
                decision = decisiones.get(clave)
                casos.append(
                    {
                        **d,
                        "id": clave,
                        "decision": decision,
                        "declarado": declarado,
                        "ajustes": ajustes,
                        "resuelto": bool(
                            decision["estado"] == "VERIFICADO"
                            if decision
                            else ajustes
                            and all(a["vigente"] for a in ajustes)
                            and d["cruce_evaluable"]
                            and not d["alertas"]
                        ),
                        "reglas": [
                            {"codigo": cod, "nombre": REGLAS[cod]}
                            for cod in (d["alertas"] or declarado["alertas"])
                        ],
                    }
                )
            for rid in ids:
                propios = [
                    d
                    for d in p["datos"]["dias"]
                    if any(f["reporte_id"] == rid for f in d["fuentes"])
                ]
                if not propios:
                    continue
                ids_dias = {d["clave"] for d in propios}
                casos_reporte = sorted(
                    [c for c in casos if c["clave"] in ids_dias],
                    key=lambda c: (c["fecha"], c["clave"]),
                )
                impedimentos = Counter(
                    d["pendiente"]
                    for d in propios
                    if d["pendiente"] and d["exclusion"] != "SIN_VIGENCIA_ESPERADA"
                )
                pendientes = sum(not c["resuelto"] for c in casos_reporte)
                estado = {
                    "contrato": CONTRATO,
                    "reglas_version": VERSION,
                    "casos": casos_reporte,
                    "pendientes": pendientes,
                    "resueltos": len(casos_reporte) - pendientes,
                    "alertas_por_regla": dict(
                        Counter(cod for c in casos_reporte for cod in c["alertas"])
                    ),
                    "no_evaluables": sum(impedimentos.values()),
                    "impedimentos": [
                        {"codigo": k, "motivo": MOTIVOS.get(k, k), "dias": n}
                        for k, n in sorted(impedimentos.items())
                    ],
                    "excluidos": sum(
                        d["exclusion"] == "SIN_VIGENCIA_ESPERADA" for d in propios
                    ),
                    "oportunidades": len(propios),
                    "evaluables": sum(d["cruce_evaluable"] for d in propios),
                    "calendario": {
                        k: cv[k]
                        for k in ("calendarizacion_version_id", "version", "estado")
                    }
                    if cv
                    else None,
                    "listo": not pendientes and not impedimentos,
                }
                estado["huella"] = huella(
                    {
                        "contrato": CONTRATO,
                        "reglas": VERSION,
                        "fuentes": propios,
                        "calendario": estado["calendario"],
                        "decisiones": [
                            (
                                c["id"],
                                c["decision"]["evento"] if c["decision"] else None,
                            )
                            for c in casos_reporte
                        ],
                    }
                )
                salida[rid] = estado
        for rid in ids:
            if rid not in salida:
                salida[rid] = {
                    "contrato": CONTRATO,
                    "reglas_version": VERSION,
                    "casos": [],
                    "pendientes": 0,
                    "resueltos": 0,
                    "alertas_por_regla": {},
                    "no_evaluables": 0,
                    "excluidos": 0,
                    "oportunidades": 0,
                    "evaluables": 0,
                    "calendario": None,
                    "listo": False,
                    "huella": huella([CONTRATO, rid, "SIN_DATOS"]),
                    "impedimentos": [
                        {
                            "codigo": "SIN_DATOS",
                            "motivo": "Sin filas diarias actuales para comprobar este reporte.",
                            "dias": 0,
                        }
                    ],
                }
    return salida


def exportar_control(wb, coherencias):
    """Anexo del corte exportado; señales originales y decisión separadas."""
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet("Control calendario")
    ws.append(["COHERENCIA REPORTE / CALENDARIZACIÓN · " + CONTRATO])
    ws.append(
        [
            "Las alertas requieren revisión; no prueban presencia física ni deciden descuentos. La declaración se conserva."
        ]
    )
    ws.append(
        [
            "Reporte",
            "Calendario",
            "Fecha",
            "Persona/rol (ID)",
            "Declarado",
            "Programado",
            "Señales",
            "Estado",
            "Revisor",
            "Valor ajustado por RRHH",
        ]
    )
    for rid, c in coherencias.items():
        cv = c["calendario"]
        version = f"v{cv['version']} · {cv['estado']}" if cv else "Sin calendario"
        for caso in c["casos"]:
            nombres = sorted(
                {
                    f.get("codigo_interpretado")
                    or f.get("codigo_recibido")
                    or f.get("estado_captura")
                    or "Sin marca"
                    for f in caso.get("declarado", caso)["fuentes"]
                }
            )
            decision = caso["decision"]
            ws.append(
                [
                    rid,
                    version,
                    caso["fecha"],
                    f"{caso['trabajador_id']} / {caso['rol_id']}",
                    ", ".join(nombres),
                    caso["calendario"].get("tipo_dia"),
                    " · ".join(r["nombre"] for r in caso["reglas"]),
                    DECISIONES[decision["estado"]]
                    if decision
                    else (
                        "Ajustado y sin diferencias"
                        if caso["resuelto"]
                        else "Pendiente"
                    ),
                    decision["autor"]
                    if decision
                    else ", ".join(a["autor"] for a in caso.get("ajustes", [])),
                    "; ".join(a["etiqueta"] for a in caso.get("ajustes", [])),
                ]
            )
        for impedimento in c["impedimentos"]:
            ws.append(
                [
                    rid,
                    version,
                    None,
                    None,
                    None,
                    None,
                    impedimento["motivo"],
                    f"Sin evaluar: {impedimento['dias']} oportunidades",
                    None,
                ]
            )
        if not c["casos"] and not c["impedimentos"]:
            ws.append(
                [
                    rid,
                    version,
                    None,
                    None,
                    None,
                    None,
                    "Sin señales en los cruces evaluables",
                    "Comprobado",
                    None,
                ]
            )
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {
        "A": 38,
        "B": 24,
        "C": 14,
        "D": 22,
        "E": 24,
        "J": 32,
        "F": 30,
        "G": 65,
        "H": 50,
        "I": 26,
    }.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:J{ws.max_row}"

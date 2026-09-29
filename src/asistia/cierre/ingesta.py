"""Lote documental real. Cada archivo conserva estado, resultado y fuente al reintentar."""

from __future__ import annotations

import calendar
import hashlib
import json
import mimetypes
import re
from collections import Counter
from dataclasses import asdict
from datetime import UTC, date, datetime
from pathlib import Path

import openpyxl
from psycopg.types.json import Jsonb

from asistia.db import registrar_documento, sha256_de
from asistia.identificacion import MESES, resolver_institucion, sin_acentos
from asistia.importar.asistencia import (
    ResultadoImportacionAsistencia,
    _dias_grid,
    _localizar_fila_encabezado_columnas,
    _procesar_persona,
    importar_asistencia,
)
from asistia.importar.calendario import (
    ResultadoImportacionCalendario,
    _escribir_dia,
    importar_calendario,
)
from asistia.importar.nexus import importar_nexus
from asistia.leyendas import (
    aplicar_leyenda_calendario,
    finalizar_asistencia,
    normalizar_leyenda,
)
from asistia.originales import conservar_original

from .gemini import SinCupo

RAIZ = Path("data/raw/data_brindada_por_ugel")
SALIDA = Path("data/cierre")
NEXUS = {
    "LUYA NEXUS AL 23-01-2025.xls": date(2025, 1, 23),
    "NEXUS LUYA AL 01-04-2026.xls": date(2026, 4, 1),
    "nexus 2026-06-01.xlsx": date(2026, 6, 1),
}
EXTENSIONES = {
    ".xlsx",
    ".xls",
    ".xlsm",
    ".pdf",
    ".docx",
    ".jpg",
    ".jpeg",
    ".png",
    ".xlsx#",
}


def normal(texto):
    return re.sub(r"[^A-Z0-9]+", " ", sin_acentos(str(texto or "")).upper()).strip()


def tipo_archivo(ruta, raiz=RAIZ):
    rel = ruta.relative_to(raiz).as_posix()
    if ruta.suffix.lower() not in EXTENSIONES or ruta.name.startswith(
        ("~$", ".~lock.")
    ):
        return None
    if rel.startswith("nexus/") and ruta.name in NEXUS:
        return "NEXUS"
    if rel.startswith("asistencia_marzo_2026_a_junio_2026/"):
        return "DRE"
    if rel.startswith(("ASISTENCIAS JULIO 2026/", "asistencia por ie/")):
        return "ASISTENCIA"
    if rel.startswith("calendarizaciones/CALENDARIZACIONES 2026 UGEL LUYA/"):
        return "CALENDARIO_2026"
    if rel.startswith("calendarizaciones/CALENDARIZACIONES 2025 UGEL LUYA"):
        return "CALENDARIO_2025"
    return None


def escribir_json(ruta, datos):
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    temporal = ruta.with_suffix(ruta.suffix + ".tmp")
    with temporal.open("w") as f:
        temporal.chmod(0o600)
        json.dump(datos, f, ensure_ascii=False, indent=2, default=str)
    temporal.replace(ruta)


def inventariar(conn, raiz=RAIZ):
    archivos = []
    for ruta in sorted(raiz.rglob("*")):
        if not ruta.is_file():
            continue
        tipo = tipo_archivo(ruta, raiz)
        sha = sha256_de(ruta)
        archivos.append(
            {
                "ruta": str(ruta.resolve()),
                "sha256": sha,
                "bytes": ruta.stat().st_size,
                "tipo": tipo or "FUERA_DE_LOTE",
                "extension": ruta.suffix.lower(),
            }
        )
        if tipo:
            conn.execute(
                """INSERT INTO cierre_documento(sha256,tipo,ruta,rutas) VALUES(%s,%s,%s,%s)
              ON CONFLICT(sha256,tipo) DO UPDATE SET rutas=CASE WHEN cierre_documento.rutas @> EXCLUDED.rutas
              THEN cierre_documento.rutas ELSE cierre_documento.rutas || EXCLUDED.rutas END""",
                (sha, tipo, str(ruta.resolve()), Jsonb([str(ruta.resolve())])),
            )
    inicial = SALIDA / "universo_inicial.json"
    if not inicial.exists():
        filas = conn.execute(
            "SELECT institucion_educativa_id,cod_mod,anexo,nombre_ie,nivel_modalidad FROM institucion_educativa ORDER BY institucion_educativa_id"
        ).fetchall()
        escribir_json(inicial, filas)
    conn.commit()
    escribir_json(
        SALIDA / "inventario.json",
        {
            "archivos": archivos,
            "por_tipo": dict(Counter(a["tipo"] for a in archivos)),
            "unicos": len({(a["sha256"], a["tipo"]) for a in archivos}),
        },
    )
    return dict(Counter(a["tipo"] for a in archivos))


def registrar_fuente(conn, item):
    if item["documento_recibido_id"]:
        return item["documento_recibido_id"]
    ruta = Path(item["ruta"])
    conservar_original(ruta, Path("data/web"))
    with conn.cursor() as cur:
        did = registrar_documento(
            cur, ruta, mimetypes.guess_type(str(ruta))[0] or "application/octet-stream"
        )
    conn.execute(
        "UPDATE cierre_documento SET documento_recibido_id=%s WHERE cierre_documento_id=%s",
        (did, item["cierre_documento_id"]),
    )
    conn.commit()
    item["documento_recibido_id"] = did
    return did


def estado(conn, item, valor, resultado=None, error=None, metodo=None):
    conn.execute(
        """UPDATE cierre_documento SET estado=%s,resultado=COALESCE(%s,resultado),error=%s,
                 metodo=COALESCE(%s,metodo),actualizado_en=now() WHERE cierre_documento_id=%s""",
        (
            valor,
            Jsonb(resultado) if resultado is not None else None,
            error,
            metodo,
            item["cierre_documento_id"],
        ),
    )
    if item["tipo"] in {"ASISTENCIA", "CALENDARIO_2026", "CALENDARIO_2025"}:
        estado_web = "PROCESADO" if valor == "PROCESADO" else "PARCIAL"
        conn.execute(
            """UPDATE web_carga c SET estado=%s,resultado=c.resultado||%s::jsonb,
          mensaje=%s,actualizado_en=now() FROM objeto_archivo o
          WHERE c.objeto_archivo_id=o.objeto_archivo_id AND o.sha256=%s AND c.tipo=%s""",
            (
                estado_web,
                Jsonb({"cierre_documento_id": str(item["cierre_documento_id"])}),
                "Consulta los bloques y pendientes en la lectura asistida. La extracción requiere cotejo."
                if valor == "PROCESADO"
                else "El original está conservado. Consulta el estado y continúa desde la lectura asistida.",
                item["sha256"],
                "asistencia" if item["tipo"] == "ASISTENCIA" else "calendario",
            ),
        )
    conn.commit()


def resolver_ie(conn, nombre, nivel, cod_mod=None, anexo="0", lugar=None):
    # Código explícito prevalece; el número de escuela nunca es cod_mod por longitud inferida.
    if cod_mod and re.fullmatch(r"\d{7}", str(cod_mod)):
        fila = conn.execute(
            "SELECT institucion_educativa_id,fn_nivel_canonico(nivel_modalidad) AS nivel FROM institucion_educativa WHERE cod_mod=%s AND anexo=%s",
            (str(cod_mod), str(anexo or "0")),
        ).fetchone()
        nivel_doc = conn.execute(
            "SELECT fn_nivel_canonico(%s) AS n", (nivel or "",)
        ).fetchone()["n"]
        if fila and (not nivel_doc or fila["nivel"] == nivel_doc):
            return fila["institucion_educativa_id"], "COD_MOD_EXPLICITO"
    with conn.cursor() as cur:
        iid = resolver_institucion(cur, str(nombre or ""), str(nivel or ""))
    if iid:
        return iid, "NOMBRE_NIVEL"
    # Todos los niveles canónicos, con nombre exacto o número institucional único + nivel.
    canon = conn.execute(
        "SELECT fn_nivel_canonico(%s) AS nivel", (nivel or nombre or "",)
    ).fetchone()["nivel"]
    candidatos = (
        conn.execute(
            "SELECT institucion_educativa_id,nombre_ie,distrito,localidad,centro_poblado FROM institucion_educativa WHERE fn_nivel_canonico(nivel_modalidad)=%s",
            (canon,),
        ).fetchall()
        if canon
        else conn.execute(
            "SELECT institucion_educativa_id,nombre_ie,distrito,localidad,centro_poblado FROM institucion_educativa"
        ).fetchall()
    )

    def nombre_servicio(valor):
        return re.sub(r"^(?:CETPRO|CEBA|CEBE|PRITE)\s+", "", normal(valor))

    objetivo = nombre_servicio(nombre)
    numero = re.search(r"\b\d{1,7}\b", objetivo)
    coincide = [
        c
        for c in candidatos
        if nombre_servicio(c["nombre_ie"]) == objetivo
        or (
            numero
            and re.search(r"\b" + re.escape(numero[0]) + r"\b", normal(c["nombre_ie"]))
        )
    ]
    if len(coincide) > 1 and lugar:
        coincide = [
            c
            for c in coincide
            if normal(lugar)
            in {
                normal(c["distrito"]),
                normal(c["localidad"]),
                normal(c["centro_poblado"]),
            }
        ]
    return (
        (
            coincide[0]["institucion_educativa_id"],
            "NOMBRE_NIVEL_UNICO" if canon else "NOMBRE_UNICO_EN_PADRON",
        )
        if len(coincide) == 1
        else (None, "PENDIENTE")
    )


def indice_personas(conn):
    indice = {}
    for f in conn.execute(
        "SELECT trabajador_id,apellido_paterno,apellido_materno,nombres FROM trabajador"
    ).fetchall():
        nombre = " ".join(
            str(f[k] or "") for k in ("apellido_paterno", "apellido_materno", "nombres")
        )
        opciones = {normal(nombre)}
        if f["apellido_paterno"] == f["nombres"]:
            opciones.add(normal(f["nombres"]))
        for n in opciones:
            indice.setdefault(n, set()).add(f["trabajador_id"])
    return indice


def preseleccionar_calendarios_2025(conn):
    """Ahorrar lecturas del año anterior cuando número y nivel identifican un servicio ya cubierto."""
    decisiones = []
    for (
        item
    ) in conn.execute("""SELECT * FROM cierre_documento WHERE tipo='CALENDARIO_2025'
      AND estado IN ('PENDIENTE','ERROR','ESPERANDO_CUPO') ORDER BY ruta""").fetchall():
        ruta = Path(item["ruta"])
        nivel = "PRIMARIA" if "PRIMARIA" in normal(str(ruta.parent)) else "INICIAL"
        nombre = ruta.parent.name if ruta.name.startswith("WhatsApp") else ruta.stem
        numero = re.match(r"^(\d{3,5})\b", nombre)
        if not numero:
            continue
        iid, regla = resolver_ie(conn, numero[1], nivel)
        if not iid:
            continue
        actual = conn.execute(
            """SELECT cv.calendarizacion_version_id FROM calendarizacion_local cl
          JOIN calendarizacion_version cv USING(calendarizacion_local_id)
          WHERE cl.institucion_educativa_id=%s AND cl.anio=2026 AND cv.estado NOT IN ('RECHAZADA','HISTORICA')
          AND COALESCE(cv.procedencia_extraccion->>'tipo','') <> 'DERIVADO_2025'
          ORDER BY cv.version DESC LIMIT 1""",
            (iid,),
        ).fetchone()
        if actual:
            decision = {
                "institucion_id": iid,
                "fuente_2026": str(actual["calendarizacion_version_id"]),
                "autor": "AGENTE_TECNICO",
                "regla": regla,
                "motivo": "Número del archivo y nivel de la carpeta coinciden con un único servicio del padrón que ya tiene calendario 2026. Se omite la lectura 2025; no se proyecta.",
            }
            estado(conn, item, "NO_NECESARIO", decision, metodo="PRESELECCION_2025")
            decisiones.append(
                {"documento_id": str(item["cierre_documento_id"]), **decision}
            )
    if decisiones:
        escribir_json(SALIDA / "preseleccion_calendarios_2025.json", decisiones)
    return decisiones


def padron_dre(conn, item):
    ruta = Path(item["ruta"])
    libro = openpyxl.load_workbook(ruta, data_only=True)
    indice = indice_personas(conn)
    total = identificados = ie_resueltas = 0
    for ws in libro:
        titulo = " ".join(
            str(c.value or "")
            for row in ws.iter_rows(max_row=min(40, ws.max_row))
            for c in row
        )
        texto = normal(titulo)
        mes = next(
            (v for k, v in MESES.items() if re.search(r"\b" + k + r"\b", texto)), None
        )
        anio = re.search(r"\b2026\b", texto)
        if not mes or not anio:
            raise ValueError("DRE sin período 2026 identificable en encabezado")
        nivel = next(
            (
                n
                for n in ("INICIAL", "PRIMARIA", "SECUNDARIA")
                if n in normal(ruta.stem + " " + titulo[:700])
            ),
            None,
        )
        encabezado = None
        columnas = {}
        for row in ws.iter_rows(max_row=min(ws.max_row, 40)):
            for c in row:
                if "APELLIDOS Y NOMBRES" in normal(c.value):
                    encabezado = c.row
                    columnas = {"nombre": c.column}
                    for h in ws[c.row]:
                        n = normal(h.value)
                        if "INSTITUCION EDUCATIVA" in n:
                            columnas["ie"] = h.column
                        elif n == "CARGO":
                            columnas["cargo"] = h.column
                        elif n == "LUGAR":
                            columnas["lugar"] = h.column
                        elif "CENTRO POBLADO" in n:
                            columnas["centro"] = h.column
                        elif n == "DNI":
                            columnas["dni"] = h.column
                    break
            if encabezado:
                break
        if not encabezado or "ie" not in columnas:
            raise ValueError("DRE sin tabla nominal reconocible")
        actual = {}
        for fila in range(encabezado + 1, ws.max_row + 1):
            for campo in ("ie", "lugar", "centro"):
                if campo in columnas:
                    valor = ws.cell(fila, columnas[campo]).value
                    if valor not in (None, ""):
                        actual[campo] = str(valor)
            nombre = ws.cell(fila, columnas["nombre"]).value
            if (
                not nombre
                or not re.search(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]", str(nombre))
                or normal(nombre) in {"APELLIDOS Y NOMBRES", "NOMBRES Y APELLIDOS"}
            ):
                continue
            cargo = (
                ws.cell(fila, columnas["cargo"]).value if "cargo" in columnas else None
            )
            # Filas de firma/pie carecen de número correlativo y cargo reconocible.
            orden = ws.cell(fila, 1).value
            if not (
                isinstance(orden, (int, float))
                or str(orden or "").strip().isdigit()
                or (
                    cargo
                    and re.search(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]", str(cargo))
                    and normal(cargo) != "CARGO"
                )
            ):
                continue
            nivel_fila = nivel or next(
                (
                    n
                    for n in ("CEBA", "CEBE", "CETPRO", "PRITE")
                    if n in normal(actual.get("ie"))
                ),
                None,
            )
            iid, metodo = resolver_ie(
                conn, actual.get("ie"), nivel_fila, lugar=actual.get("lugar")
            )
            candidatos = indice.get(normal(nombre), set())
            tid = next(iter(candidatos)) if len(candidatos) == 1 else None
            if len(candidatos) > 1 and iid:
                conocidos = {
                    r["trabajador_id"]
                    for r in conn.execute(
                        "SELECT trabajador_id FROM vinculo_trabajador_ie WHERE institucion_educativa_id=%s",
                        (iid,),
                    ).fetchall()
                }
                compatibles = candidatos & conocidos
                tid = next(iter(compatibles)) if len(compatibles) == 1 else None
            raw = {
                "nivel": nivel_fila,
                "lugar": actual.get("lugar"),
                "centro_poblado": actual.get("centro"),
                "tipo_fuente": "PADRON_DERIVADO_DRE",
                "valores_fila": json.loads(
                    json.dumps([c.value for c in ws[fila]], default=str)
                ),
                "metodo_ie": metodo,
                "limitacion": "No contiene DNI ni fechas diarias suficientes; no genera asistencia ni cese.",
            }
            conn.execute(
                """INSERT INTO padron_evidencia(cierre_documento_id,institucion_educativa_id,trabajador_id,
              periodo,hoja,fila,institucion_raw,nombres_raw,cargo_raw,datos_raw,metodo_identidad)
              VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(cierre_documento_id,hoja,fila) DO UPDATE SET
                datos_raw=padron_evidencia.datos_raw||EXCLUDED.datos_raw,
                trabajador_id=COALESCE(padron_evidencia.trabajador_id,EXCLUDED.trabajador_id),
                institucion_educativa_id=COALESCE(padron_evidencia.institucion_educativa_id,EXCLUDED.institucion_educativa_id),
                metodo_identidad=CASE WHEN padron_evidencia.trabajador_id IS NULL THEN EXCLUDED.metodo_identidad ELSE padron_evidencia.metodo_identidad END""",
                (
                    item["cierre_documento_id"],
                    iid,
                    tid,
                    date(2026, mes, 1),
                    ws.title,
                    fila,
                    actual.get("ie"),
                    str(nombre),
                    str(cargo or ""),
                    Jsonb(raw),
                    "NOMBRE_EXACTO_UNICO" if tid else "PENDIENTE",
                ),
            )
            total += 1
            identificados += bool(tid)
            ie_resueltas += bool(iid)
    libro.close()
    if not total:
        raise ValueError("No se extrajeron filas nominales del DRE")
    conn.commit()
    return {
        "filas": total,
        "personas_identificadas": identificados,
        "instituciones_identificadas_en_filas": ie_resueltas,
        "tipo_fuente": "PADRON_DERIVADO_DRE",
        "asistencia_generada": 0,
    }


def existentes(conn, item):
    tabla, col = (
        ("calendarizacion_version", "calendarizacion_version_id")
        if item["tipo"].startswith("CALENDARIO")
        else ("reporte_asistencia", "reporte_asistencia_id")
    )
    return conn.execute(
        f"""SELECT t.{col} AS id FROM {tabla} t JOIN documento_recibido d USING(documento_recibido_id)
                           JOIN objeto_archivo o USING(objeto_archivo_id) WHERE o.sha256=%s ORDER BY t.creado_en""",
        (item["sha256"],),
    ).fetchall()


def validar_bloque(b):
    if not isinstance(b, dict):
        raise TypeError("Bloque no es un objeto")
    if b.get("error_reconstruccion"):
        raise ValueError(b["error_reconstruccion"])
    normalizar_leyenda(b.get("leyenda"), b.get("tipo"))
    if b.get("tipo") not in {"calendario", "asistencia", "solo_resumen", "otro"}:
        raise ValueError("Tipo documental no reconocido")
    if b["tipo"] in {"solo_resumen", "otro"}:
        return
    if not isinstance(b.get("anio"), int) or not 2020 <= b["anio"] <= 2100:
        raise ValueError("Año inválido o ausente")
    if b["tipo"] == "calendario":
        vistos = set()
        if not b.get("meses"):
            raise ValueError("Calendario sin meses")
        for m in b["meses"]:
            if (
                not isinstance(m.get("mes"), int)
                or not 1 <= m["mes"] <= 12
                or m["mes"] in vistos
            ):
                raise ValueError("Mes inválido o repetido")
            vistos.add(m["mes"])
            if len(m.get("codigos", [])) != calendar.monthrange(b["anio"], m["mes"])[1]:
                raise ValueError("Calendario con columnas diarias incompletas")
    else:
        if not isinstance(b.get("mes"), int) or not 1 <= b["mes"] <= 12:
            raise ValueError("Mes de asistencia ausente")
        if not b.get("personas"):
            raise ValueError("Asistencia sin personas")
        dias = calendar.monthrange(b["anio"], b["mes"])[1]
        for p in b["personas"]:
            if not p.get("nombre") and not p.get("dni"):
                raise ValueError("Fila sin identidad ni nombre")
            if len(p.get("marcas", [])) != dias:
                raise ValueError("Fila de asistencia con columnas diarias incompletas")
            if any(
                not isinstance(v, (str, int, float, type(None))) for v in p["marcas"]
            ):
                raise ValueError("Marca diaria no escalar")


def aplicar_calendario(conn, item, b, meta, iid):
    meta["huella_extraccion"] = hashlib.sha256(
        json.dumps(
            {"institucion_id": iid, "bloque": b, "procedencia": meta},
            sort_keys=True,
            default=str,
        ).encode()
    ).hexdigest()
    previo = conn.execute(
        """SELECT calendarizacion_version_id FROM calendarizacion_version
      WHERE documento_recibido_id=%s AND procedencia_extraccion->>'huella_extraccion'=%s LIMIT 1""",
        (item["documento_recibido_id"], meta["huella_extraccion"]),
    ).fetchone()
    if previo:
        return str(previo["calendarizacion_version_id"])
    anio = b["anio"]
    local = conn.execute(
        """INSERT INTO calendarizacion_local(institucion_educativa_id,anio) VALUES(%s,%s)
      ON CONFLICT(institucion_educativa_id,anio) DO UPDATE SET anio=EXCLUDED.anio RETURNING calendarizacion_local_id""",
        (iid, anio),
    ).fetchone()["calendarizacion_local_id"]
    version = conn.execute(
        "SELECT COALESCE(max(version),0)+1 AS n FROM calendarizacion_version WHERE calendarizacion_local_id=%s",
        (local,),
    ).fetchone()["n"]
    cvid = conn.execute(
        """INSERT INTO calendarizacion_version(calendarizacion_local_id,version,documento_recibido_id,
      origen,motivo_version,procedencia_extraccion) VALUES(%s,%s,%s,'IMPORTACION_IE',%s,%s) RETURNING calendarizacion_version_id""",
        (
            local,
            version,
            item["documento_recibido_id"],
            "Extracción documental; revisión técnica del cierre",
            Jsonb(meta),
        ),
    ).fetchone()["calendarizacion_version_id"]
    resultado = ResultadoImportacionCalendario(calendarizacion_version_id=str(cvid))
    with conn.cursor() as cur:
        for m in b["meses"]:
            for dia, codigo in enumerate(m["codigos"], 1):
                _escribir_dia(
                    cur,
                    cvid,
                    "UGEL_LUYA_CALENDARIO_2026",
                    date(anio, m["mes"], dia),
                    codigo,
                    str(b.get("hoja") or b.get("pagina") or "1"),
                    f"mes:{m['mes']};dia:{dia}",
                    resultado,
                )
    aplicar_leyenda_calendario(
        conn, cvid, normalizar_leyenda(b.get("leyenda"), "calendario")
    )
    return str(cvid)


def aplicar_asistencia(conn, item, b, meta, iid, personas):
    meta["huella_extraccion"] = hashlib.sha256(
        json.dumps(
            {
                "institucion_id": iid,
                "bloque": b,
                "personas": personas,
                "procedencia": meta,
            },
            sort_keys=True,
            default=str,
        ).encode()
    ).hexdigest()
    previo = conn.execute(
        """SELECT reporte_asistencia_id FROM reporte_asistencia
      WHERE documento_recibido_id=%s AND procedencia_extraccion->>'huella_extraccion'=%s LIMIT 1""",
        (item["documento_recibido_id"], meta["huella_extraccion"]),
    ).fetchone()
    if previo:
        return str(previo["reporte_asistencia_id"])
    periodo = date(b["anio"], b["mes"], 1)
    ie = conn.execute(
        "SELECT * FROM institucion_educativa WHERE institucion_educativa_id=%s", (iid,)
    ).fetchone()
    turno = str(b.get("turno") or "").strip().upper()
    if not turno:
        turnos = conn.execute(
            "SELECT DISTINCT turno FROM reporte_asistencia_serie WHERE institucion_educativa_id=%s AND periodo=%s",
            (iid, periodo),
        ).fetchall()
        turno = turnos[0]["turno"] if len(turnos) == 1 else "TODOS"
        meta["supuestos"] = [
            *meta.get("supuestos", []),
            "Turno no declarado; se usa el único turno conocido del mes o TODOS si no hay uno único.",
        ]
    serie = conn.execute(
        """INSERT INTO reporte_asistencia_serie(institucion_educativa_id,periodo,nivel_modalidad,turno)
      VALUES(%s,%s,%s,%s) ON CONFLICT(institucion_educativa_id,periodo,nivel_modalidad,turno)
      DO UPDATE SET turno=EXCLUDED.turno RETURNING reporte_asistencia_serie_id""",
        (iid, periodo, ie["nivel_modalidad"], turno),
    ).fetchone()["reporte_asistencia_serie_id"]
    version = conn.execute(
        """SELECT COALESCE(max(version),0)+1 AS n FROM reporte_asistencia
        WHERE reporte_asistencia_serie_id=%s OR (documento_recibido_id=%s
        AND hoja_pagina_origen=%s AND indice_bloque=%s)""",
        (
            serie,
            item["documento_recibido_id"],
            str(b.get("hoja") or b.get("pagina") or "1")[:60],
            meta["bloque"],
        ),
    ).fetchone()["n"]
    sufijo = Path(item["ruta"]).suffix.lower()
    tipo = (
        "PDF_ESCANEADO"
        if sufijo == ".pdf"
        else (
            "FOTO"
            if sufijo in {".jpg", ".jpeg", ".png"}
            else ("DOCX_NATIVO" if sufijo == ".docx" else "EXCEL_NATIVO")
        )
    )
    rid = conn.execute(
        """INSERT INTO reporte_asistencia(reporte_asistencia_serie_id,institucion_educativa_id,
      documento_recibido_id,periodo,version,tipo_fuente,institucion_reportada_raw,nivel_modalidad_reportada_raw,
      hoja_pagina_origen,indice_bloque,estado_match_ie,procedencia_extraccion)
      VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'RESUELTO',%s) RETURNING reporte_asistencia_id""",
        (
            serie,
            iid,
            item["documento_recibido_id"],
            periodo,
            version,
            tipo,
            b.get("institucion") or "",
            b.get("nivel"),
            str(b.get("hoja") or b.get("pagina") or "1")[:60],
            meta["bloque"],
            Jsonb(meta),
        ),
    ).fetchone()["reporte_asistencia_id"]
    ws = openpyxl.Workbook().active
    columnas = {"dni": 1, "nombres": 2, "cargo": 3, "condicion": 35}
    dias = calendar.monthrange(periodo.year, periodo.month)[1]
    resultado = ResultadoImportacionAsistencia()
    with conn.cursor() as cur:
        for orden, p in enumerate(personas, 1):
            fila = (
                orden  # Identidad interna; página y fila de origen se conservan abajo.
            )
            ws.cell(fila, 1, p.get("dni"))
            ws.cell(fila, 2, p.get("nombre"))
            ws.cell(fila, 3, p.get("cargo"))
            ws.cell(fila, 35, p.get("condicion"))
            for d, codigo in enumerate(p["marcas"], 1):
                ws.cell(fila, d + 3, codigo)
            _procesar_persona(
                cur,
                ws,
                fila,
                columnas,
                [(d, d + 3) for d in range(1, dias + 1)],
                periodo,
                rid,
                {},
                resultado,
                diferir_marcas=True,
            )
            localizador = f"{p.get('__hoja') or ('pagina:' + str(p.get('__pagina') or 1))};fila:{p.get('fila') or orden}"
            conn.execute(
                """UPDATE asistencia_dia a SET celda_origen=%s||';dia:'||EXTRACT(day FROM a.fecha)::int
              FROM trabajador_en_reporte t WHERE a.trabajador_en_reporte_id=t.trabajador_en_reporte_id
              AND t.reporte_asistencia_id=%s AND t.fila_detalle_origen=%s""",
                (localizador, rid, fila),
            )
            for d, loc in p.get("localizadores_marcas", {}).items():
                conn.execute(
                    """UPDATE asistencia_dia a SET celda_origen=%s,evidencia_interpretacion=%s
                    FROM trabajador_en_reporte t WHERE a.trabajador_en_reporte_id=t.trabajador_en_reporte_id
                    AND t.reporte_asistencia_id=%s AND t.fila_detalle_origen=%s AND extract(day FROM a.fecha)=%s""",
                    (
                        f"{loc.get('hoja') or ('pagina:' + str(loc.get('pagina')))};fila:{loc.get('fila') or loc['ancla']};dia:{d}",
                        Jsonb(
                            {
                                "reconstruccion": p["fuentes_fragmentos"],
                                "celda_nativa": loc,
                                **(
                                    {
                                        "valor_extraido_ocr": p["marcas_ocr"][
                                            int(d) - 1
                                        ]
                                        if int(d) <= len(p["marcas_ocr"])
                                        else None
                                    }
                                    if "marcas_ocr" in p
                                    else {}
                                ),
                            }
                        ),
                        rid,
                        fila,
                        int(d),
                    ),
                )
    categorias = normalizar_leyenda(b.get("leyenda"), "asistencia")
    if sufijo in {".xlsx", ".xlsm"}:
        from asistia.calidad_leyendas import (
            HojaLeyendaNoDeterminada,
            cotejar_excel,
            referencias,
        )

        try:
            categorias = cotejar_excel(
                item["ruta"], b.get("hoja"), categorias, referencias(conn)
            )
        except HojaLeyendaNoDeterminada as exc:
            conn.execute(
                "UPDATE reporte_asistencia SET procedencia_extraccion=procedencia_extraccion||%s WHERE reporte_asistencia_id=%s",
                (
                    Jsonb(
                        {
                            "control_leyenda": {
                                "estado": "NO_COMPROBABLE",
                                "motivo": str(exc),
                            }
                        }
                    ),
                    rid,
                ),
            )
    finalizar_asistencia(conn, rid, categorias)
    return str(rid)


def aplicar_extraccion(conn, item, extraccion):
    bloques = bloques_normalizados(item, extraccion)
    for b in bloques:
        validar_bloque(b)
    ids, pendientes, complementos = (
        [],
        list(extraccion.get("pendientes_extraccion", [])),
        [],
    )
    grupos = {}
    esperado = "calendario" if item["tipo"].startswith("CALENDARIO") else "asistencia"
    for idx, b in enumerate(bloques, 1):
        idx = b.get("indice_bloque_original", idx)
        if b["tipo"] in {"otro", "solo_resumen"}:
            complementos.append({"bloque": idx, "tipo": b["tipo"]})
            continue
        if b["tipo"] != esperado:
            pendientes.append({"bloque": idx, "motivo": "tipo_difiere_del_lote"})
            continue
        personas = b["personas"] if esperado == "asistencia" else [None]
        for p in personas:
            nivel = (p.get("nivel") if p else None) or b.get("nivel")
            iid, regla = resolver_ie(
                conn, b.get("institucion"), nivel, b.get("cod_mod"), b.get("anexo")
            )
            if (
                not iid
                and p
                and p.get("nivel")
                and b.get("nivel")
                and p["nivel"] != b["nivel"]
            ):
                cabecera, _ = resolver_ie(
                    conn,
                    b.get("institucion"),
                    b["nivel"],
                    b.get("cod_mod"),
                    b.get("anexo"),
                )
                if cabecera:
                    iid, regla = cabecera, "NIVEL_CABECERA_UNICO"
            contexto = (
                item.get("contexto_tecnico", {})
                .get("instituciones", {})
                .get(f"{idx}:{nivel or ''}")
            )
            if contexto:
                asignada = conn.execute(
                    "SELECT institucion_educativa_id FROM institucion_educativa WHERE institucion_educativa_id=%s",
                    (contexto["institucion_id"],),
                ).fetchone()
                if not asignada:
                    raise ValueError(
                        "La institución asignada al documento ya no existe"
                    )
                iid, regla = (
                    asignada["institucion_educativa_id"],
                    "CONTEXTO_DOCUMENTAL_DECLARADO",
                )
            if not iid:
                pendiente = {
                    "bloque": idx,
                    "motivo": "institucion_ambigua",
                    "institucion": b.get("institucion"),
                    "nivel": nivel,
                }
                if pendiente not in pendientes:
                    pendientes.append(pendiente)
                continue
            if (
                item["tipo"] == "CALENDARIO_2025"
                and conn.execute(
                    """SELECT 1 FROM calendarizacion_local cl JOIN calendarizacion_version cv USING(calendarizacion_local_id)
                   WHERE cl.institucion_educativa_id=%s AND cl.anio=2026
                     AND cv.estado NOT IN ('RECHAZADA','HISTORICA')
                     AND COALESCE(cv.procedencia_extraccion->>'tipo','') <> 'DERIVADO_2025' LIMIT 1""",
                    (iid,),
                ).fetchone()
            ):
                complementos.append(
                    {"bloque": idx, "tipo": "2026_disponible", "institucion_id": iid}
                )
                continue
            clave = (iid, b["anio"], b.get("mes"), b.get("turno"))
            if clave not in grupos:
                grupos[clave] = {
                    "bloque": {**b, "meses": []},
                    "personas": [],
                    "meta": {
                        "metodo": "GEMINI",
                        "modelo": extraccion["modelo"],
                        "version": extraccion["version"],
                        "sha256": extraccion["sha256"],
                        "bloque": len(grupos) + 1,
                        "pagina": b.get("pagina"),
                        "hoja": b.get("hoja"),
                        "supuestos": list(b.get("supuestos", [])),
                        "regla_identidad_ie": regla,
                        "revision": "TECNICA_PENDIENTE",
                        "cierre_documento_id": str(item["cierre_documento_id"]),
                        "fuentes_bloques": [],
                        "fragmentos": extraccion.get("fragmentos", []),
                        "reconstruccion": b.get("reconstruccion"),
                        "periodo_extraido": b.get("periodo_extraido"),
                        "ajuste_periodo": b.get("ajuste_periodo"),
                    },
                }
                if extraccion.get("conversion"):
                    grupos[clave]["meta"]["conversion"] = extraccion["conversion"]
                if contexto:
                    grupos[clave]["meta"]["supuestos"].append(contexto["motivo"])
                    grupos[clave]["meta"]["resolucion_institucion"] = contexto
                if regla == "NOMBRE_UNICO_EN_PADRON":
                    grupos[clave]["meta"]["supuestos"].append(
                        "Nivel no determinado en el documento; se usa el del único servicio educativo con ese nombre en el padrón."
                    )
            grupo = grupos[clave]
            for entrada in b.get("leyenda_sin_codigo", []):
                sin_codigo = grupo["meta"].setdefault("leyenda_sin_codigo", [])
                if entrada not in sin_codigo:
                    sin_codigo.append(entrada)
            for supuesto in b.get("supuestos", []):
                if supuesto not in grupo["meta"]["supuestos"]:
                    grupo["meta"]["supuestos"].append(supuesto)
            leyenda = grupo["bloque"].setdefault("leyenda", [])
            for entrada in b.get("leyenda", []):
                if entrada not in leyenda:
                    leyenda.append(entrada)
            normalizar_leyenda(leyenda, esperado)
            if regla == "NIVEL_CABECERA_UNICO":
                nota = "El nivel atribuido a una persona no identifica un servicio con este nombre. Se usa el nivel explícito de la cabecera, coincidente con el padrón; una especialidad profesional no cambia la institución."
                if nota not in grupo["meta"]["supuestos"]:
                    grupo["meta"]["supuestos"].append(nota)
            fuente = {"bloque": idx, "pagina": b.get("pagina"), "hoja": b.get("hoja")}
            if fuente not in grupo["meta"]["fuentes_bloques"]:
                grupo["meta"]["fuentes_bloques"].append(fuente)
            if p:
                duplicada = next(
                    (
                        prev
                        for prev in grupo["personas"]
                        if {k: v for k, v in prev.items() if not k.startswith("__")}
                        == p
                    ),
                    None,
                )
                if duplicada is not None:
                    grupo["meta"].setdefault("filas_repetidas", []).append(
                        {
                            "fila": p.get("fila"),
                            "pagina": b.get("pagina"),
                            "hoja": b.get("hoja"),
                            "pagina_conservada": duplicada.get("__pagina"),
                            "motivo": "Fila idéntica repetida en el mismo documento; se conserva una sola observación.",
                        }
                    )
                    continue
                grupo["personas"].append(
                    {**p, "__pagina": b.get("pagina"), "__hoja": b.get("hoja")}
                )
            else:
                for mes in b["meses"]:
                    previo = next(
                        (m for m in grupo["bloque"]["meses"] if m["mes"] == mes["mes"]),
                        None,
                    )
                    if previo and previo != mes:
                        raise ValueError(
                            "Dos tablas del documento contradicen el mismo mes de calendario"
                        )
                    if not previo:
                        grupo["bloque"]["meses"].append(mes)
    for (iid, *_), grupo in grupos.items():
        if esperado == "calendario":
            ids.append(
                aplicar_calendario(conn, item, grupo["bloque"], grupo["meta"], iid)
            )
        else:
            ids.append(
                aplicar_asistencia(
                    conn, item, grupo["bloque"], grupo["meta"], iid, grupo["personas"]
                )
            )
    return {
        "ids": ids,
        "pendientes": pendientes,
        "complementos": complementos,
        "bloques": len(bloques),
        "advertencias": extraccion["datos"].get("advertencias", []),
        "datos_no_aplicados": extraccion.get("datos_no_aplicados", []),
    }


def bloques_normalizados(item, extraccion):
    """Normalizar representación y declarar contexto; nunca rellenar una marca omitida."""
    bloques = json.loads(json.dumps(extraccion["datos"]["bloques"]))
    for idx, b in enumerate(bloques, 1):
        idx = b.get("indice_bloque_original", idx)
        # Una descripción sin símbolo textual (por ejemplo, solo color) no
        # autoriza asignar una letra ni dar significado a celdas vacías.
        if isinstance(b.get("leyenda"), list):
            sin_codigo = [
                e
                for e in b["leyenda"]
                if isinstance(e, dict)
                and e.get("codigo") in (None, "")
                and isinstance(e.get("descripcion"), str)
                and e["descripcion"].strip()
            ]
            if sin_codigo:
                b.setdefault("leyenda_sin_codigo", []).extend(sin_codigo)
                b["leyenda"] = [e for e in b["leyenda"] if e not in sin_codigo]
                b.setdefault("supuestos", []).append(
                    "La leyenda incluye descripciones sin código textual: se conservan sin inventar letras ni asignarlas a días vacíos."
                )
        if b.get("tipo") not in {"calendario", "asistencia"}:
            continue
        b.setdefault("supuestos", [])
        ajuste = item.get("contexto_tecnico", {}).get("periodos", {}).get(str(idx))
        if ajuste:
            b.setdefault(
                "periodo_extraido", {"anio": b.get("anio"), "mes": b.get("mes")}
            )
            b["anio"], b["mes"] = ajuste["anio"], ajuste["mes"]
            b["ajuste_periodo"] = ajuste
            b["supuestos"].append(
                f"Período extraído {b['periodo_extraido']} interpretado como {ajuste['anio']}-{ajuste['mes']:02}: {ajuste['motivo']}"
            )
        for campo in ("anio", "mes"):
            valor = b.get(campo)
            if isinstance(valor, str) and valor.strip().isdigit():
                b[campo] = int(valor.strip())
        if b.get("anio") is None and item["tipo"].startswith("CALENDARIO_"):
            b["anio"] = int(item["tipo"][-4:])
            b["supuestos"].append(
                f"Año no visible: se usa {b['anio']} por la carpeta indicada por el usuario."
            )
        if b.get("tipo") == "asistencia" and "ASISTENCIAS JULIO 2026" in item["ruta"]:
            for campo, valor in [("anio", 2026), ("mes", 7)]:
                if b.get(campo) is None:
                    b[campo] = valor
                    b["supuestos"].append(
                        f"{campo} no visible: se usa {valor} por el lote de julio 2026 indicado por el usuario."
                    )
        if isinstance(b.get("anio"), int):
            for m in b.get("meses", []):
                if isinstance(m.get("mes"), str) and m["mes"].isdigit():
                    m["mes"] = int(m["mes"])
                if not isinstance(m.get("mes"), int) or not 1 <= m["mes"] <= 12:
                    continue
                dias = calendar.monthrange(b["anio"], m["mes"])[1]
                codigos = m.get("codigos", [])
                if dias < len(codigos) <= 31 and all(v is None for v in codigos[dias:]):
                    m["codigos"] = codigos[:dias]
                    b["supuestos"].append(
                        f"Mes {m['mes']}: columnas vacías posteriores al último día no representan fechas y se omiten."
                    )
    from .tablas import reconstruir_columnas

    reconstruidos = []
    for b in bloques:
        try:
            reconstruidos.append(reconstruir_columnas(b))
        except (ValueError, TypeError, KeyError) as exc:
            reconstruidos.append({**b, "error_reconstruccion": str(exc)})
    return reconstruidos


def procesar_documento(conn, item, gemini=None, solo_nativo=False):
    ruta = Path(item["ruta"])
    if sha256_de(ruta) != item["sha256"]:
        raise ValueError("El original cambió desde el inventario")
    registrar_fuente(conn, item)
    conn.execute(
        "UPDATE cierre_documento SET estado='PROCESANDO',intentos=intentos+1 WHERE cierre_documento_id=%s",
        (item["cierre_documento_id"],),
    )
    conn.commit()
    if item["tipo"] == "NEXUS":
        corte = NEXUS[ruta.name]
        carga = conn.execute(
            """SELECT n.* FROM nexus_carga n JOIN documento_recibido d USING(documento_recibido_id)
          JOIN objeto_archivo o USING(objeto_archivo_id) WHERE o.sha256=%s AND n.fecha_corte=%s
          AND (SELECT count(*) FROM nexus_registro r WHERE r.nexus_carga_id=n.nexus_carga_id)=n.total_filas
          ORDER BY n.cargado_en DESC LIMIT 1""",
            (item["sha256"], corte),
        ).fetchone()
        if carga:
            datos = {
                "nexus_carga_id": str(carga["nexus_carga_id"]),
                "total_filas": carga["total_filas"],
                "reutilizado": True,
            }
        else:
            datos = json.loads(
                json.dumps(asdict(importar_nexus(conn, ruta, corte)), default=str)
            )
        estado(
            conn,
            item,
            "PARCIAL" if datos.get("errores") else "PROCESADO",
            datos,
            metodo="NEXUS_NATIVO",
        )
        return
    if item["tipo"] == "DRE":
        datos = padron_dre(conn, item)
        estado(conn, item, "PROCESADO", datos, metodo="PADRON_DERIVADO_DRE")
        return
    previo = existentes(conn, item)
    # Fuentes ya cargadas se conservan. Sus pendientes de contenido van a revisión institucional.
    if previo and item["intentos"] == 0:
        estado(
            conn,
            item,
            "PROCESADO",
            {"ids": [str(r["id"]) for r in previo], "reutilizado": True},
            metodo="NATIVO_PREEXISTENTE",
        )
        return
    if item.get("metodo") not in {
        "NATIVO_FALLIDO",
        "GEMINI",
    } and ruta.suffix.lower() in {".xlsx", ".pdf"}:
        try:
            resultado = None
            if item["tipo"].startswith("CALENDARIO"):
                resultado = importar_calendario(conn, ruta)
            elif ruta.suffix.lower() == ".xlsx" and len(hojas_diarias(ruta)) <= 1:
                resultado = importar_asistencia(conn, ruta)
            if resultado:
                datos = json.loads(json.dumps(asdict(resultado), default=str))
                correcto = not (
                    datos.get("error")
                    or datos.get("errores_fila")
                    or datos.get("meses_sin_alinear")
                )
                if item["tipo"].startswith("CALENDARIO"):
                    correcto = correcto and datos.get("dias_registrados", 0) >= 200
                if correcto:
                    estado(conn, item, "PROCESADO", datos, metodo="NATIVO")
                    return
                conn.rollback()
        except Exception:  # noqa: BLE001 -- el OCR recupera fallos del lector nativo por documento
            conn.rollback()
    if solo_nativo:
        estado(
            conn,
            item,
            "PENDIENTE",
            error="Requiere extracción documental adicional",
            metodo="NATIVO_FALLIDO",
        )
        return
    if gemini is None:
        raise SinCupo("Extracción Gemini pendiente de configurar")
    contexto = f"Tipo esperado: {item['tipo']}. Nombre/ruta proporcionados por usuario: {ruta.relative_to(RAIZ.resolve()) if ruta.is_relative_to(RAIZ.resolve()) else ruta.name}."
    from .extraccion import extraer_documento

    extraccion = extraer_documento(gemini, item, contexto)
    datos = aplicar_extraccion(conn, item, extraccion)
    nuevo = estado_extraccion(datos)
    estado(conn, item, nuevo, datos, metodo="GEMINI")


def estado_extraccion(datos):
    if datos["ids"]:
        return "PARCIAL" if datos["pendientes"] else "PROCESADO"
    if not datos["pendientes"] and any(
        c.get("tipo") == "2026_disponible" for c in datos.get("complementos", [])
    ):
        return "NO_NECESARIO"
    return "SIN_DATOS"


def ejecutar(conn, tipos=None, solo_nativo=False, gemini=None, limite=None):
    if not conn.execute("SELECT pg_try_advisory_lock(26091206) AS obtenido").fetchone()[
        "obtenido"
    ]:
        raise RuntimeError(
            "Hay otra carga de cierre activa; revisar su progreso antes de reintentar"
        )
    try:
        items = conn.execute(
            """SELECT * FROM cierre_documento WHERE tipo=ANY(%s)
          AND estado IN ('PENDIENTE','ERROR','ESPERANDO_CUPO','PROCESANDO')
          ORDER BY CASE tipo WHEN 'NEXUS' THEN 0 WHEN 'DRE' THEN 1 WHEN 'CALENDARIO_2026' THEN 2
            WHEN 'ASISTENCIA' THEN 3 ELSE 4 END,ruta""",
            (tipos or ["NEXUS", "DRE", "CALENDARIO_2026", "ASISTENCIA"],),
        ).fetchall()
        items.sort(
            key=lambda i: (
                (0, str(NEXUS[Path(i["ruta"]).name]))
                if i["tipo"] == "NEXUS"
                else (
                    {
                        "DRE": 1,
                        "CALENDARIO_2026": 2,
                        "ASISTENCIA": 3,
                        "CALENDARIO_2025": 4,
                    }[i["tipo"]],
                    i["ruta"],
                )
            )
        )
        resumen = Counter()
        for i, item in enumerate(items[:limite] if limite else items, 1):
            try:
                procesar_documento(conn, item, gemini, solo_nativo)
            except SinCupo as exc:
                conn.rollback()
                estado(conn, item, "ESPERANDO_CUPO", error=str(exc), metodo="GEMINI")
            except Exception as exc:  # noqa: BLE001 -- aislar cada documento y conservar el error local
                conn.rollback()
                # Mensaje local validado; SQL/tracebacks no llegan a interfaz ni logs públicos.
                mensaje = (
                    str(exc)[:500]
                    if isinstance(exc, ValueError)
                    else type(exc).__name__
                )
                estado(conn, item, "ERROR", error=mensaje)
            actual = conn.execute(
                "SELECT estado FROM cierre_documento WHERE cierre_documento_id=%s",
                (item["cierre_documento_id"],),
            ).fetchone()["estado"]
            resumen[actual] += 1
            print(
                f"{i}/{len(items)} {item['tipo']} {actual} {str(item['cierre_documento_id'])[:8]}",
                flush=True,
            )
        return dict(resumen)
    finally:
        conn.rollback()
        conn.execute("SELECT pg_advisory_unlock(26091206)")
        conn.commit()


def hojas_diarias(ruta):
    """Detectar libros que el importador de una sola tabla no cubre completos."""
    with Path(ruta).open("rb") as archivo:
        libro = openpyxl.load_workbook(archivo, data_only=True)
        hojas = []
        for ws in libro:
            fila = _localizar_fila_encabezado_columnas(ws)
            if fila is not None and _dias_grid(ws, fila + 1):
                hojas.append(ws.title)
        libro.close()
    return hojas


def asignar_contexto(conn, did, bloque, iid, nivel, motivo, autor):
    """Una asignación solo resuelve el bloque/nivel señalado; el historial conserva el supuesto."""
    item = conn.execute(
        "SELECT * FROM cierre_documento WHERE cierre_documento_id=%s FOR UPDATE", (did,)
    ).fetchone()
    if not item or not any(
        p.get("bloque") == bloque
        and (p.get("nivel") or "") == (nivel or "")
        and p.get("motivo") == "institucion_ambigua"
        for p in item["resultado"].get("pendientes", [])
    ):
        raise ValueError(
            "Ese bloque ya no tiene una institución pendiente. Actualiza el documento."
        )
    if not motivo or len(motivo) > 4000:
        raise ValueError("Escribe el motivo y la evidencia de la asignación.")
    ie = conn.execute(
        "SELECT fn_nivel_canonico(nivel_modalidad) AS nivel FROM institucion_educativa WHERE institucion_educativa_id=%s",
        (iid,),
    ).fetchone()
    esperado = conn.execute(
        "SELECT fn_nivel_canonico(%s) AS nivel", (nivel or "",)
    ).fetchone()["nivel"]
    if not ie or (esperado and ie["nivel"] != esperado):
        raise ValueError(
            "La institución seleccionada no corresponde al nivel del bloque."
        )
    contexto = item["contexto_tecnico"]
    evento = {
        "institucion_id": iid,
        "bloque": bloque,
        "nivel": nivel,
        "motivo": motivo,
        "autor": autor,
        "fecha": datetime.now(UTC).isoformat(),
    }
    contexto.setdefault("instituciones", {})[f"{bloque}:{nivel or ''}"] = evento
    contexto.setdefault("historial", []).append(evento)
    conn.execute(
        "UPDATE cierre_documento SET contexto_tecnico=%s,estado='PENDIENTE',actualizado_en=now() WHERE cierre_documento_id=%s",
        (Jsonb(contexto), did),
    )
    conn.commit()

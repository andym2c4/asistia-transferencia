"""Endpoints HTML del recorrido de RRHH; todas las escrituras son POST + CSRF."""

from __future__ import annotations

import csv
import io
import json
import uuid
from datetime import date, datetime
from zoneinfo import ZoneInfo

import openpyxl
import xlrd
from flask import (
    Blueprint,
    Response,
    current_app,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)

from asistia.consolidado.estimacion import DECLARACION_VIGENTE, estimar
from asistia.consolidado.resolucion import (
    ESTADOS_RESOLUCION,
    estado_resolucion,
    instituciones_sin_nivel_canonico,
    resumen_ugel,
)
from asistia.importar.calendario import _FAMILIA_FORMATO_DEFECTO
from asistia.niveles import NIVELES

# Texto que ve RRHH para cada estado del desglose. El codigo interno no se muestra crudo.
ETIQUETAS_ESTADO = {
    "SIN_RECEPCION": "Sin recepción",
    "RECIBIDO_SIN_EXTRAER": "Recibido, sin extraer",
    "EXTRAIDO_PENDIENTE": "Extraído, pendiente de revisión",
    "REVISADO": "Revisado",
}

from asistia.consolidado.remuneracion import cruce_persona
from asistia.leyendas import categoria_dia

from . import ErrorDeTrabajo, instituciones_edicion, lecturas, leyendas
from .archivos import ruta_verificada
from .calendario_mensual import MESES, guardar_dias, meses_calendario, opciones_dias
from .calendarios import ESTADOS_APROBABLES, dias_bloqueantes, marcar_vigente
from .cargas import procesar, recibir
from .codigos_dia import (
    clasificaciones_vigentes,
    clasificar_codigo,
    codigos_desconocidos,
    tipos_dia,
)
from .db import base, escritura
from .instituciones import (
    calendario_institucion,
    directorio,
    reportes_institucion,
)
from .instituciones import (
    filtros as filtros_directorio,
)
from .revision import (
    confirmar_reporte,
    corregir_marca,
    decidir_alerta,
    resolver_identidad,
)
from .revision_vacios import clasificar_vacios
from .salidas import preparar

pages = Blueprint("pages", __name__)


def operacion():
    try:
        return uuid.UUID(request.form["operacion"])
    except (KeyError, ValueError):
        raise ErrorDeTrabajo(
            "La solicitud no tiene una referencia válida. Actualiza el formulario."
        ) from None


def motivo(predeterminado=None):
    texto = request.form.get("motivo", "").strip()
    if (not texto and not predeterminado) or len(texto) > 4000:
        raise ErrorDeTrabajo(
            "Escribe el motivo y sustento de la decisión (hasta 4000 caracteres)."
        )
    return texto or predeterminado


def ambito():
    mes = request.values.get(
        "periodo",
        session.get(
            "periodo", datetime.now(ZoneInfo("America/Lima")).strftime("%Y-%m")
        ),
    )
    nivel = request.values.get("nivel", session.get("nivel", "PRIMARIA"))
    try:
        periodo = date.fromisoformat(mes + "-01")
    except ValueError:
        raise ErrorDeTrabajo("Selecciona un mes y año válidos.") from None
    if nivel not in NIVELES:
        raise ErrorDeTrabajo("Selecciona uno de los niveles educativos disponibles.")
    session.update(periodo=mes, nivel=nivel)
    return periodo, nivel


def vista(nombre, **kwargs):
    periodo, nivel = ambito()
    return render_template(nombre, periodo=periodo, nivel=nivel, **kwargs)


@pages.get("/salud")
def salud():
    return {"estado": "disponible"}


@pages.get("/sesion")
def sesion():
    return {"csrf": session["csrf"]}


@pages.get("/")
def mes():
    from .dashboard import preparar as preparar_dashboard

    periodo, nivel = ambito()
    conn = base()
    # La autenticación previa solo leyó. Todos los indicadores comparten ahora
    # una instantánea; consultar el Dashboard no registra cortes ni decisiones.
    conn.rollback()
    conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
    datos = preparar_dashboard(conn, periodo, nivel, request.args.get("corte"))
    estado = request.args.get("estado", "todos")
    datos["reportes_visibles"] = [
        r
        for r in datos["reportes"]
        if estado == "todos"
        or (estado == "pendientes" and r["estado"] != "VALIDADO")
        or (estado == "revisados" and r["estado"] == "VALIDADO")
    ]
    return vista("mes.html", filtro=estado, **datos)


@pages.post("/dashboard/analisis")
def actualizar_analisis_dashboard():
    from asistia.monitoreo.cortes import crear

    periodo, nivel = ambito()
    try:
        with escritura() as conn:
            cid, nuevo = crear(
                conn,
                periodo.year,
                autor=str(g.usuario["usuario_id"]),
                motivo=f"Actualización solicitada desde Dashboard para {periodo:%Y-%m}, {nivel}. Análisis de fuentes del año.",
            )
    except ValueError as exc:
        raise ErrorDeTrabajo(str(exc)) from None
    flash(
        "Análisis actualizado. Consulta las sugerencias y sus documentos."
        if nuevo
        else "El análisis ya corresponde a las fuentes actuales."
    )
    return redirect(
        url_for(
            "pages.mes",
            periodo=periodo.strftime("%Y-%m"),
            nivel=nivel,
            corte=cid,
            _anchor="analisis",
        ),
        code=303,
    )


@pages.get("/desarrollo")
def desarrollo():
    from .desarrollo import HERRAMIENTAS

    return vista("desarrollo.html", herramientas=HERRAMIENTAS, mostrar_ambito=False)


@pages.get("/revision/subir", endpoint="subir_reporte")
@pages.get("/documentos")
def documentos():
    from .revision_bandeja import contexto_carga

    destino = contexto_carga(base(), request.args.get("institucion_id"))
    if request.endpoint == "pages.subir_reporte":
        if destino is None:
            raise ErrorDeTrabajo("Selecciona una institución desde Revisión.")
        return vista(
            "documentos.html",
            institucion_destino=destino,
            tiene_nexus=True,
            destino_ambito="pages.revision",
        )
    try:
        pagina = max(1, int(request.args.get("pagina", 1)))
    except ValueError:
        raise ErrorDeTrabajo("La página solicitada no es válida.") from None
    cargas = (
        base()
        .execute(
            """SELECT c.*,o.sha256,r.nombre_original,r.recibido_real_en,r.cargado_en,
                  (SELECT count(*) FROM web_recepcion x WHERE x.web_carga_id=c.web_carga_id) AS recepciones
           FROM web_carga c JOIN objeto_archivo o USING(objeto_archivo_id)
           JOIN LATERAL (SELECT * FROM web_recepcion r WHERE r.web_carga_id=c.web_carga_id ORDER BY cargado_en DESC LIMIT 1) r ON true
           ORDER BY r.cargado_en DESC LIMIT 25 OFFSET %s""",
            ((pagina - 1) * 25,),
        )
        .fetchall()
    )
    total = base().execute("SELECT count(*) AS n FROM web_carga").fetchone()["n"]
    tiene_nexus = bool(
        base().execute("SELECT 1 FROM institucion_educativa LIMIT 1").fetchone()
    )
    return vista(
        "documentos.html",
        cargas=cargas,
        pagina=pagina,
        total=total,
        tiene_nexus=tiene_nexus,
        institucion_destino=destino,
    )


@pages.get("/instituciones/<int:iid>/calendario/subir")
def subir_calendario(iid):
    from .revision_bandeja import contexto_carga

    return vista(
        "calendario_carga.html",
        institucion_destino=contexto_carga(base(), iid),
        mostrar_ambito=False,
    )


@pages.post("/instituciones/<int:iid>/calendario/subir", endpoint="cargar_calendario")
@pages.post("/documentos")
@pages.post("/revision/cargar", endpoint="cargar_revision")
def cargar(iid=None):
    from .revision_bandeja import contexto_carga
    from .revision_cargas import resultado_carga

    desde_revision = request.endpoint == "pages.cargar_revision"
    desde_calendario = request.endpoint == "pages.cargar_calendario"
    cola_json = desde_revision and request.accept_mimetypes.best == "application/json"
    periodo, _nivel = ambito()
    destino = contexto_carga(
        base(), iid if desde_calendario else request.form.get("institucion_id")
    )
    tipo = request.form.get("tipo")
    if tipo not in {"nexus", "calendario", "asistencia"}:
        raise ErrorDeTrabajo("Selecciona qué tipo de documento vas a cargar.")
    if desde_calendario and tipo != "calendario":
        raise ErrorDeTrabajo("Selecciona un archivo de calendarización.")
    if not desde_calendario and (destino or desde_revision) and tipo != "asistencia":
        raise ErrorDeTrabajo(
            "Desde la bandeja de revisión, carga un reporte de asistencia."
        )
    corte = None
    recibido = None
    try:
        if tipo == "nexus":
            corte = date.fromisoformat(request.form.get("fecha_corte", ""))
        if request.form.get("recibido_real_en"):
            recibido = datetime.fromisoformat(request.form["recibido_real_en"]).replace(
                tzinfo=ZoneInfo("America/Lima")
            )
            if recibido > datetime.now(ZoneInfo("America/Lima")):
                raise ValueError("recepción futura")
    except ValueError:
        raise ErrorDeTrabajo(
            "Revisa la fecha de corte y la fecha real de recepción. La recepción no puede ser futura; déjala vacía si no se conoce."
        ) from None
    archivos = [f for f in request.files.getlist("archivos") if f.filename]
    if (destino or cola_json) and len(archivos) != 1:
        raise ErrorDeTrabajo(
            "Selecciona un único archivo de calendario."
            if desde_calendario
            else "Selecciona un único archivo de reporte mensual."
        )
    if not 1 <= len(archivos) <= 10:
        raise ErrorDeTrabajo("Selecciona entre uno y diez archivos.")
    op = operacion()
    recibidas = []
    resultados = []
    with escritura() as conn:
        for index, archivo in enumerate(archivos):
            receipt = uuid.uuid5(op, str(index))
            cid = None
            duplicado = False
            try:
                cid, duplicado = recibir(
                    conn,
                    archivo,
                    tipo,
                    corte,
                    recibido,
                    periodo,
                    g.usuario["usuario_id"],
                    receipt,
                )
                recibidas.append(cid)
                if not duplicado:
                    procesar(
                        conn,
                        cid,
                        g.usuario["usuario_id"],
                        uuid.uuid5(receipt, "procesar"),
                    )
                if cola_json:
                    resultados.append(resultado_carga(conn, cid, duplicado))
                else:
                    flash(
                        "Contenido ya recibido; consulta su resultado existente."
                        if duplicado
                        else "Archivo recibido. Consulta su resultado y los pendientes en la lista."
                    )
            except ErrorDeTrabajo as exc:
                conn.rollback()
                if cola_json:
                    if cid is None:
                        return jsonify(error=exc.mensaje), exc.status
                    resultados.append(
                        resultado_carga(conn, cid, duplicado, exc.mensaje)
                    )
                else:
                    flash(exc.mensaje, "error")
    if cola_json:
        return jsonify(resultados[0])
    if destino:
        contexto = {
            "institucion_id": destino["institucion_educativa_id"],
            "periodo": periodo.strftime("%Y-%m"),
            "nivel": _nivel,
        }
        if len(archivos) == 1 and recibidas:
            return redirect(
                url_for("pages.documento", cid=recibidas[0], **contexto), code=303
            )
        if desde_calendario:
            return redirect(url_for("pages.subir_calendario", iid=iid), code=303)
        return redirect(url_for("pages.subir_reporte", **contexto), code=303)
    if desde_revision:
        return redirect(
            url_for(
                "pages.revision",
                periodo=periodo.strftime("%Y-%m"),
                nivel=_nivel,
                _anchor="carga-reportes",
            ),
            code=303,
        )
    return redirect(url_for("pages.documentos"), code=303)


@pages.get("/documentos/<uuid:cid>")
def documento(cid):
    from .revision_bandeja import contexto_carga

    destino = contexto_carga(base(), request.args.get("institucion_id"))
    periodo, _nivel = ambito()
    c = lecturas.uno(
        base(),
        "SELECT c.*,o.sha256,o.mime_type FROM web_carga c JOIN objeto_archivo o USING(objeto_archivo_id) WHERE web_carga_id=%s",
        (cid,),
    )
    recepciones = (
        base()
        .execute(
            "SELECT r.*,u.nombre AS autor FROM web_recepcion r JOIN usuario u ON u.usuario_id=r.cargado_por WHERE web_carga_id=%s ORDER BY cargado_en DESC",
            (cid,),
        )
        .fetchall()
    )
    intentos = (
        base()
        .execute(
            "SELECT * FROM web_carga_intento WHERE web_carga_id=%s ORDER BY iniciado_en DESC",
            (cid,),
        )
        .fetchall()
    )
    reportes = (
        base()
        .execute(
            "SELECT r.*,ie.nombre_ie FROM reporte_asistencia r JOIN documento_recibido d USING(documento_recibido_id) LEFT JOIN institucion_educativa ie USING(institucion_educativa_id) WHERE d.objeto_archivo_id=%s ORDER BY r.periodo,r.indice_bloque,r.version DESC",
            (c["objeto_archivo_id"],),
        )
        .fetchall()
    )
    calendarios_carga = (
        base()
        .execute(
            """SELECT v.calendarizacion_version_id,v.version,v.estado,l.anio,
                  l.institucion_educativa_id,ie.nombre_ie
           FROM calendarizacion_version v JOIN calendarizacion_local l USING(calendarizacion_local_id)
           JOIN institucion_educativa ie USING(institucion_educativa_id)
           JOIN documento_recibido d USING(documento_recibido_id)
           WHERE d.objeto_archivo_id=%s ORDER BY l.anio,v.version DESC""",
            (c["objeto_archivo_id"],),
        )
        .fetchall()
        if c["tipo"] == "calendario"
        else []
    )
    return vista(
        "documento.html",
        carga=c,
        recepciones=recepciones,
        intentos=intentos,
        reportes=reportes,
        calendarios_carga=calendarios_carga,
        institucion_destino=destino,
        coincide_destino=bool(
            destino
            and any(
                r["institucion_educativa_id"] == destino["institucion_educativa_id"]
                and r["periodo"] == periodo
                and r["estado"] not in {"RECHAZADO", "HISTORICA"}
                for r in reportes
            )
        ),
    )


@pages.post("/documentos/<uuid:cid>/reintentar")
def reintentar(cid):
    with escritura() as conn:
        procesar(conn, cid, g.usuario["usuario_id"], operacion())
    flash("Resultado comprobado. Revisa el estado de este archivo.")
    return redirect(url_for("pages.documento", cid=cid), code=303)


@pages.get("/documentos/<uuid:cid>/original")
def original_carga(cid):
    c = lecturas.uno(
        base(),
        "SELECT c.ruta_copia,o.sha256 FROM web_carga c JOIN objeto_archivo o USING(objeto_archivo_id) WHERE web_carga_id=%s",
        (cid,),
    )
    return send_file(ruta_verificada(c["ruta_copia"], c["sha256"]), as_attachment=True)


@pages.get("/revision")
def revision():
    from .revision_bandeja import instituciones_revision
    from .revision_cargas import cargas_recientes

    periodo, nivel = ambito()
    return vista(
        "revision.html",
        instituciones=instituciones_revision(base(), periodo, nivel),
        cargas_recientes=cargas_recientes(base(), g.usuario["usuario_id"]),
        destino_ambito="pages.revision",
    )


@pages.get("/reportes/<uuid:rid>")
def reporte(rid):
    from asistia.coherencia import revisar_reportes
    from asistia.monitoreo.reglas import REGLAS

    from .ajustes_asistencia import opciones as opciones_ajuste
    from .reporte_calendario import calendario_reporte
    from .reporte_reglas import contexto_catalogo
    from .reporte_revision import preparar as preparar_revision

    r = lecturas.reporte(base(), rid)
    filas = lecturas.detalle_personas(base(), rid)
    historico = bool(
        base()
        .execute(
            "SELECT 1 FROM reporte_asistencia WHERE reporte_asistencia_serie_id=%s AND version>%s",
            (r["reporte_asistencia_serie_id"], r["version"]),
        )
        .fetchone()
    )
    historia = (
        base()
        .execute(
            "SELECT e.*,u.nombre AS autor FROM web_revision_evento e JOIN usuario u USING(usuario_id) WHERE reporte_asistencia_id=%s ORDER BY creado_en DESC",
            (rid,),
        )
        .fetchall()
    )
    dias = (
        base()
        .execute(
            "SELECT a.*,c.codigo,c.nombre FROM asistencia_dia a "
            "JOIN trabajador_en_reporte t USING(trabajador_en_reporte_id) "
            "LEFT JOIN catalogo_estado_asistencia c USING(estado_asistencia_id) "
            "WHERE t.reporte_asistencia_id=%s ORDER BY t.fila_detalle_origen,a.fecha",
            (rid,),
        )
        .fetchall()
    )
    categorias = leyendas.codigos_documento(
        r, [d for d in dias if d["estado_captura"] in {"REGISTRADO", "DERIVADO"}]
    )
    alertas = lecturas.alertas(base(), rid)
    revision = preparar_revision(r, filas, dias, categorias, alertas, historia)
    coherencia = revisar_reportes(base(), [r])[str(rid)]
    personas_por_id = {str(p["trabajador_en_reporte_id"]): p for p in filas}
    alertados = set()
    for caso in coherencia["casos"]:
        fuente = next(f for f in caso["fuentes"] if f["reporte_id"] == str(rid))
        caso["fuente_reporte"] = fuente
        caso["fuente_declarada"] = next(
            f for f in caso["declarado"]["fuentes"] if f["reporte_id"] == str(rid)
        )
        caso["persona"] = personas_por_id.get(
            fuente["trabajador_en_reporte_id"], {}
        ).get("nombres_reportados_raw", "Personal del reporte")
        if not caso["resuelto"]:
            alertados.update(
                f["asistencia_dia_id"]
                for f in caso["fuentes"]
                if f["asistencia_dia_id"]
            )
    for dia in dias:
        dia["requiere_coherencia"] = str(dia["asistencia_dia_id"]) in alertados
    codigos = (
        base()
        .execute(
            "SELECT codigo,nombre FROM catalogo_estado_asistencia WHERE activo AND codigo<>'OTRO_REPORTADO' ORDER BY codigo"
        )
        .fetchall()
    )
    return vista(
        "reporte.html",
        reporte=r,
        personas=filas,
        alertas=alertas,
        revision=revision,
        coherencia=coherencia,
        ajustes_asistencia=opciones_ajuste(base(), r),
        reglas_coherencia=REGLAS,
        calendario_mes=calendario_reporte(base(), r, revision["fechas"]),
        reglas_revision=contexto_catalogo(base()),
        categorias=categorias,
        codigos=[
            {
                "codigo": c["codigo"],
                "nombre": c.get("tipo_dia") or "Significado pendiente",
            }
            for c in categorias
        ]
        + [c for c in codigos if c["codigo"] not in {k["codigo"] for k in categorias}],
        huella=lecturas.huella_reporte(base(), rid),
        historico=historico,
        mostrar_ambito=False,
    )


@pages.post("/reportes/<uuid:rid>/revisar")
def revisar_reporte(rid):
    with escritura() as conn:
        confirmar_reporte(
            conn,
            rid,
            motivo("Revisión completa registrada; sin comentario adicional."),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Revisión del reporte registrada. Los pendientes de otros reportes siguen visibles."
    )
    return redirect(url_for("pages.reporte", rid=rid), code=303)


@pages.post("/reportes/<uuid:rid>/coherencia")
def revisar_coherencia(rid):
    from .coherencia import decidir

    with escritura() as conn:
        decidir(
            conn,
            rid,
            request.form,
            motivo("Revisión de coherencia registrada; sin comentario adicional."),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Decisión de coherencia registrada. La declaración del documento se conserva."
    )
    return redirect(url_for("pages.reporte", rid=rid) + "#coherencia", code=303)


@pages.get("/reportes/<uuid:rid>/comparacion-dia")
def comparar_asistencia_dia(rid):
    from .comparacion_asistencia import comparar

    conn = base()
    conn.rollback()
    conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
    return jsonify(comparar(conn, rid, request.args))


@pages.post("/reportes/<uuid:rid>/ajustes-asistencia")
def ajustar_asistencia(rid):
    from .ajustes_asistencia import guardar

    with escritura() as conn:
        guardar(
            conn,
            rid,
            request.form,
            motivo(
                "Ajuste de asistencia confirmado por RRHH; sin comentario adicional."
            ),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Ajuste guardado y cruce recalculado. La digitalización se conserva; cualquier diferencia restante sigue pendiente."
    )
    retorno = (
        {"dia_coherencia": request.form.get("dia")}
        if request.form.get("volver_dia")
        else {}
    )
    return redirect(
        url_for("pages.reporte", rid=rid, **retorno) + "#coherencia", code=303
    )


@pages.post("/reportes/<uuid:rid>/alertas/<uuid:aid>")
def resolver(rid, aid):
    with escritura() as conn:
        decidir_alerta(
            conn,
            rid,
            aid,
            request.form.get("accion"),
            motivo("Decisión de observación registrada; sin comentario adicional."),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Decisión y motivo registrados. Esto no revisa automáticamente el reporte completo."
    )
    return redirect(url_for("pages.reporte", rid=rid), code=303)


@pages.get("/reportes/<uuid:rid>/personas/<uuid:tid>")
def persona(rid, tid):
    r = lecturas.reporte(base(), rid)
    p = lecturas.uno(
        base(),
        "SELECT * FROM trabajador_en_reporte WHERE trabajador_en_reporte_id=%s AND reporte_asistencia_id=%s",
        (tid, rid),
    )
    dias = (
        base()
        .execute(
            "SELECT a.*,c.codigo,c.nombre FROM asistencia_dia a LEFT JOIN catalogo_estado_asistencia c USING(estado_asistencia_id) WHERE trabajador_en_reporte_id=%s ORDER BY fecha",
            (tid,),
        )
        .fetchall()
    )
    codigos = (
        base()
        .execute(
            "SELECT codigo,nombre FROM catalogo_estado_asistencia WHERE activo ORDER BY codigo"
        )
        .fetchall()
    )
    consulta = request.args.get("dni", "").strip()
    nombre_consulta = request.args.get("nombre", "").strip()
    candidatos = []
    # Busca en toda la UGEL, no solo en la institución del reporte: encontrar a la persona
    # registrada en otra institución es justamente lo que confirma o descarta una alerta de
    # vínculo automático o "sin coincidencia en NEXUS". Vincularla aquí solo es posible dentro de
    # la institución del reporte (resolver_identidad lo exige); las demás quedan solo como consulta.
    if consulta:
        if len(consulta) != 8 or not consulta.isdecimal():
            raise ErrorDeTrabajo(
                "Busca con los ocho dígitos del DNI que consta en la fuente."
            )
        candidatos = (
            base()
            .execute(
                "SELECT v.*,t.apellido_paterno,t.apellido_materno,t.nombres,p.codigo_plaza,c.nombre AS rol,"
                "ie.nombre_ie,(v.institucion_educativa_id=%s) AS misma_institucion "
                "FROM vinculo_trabajador_ie v JOIN trabajador t USING(trabajador_id) LEFT JOIN plaza p USING(plaza_id) "
                "JOIN catalogo_rol_laboral c USING(rol_laboral_id) "
                "JOIN institucion_educativa ie ON ie.institucion_educativa_id=v.institucion_educativa_id "
                "WHERE t.dni=%s ORDER BY misma_institucion DESC",
                (r["institucion_educativa_id"], consulta),
            )
            .fetchall()
        )
    elif nombre_consulta:
        if len(nombre_consulta) < 3:
            raise ErrorDeTrabajo("Escribe al menos tres letras del apellido o nombre.")
        candidatos = (
            base()
            .execute(
                "SELECT v.*,t.apellido_paterno,t.apellido_materno,t.nombres,p.codigo_plaza,c.nombre AS rol,"
                "ie.nombre_ie,(v.institucion_educativa_id=%s) AS misma_institucion "
                "FROM vinculo_trabajador_ie v JOIN trabajador t USING(trabajador_id) LEFT JOIN plaza p USING(plaza_id) "
                "JOIN catalogo_rol_laboral c USING(rol_laboral_id) "
                "JOIN institucion_educativa ie ON ie.institucion_educativa_id=v.institucion_educativa_id "
                "WHERE (t.apellido_paterno||' '||t.apellido_materno||' '||t.nombres) ILIKE %s "
                "ORDER BY misma_institucion DESC,t.apellido_paterno LIMIT 30",
                (r["institucion_educativa_id"], f"%{nombre_consulta}%"),
            )
            .fetchall()
        )
    return vista(
        "persona.html",
        reporte=r,
        persona=p,
        cruce=cruce_persona(base(), tid),
        dias=[{**d, "clasificacion": categoria_dia(r, d)} for d in dias],
        codigos=[
            {"codigo": k, "nombre": v["tipo_dia"]}
            for k, v in r["clasificacion_codigos"].items()
        ]
        + [
            c
            for c in codigos
            if c["codigo"] not in r["clasificacion_codigos"]
            and c["codigo"] != "OTRO_REPORTADO"
        ],
        candidatos=candidatos,
        dni=consulta,
        nombre_consulta=nombre_consulta,
        huella=lecturas.huella_reporte(base(), rid),
    )


@pages.post("/reportes/<uuid:rid>/vacios")
def revisar_vacios(rid):
    with escritura() as conn:
        cantidad = clasificar_vacios(
            conn,
            rid,
            request.form.getlist("dias"),
            motivo(
                "Vacíos clasificados como no correspondía asistir; sin comentario adicional."
            ),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        f"{cantidad} vacíos clasificados: no correspondía asistir. "
        "Se conservan el original y el historial; no se registraron faltas."
    )
    return redirect(url_for("pages.reporte", rid=rid), code=303)


@pages.post("/reportes/<uuid:rid>/dias/<uuid:did>")
def corregir_dia(rid, did):
    with escritura() as conn:
        corregir_marca(
            conn,
            rid,
            did,
            request.form.get("codigo", ""),
            request.form.get("captura", ""),
            motivo("Lectura registrada desde revisión; sin comentario adicional."),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
        fila = lecturas.uno(
            conn,
            "SELECT trabajador_en_reporte_id FROM asistencia_dia WHERE asistencia_dia_id=%s",
            (did,),
        )
    flash(
        "Corrección guardada con el valor recibido y su historial. El reporte necesita revisarse otra vez."
    )
    if request.form.get("volver") == "reporte":
        return redirect(url_for("pages.reporte", rid=rid), code=303)
    return redirect(
        url_for("pages.persona", rid=rid, tid=fila["trabajador_en_reporte_id"]),
        code=303,
    )


@pages.post("/reportes/<uuid:rid>/personas/<uuid:tid>/identidad")
def identificar(rid, tid):
    try:
        vid = int(request.form.get("vinculo", ""))
    except ValueError:
        raise ErrorDeTrabajo(
            "Selecciona el vínculo sustentado por el original."
        ) from None
    with escritura() as conn:
        resolver_identidad(
            conn,
            rid,
            tid,
            vid,
            motivo(
                "Identidad y vínculo asignados desde revisión; sin comentario adicional."
            ),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Identidad y vínculo registrados con sustento. Revisa los días y las observaciones del reporte."
    )
    return redirect(url_for("pages.persona", rid=rid, tid=tid), code=303)


def ruta_original_reporte(rid):
    r = lecturas.reporte(base(), rid)
    copia = (
        base()
        .execute(
            "SELECT c.ruta_copia FROM web_carga c JOIN objeto_archivo o USING(objeto_archivo_id) WHERE o.sha256=%s LIMIT 1",
            (r["sha256"],),
        )
        .fetchone()
    )
    path = ruta_verificada(
        copia["ruta_copia"] if copia else r["ruta_objeto"], r["sha256"]
    )
    return r, path


@pages.get("/reportes/<uuid:rid>/original")
def original_reporte(rid):
    r, path = ruta_original_reporte(rid)
    return send_file(path, as_attachment=True, download_name=r["nombre_original"])


@pages.get("/reportes/<uuid:rid>/fuente")
def fuente(rid):
    r, path = ruta_original_reporte(rid)
    if request.args.get("vista") == "celdas":
        return fuente_celdas(r, path)
    if request.args.get("archivo") == "1" and path.suffix.lower() in {
        ".jpg",
        ".jpeg",
        ".png",
    }:
        return send_file(path, as_attachment=False)
    from .visor_original import EXCEL

    return render_template(
        "visor_original.html",
        manifest_url=url_for("pages.vista_original", rid=rid),
        original_url=url_for("pages.original_reporte", rid=rid),
        textual_url=url_for(
            "pages.fuente",
            rid=rid,
            vista="celdas",
            **{
                k: request.args[k]
                for k in ("hoja", "celda", "fila", "columna")
                if k in request.args
            },
        )
        if path.suffix.lower() in EXCEL
        else None,
    )


def contexto_visor(rid):
    from .visor_original import EXCEL, hoja_inicial, hojas_excel, preparar

    r, path = ruta_original_reporte(rid)
    hojas = hojas_excel(path) if path.suffix.lower() in EXCEL else []
    hoja = (
        hoja_inicial(hojas, request.args.get("hoja"), r["hoja_pagina_origen"])
        if hojas
        else None
    )
    if path.suffix.lower() in EXCEL and not hoja:
        raise ErrorDeTrabajo("El libro no tiene hojas visibles para consultar.", 422)
    carpeta, meta = preparar(path, r["sha256"], hoja)
    return r, path, hojas, hoja, carpeta, meta


@pages.get("/reportes/<uuid:rid>/fuente/vista")
def vista_original(rid):
    _r, _path, hojas, hoja, _carpeta, meta = contexto_visor(rid)
    return jsonify(
        paginas=meta["paginas"],
        tipo=meta["tipo"],
        hojas=[h["nombre"] for h in hojas],
        hoja=hoja["nombre"] if hoja else None,
        imagen_url=url_for(
            "pages.pagina_original",
            rid=rid,
            pagina=1,
            **({"hoja": hoja["nombre"]} if hoja else {}),
        ),
    )


@pages.get("/reportes/<uuid:rid>/fuente/pagina/<int:pagina>")
def pagina_original(rid, pagina):
    from .visor_original import pagina_png

    _r, path, _hojas, _hoja, carpeta, meta = contexto_visor(rid)
    if not 1 <= pagina <= meta["paginas"]:
        raise ErrorDeTrabajo("Esa página no existe en el documento.", 404)
    if meta["tipo"] == "imagen":
        return send_file(path, as_attachment=False)
    return send_file(
        pagina_png(carpeta, meta, pagina), mimetype="image/png", as_attachment=False
    )


def fuente_celdas(r, path):
    """Alternativa textual explícita; conserva localizadores del recorrido anterior."""
    if path.suffix.lower() == ".pdf":
        return send_file(path, as_attachment=False)
    if path.suffix.lower() in {".jpg", ".jpeg", ".png"}:
        if request.args.get("archivo") == "1":
            return send_file(path, as_attachment=False)
        return render_template("fuente_imagen.html", reporte=r)
    if path.suffix.lower() == ".docx":
        from asistia.cierre.conversion import representar_docx

        try:
            pdf, _ = representar_docx(path)
        except ValueError:
            return render_template("fuente_documental.html", reporte=r)
        return send_file(pdf, mimetype="application/pdf", as_attachment=False)
    # Vista de celdas del original, sin fórmulas activas ni ejecución de macros.
    try:
        fila = max(1, min(100000, int(request.args.get("fila", 1))))
        columna = max(1, min(1000, int(request.args.get("columna", 1))))
    except ValueError:
        raise ErrorDeTrabajo("La posición de la fuente no es válida.") from None
    from .reporte_revision import posicion_celda

    marcada = posicion_celda(request.args.get("celda"))
    if request.args.get("celda") and not marcada:
        raise ErrorDeTrabajo("La celda de la fuente no es válida.")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            hoja = request.args.get("hoja") or (
                r["hoja_pagina_origen"]
                if r["hoja_pagina_origen"] in wb.sheetnames
                else next(
                    (
                        h
                        for h in wb.sheetnames
                        if h.upper().replace(" ", "") in {"ANEXO3", "ANEXO03"}
                    ),
                    wb.sheetnames[0],
                )
            )
            if hoja not in wb.sheetnames:
                raise ErrorDeTrabajo("No se encontró esa hoja en el original.", 404)
            ws = wb[hoja]
            celdas = [
                ["" if c is None else str(c) for c in row]
                for row in ws.iter_rows(
                    min_row=fila,
                    max_row=fila + 39,
                    min_col=columna,
                    max_col=columna + 59,
                    values_only=True,
                )
            ]
            hojas = wb.sheetnames
        finally:
            wb.close()
    elif path.suffix.lower() == ".xls":
        wb = xlrd.open_workbook(path)
        try:
            hojas = wb.sheet_names()
            hoja = request.args.get("hoja") or hojas[0]
            if hoja not in hojas:
                raise ErrorDeTrabajo("No se encontró esa hoja.", 404)
            ws = wb.sheet_by_name(hoja)
            celdas = [
                [
                    str(ws.cell_value(rr, cc))
                    for cc in range(columna - 1, min(ws.ncols, columna + 59))
                ]
                for rr in range(fila - 1, min(ws.nrows, fila + 39))
            ]
        finally:
            wb.release_resources()
    else:
        raise ErrorDeTrabajo(
            "La vista de celdas corresponde a Excel. Descarga el original para revisarlo.",
            409,
        )
    columnas = [
        openpyxl.utils.get_column_letter(c) for c in range(columna, columna + 60)
    ]
    return render_template(
        "fuente.html",
        reporte=r,
        hojas=hojas,
        hoja=hoja,
        fila=fila,
        columna=columna,
        celdas=celdas,
        columnas=columnas,
        marcada=marcada,
    )


@pages.get("/consolidado")
def consolidado():
    from asistia.coherencia import revisar_reportes

    periodo, nivel = ambito()
    reportes = lecturas.reportes_mes(base(), periodo, nivel)
    coherencias = revisar_reportes(base(), reportes)
    return vista(
        "consolidado.html",
        coherencias=coherencias,
        reportes=reportes,
        calendarios_aplicables=lecturas.calendarios_actuales(
            base(), reportes, periodo.year
        ),
        salidas=lecturas.salidas(base(), periodo, nivel),
        revisable=bool(reportes)
        and all(c["listo"] for c in coherencias.values())
        and all(
            r["estado"] == "VALIDADO" and not r["criticos"] and not r["sin_identidad"]
            for r in reportes
        ),
    )


@pages.post("/consolidado")
def generar():
    if not request.form.get("periodo") or not request.form.get("nivel"):
        raise ErrorDeTrabajo(
            "El formulario debe indicar su mes y nivel. Actualiza la página del consolidado."
        )
    periodo, nivel = ambito()
    estado = request.form.get("estado")
    if estado not in {"BORRADOR", "REVISADO"}:
        raise ErrorDeTrabajo("Selecciona el estado de la salida.")
    with escritura() as conn:
        sid = preparar(
            conn,
            periodo,
            nivel,
            estado == "REVISADO",
            operacion(),
            g.usuario["usuario_id"],
        )
    flash("Salida conservada. La descarga corresponde a esta versión.")
    return redirect(url_for("pages.salida", sid=sid), code=303)


@pages.get("/salidas/<uuid:sid>")
def salida(sid):
    s = lecturas.salida(base(), sid)
    return vista(
        "salida.html",
        salida=s,
        cambios=lecturas.cambios_posteriores(base(), s),
        descarga_desarrollo=current_app.config["DEVELOPMENT_EXPORTS"],
    )


@pages.get("/salidas/<uuid:sid>/descargar")
def descargar(sid):
    from .exportacion import primera_hoja

    detalle = request.args.get("detalle", "0")
    if detalle not in {"0", "1"}:
        raise ErrorDeTrabajo("Selecciona una modalidad de descarga válida.")
    if detalle == "1" and not current_app.config["DEVELOPMENT_EXPORTS"]:
        raise ErrorDeTrabajo("Esta descarga no está disponible.", 404)
    s = lecturas.salida(base(), sid)
    path = ruta_verificada(s["ruta_objeto"], s["sha256"])
    sufijo = "_completo_desarrollo" if detalle == "1" else ""
    return send_file(
        path if detalle == "1" else primera_hoja(path),
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        download_name=f"ASISTIA_{s['nivel_modalidad']}_{s['periodo']:%Y-%m}_v{s['version']}_{s['estado_revision']}{sufijo}.xlsx",
    )


@pages.get("/instituciones")
def instituciones():
    conn = base()
    # Totales, detalle y movimientos pertenecen a un mismo snapshot de lectura.
    conn.rollback()
    conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
    datos = directorio(conn, **filtros_directorio(request.args))
    return render_template("instituciones.html", mostrar_ambito=False, **datos)


@pages.get("/instituciones/<int:iid>")
def institucion(iid):
    periodo, _nivel = ambito()
    ie = lecturas.uno(
        base(),
        "SELECT * FROM institucion_educativa WHERE institucion_educativa_id=%s",
        (iid,),
    )
    personas = (
        base()
        .execute(
            "SELECT v.*,t.dni,t.apellido_paterno,t.apellido_materno,t.nombres,p.codigo_plaza,c.nombre AS rol "
            "FROM vinculo_trabajador_ie v LEFT JOIN trabajador t USING(trabajador_id) LEFT JOIN plaza p USING(plaza_id) "
            "JOIN catalogo_rol_laboral c USING(rol_laboral_id) WHERE v.institucion_educativa_id=%s ORDER BY t.apellido_paterno",
            (iid,),
        )
        .fetchall()
    )
    revision_cierre = (
        base()
        .execute(
            """SELECT * FROM cierre_revision_institucion
      WHERE institucion_educativa_id=%s AND periodo=%s ORDER BY revisado_en DESC LIMIT 1""",
            (iid, periodo),
        )
        .fetchone()
    )
    contexto_directorio = {
        k: request.args[k]
        for k in ("q", "nivel_ie", "actualizado", "fc")
        if k in request.args
    }
    calendario = calendario_institucion(
        base(),
        iid,
        anio=request.args.get("anio_calendario"),
        mes=request.args.get("mes"),
        mes_preferido=periodo.month,
    )
    return vista(
        "institucion.html",
        retorno_directorio=url_for("pages.instituciones", **contexto_directorio),
        contexto_directorio=contexto_directorio,
        institucion=ie,
        personas=personas,
        revision_cierre=revision_cierre,
        reportes_ie=reportes_institucion(base(), iid),
        nombres_meses=MESES,
        mostrar_ambito=False,
        hoy=datetime.now(ZoneInfo("America/Lima")).date(),
        **calendario,
    )


@pages.route("/instituciones/nueva", methods=["GET", "POST"])
@pages.route("/instituciones/<int:iid>/editar", methods=["GET", "POST"])
def institucion_editar(iid=None):
    contexto = {
        k: request.args[k]
        for k in ("q", "nivel_ie", "actualizado", "fc")
        if k in request.args
    }
    ie = instituciones_edicion.leer(base(), iid) if iid is not None else None
    error, status = None, 200
    if request.method == "POST":
        try:
            with escritura() as conn:
                destino = instituciones_edicion.guardar(
                    conn,
                    iid=iid,
                    form=request.form,
                    op=operacion(),
                    autor=g.usuario["usuario_id"],
                    responsable=g.usuario["nombre"],
                )
            flash(
                "Institución creada."
                if iid is None
                else "Datos de la institución actualizados."
            )
            return redirect(
                url_for("pages.institucion", iid=destino, **contexto), code=303
            )
        except ErrorDeTrabajo as exc:
            error, status = exc.mensaje, exc.status
    return render_template(
        "institucion_editar.html",
        institucion=ie,
        campos=instituciones_edicion.CAMPOS,
        datos=request.form
        if request.method == "POST"
        else ie or {"anexo": "0", "departamento": "AMAZONAS", "provincia": "LUYA"},
        huella=request.form.get("huella", "")
        if request.method == "POST"
        else instituciones_edicion.huella(ie)
        if ie
        else "",
        op=request.form.get("operacion")
        if request.method == "POST"
        else str(uuid.uuid4()),
        opciones_nivel=[
            r["nivel_modalidad"]
            for r in base().execute(
                "SELECT DISTINCT nivel_modalidad FROM institucion_educativa ORDER BY 1"
            )
        ],
        cambios=instituciones_edicion.historial(base(), iid) if iid is not None else [],
        volver=url_for("pages.institucion", iid=iid, **contexto)
        if iid is not None
        else url_for("pages.instituciones", **contexto),
        error=error,
        mostrar_ambito=False,
    ), status


@pages.get("/cobertura")
def cobertura():
    """Estado de resolucion del mes por institucion, con la estimacion declarada de esfuerzo.

    El desglose es exhaustivo sobre el universo operativo de NEXUS; la plantilla debe mostrar esa
    limitacion junto al numero (no es la cobertura validada de M04) y la estimacion siempre con
    sus supuestos.
    """
    periodo, nivel = ambito()
    resumen = estado_resolucion(base(), periodo, nivel)
    # Se pasa la declaracion tambien al nivel: asi el resultado dice que NO APLICA a este
    # grano (es por UGEL-mes) en vez de afirmar que falta una declaracion que si existe.
    estimacion = estimar(resumen, declaracion=DECLARACION_VIGENTE)
    filtro = request.args.get("estado", "todos")
    visibles = [
        i for i in resumen.instituciones if filtro == "todos" or i.estado == filtro
    ]
    # El tiempo manual declarado es de toda la UGEL: solo el agregado puede convertirlo en
    # tiempo absoluto sin multiplicarlo por cada nivel.
    agregado = resumen_ugel(base(), periodo)
    return vista(
        "cobertura.html",
        resumen=resumen,
        estimacion=estimacion,
        agregado=agregado,
        estimacion_ugel=estimar(agregado, declaracion=DECLARACION_VIGENTE),
        instituciones=visibles,
        estados=ESTADOS_RESOLUCION,
        etiquetas_estado=ETIQUETAS_ESTADO,
        filtro=filtro,
        sin_nivel_canonico=len(instituciones_sin_nivel_canonico(base())),
    )


@pages.get("/cobertura.csv")
def cobertura_csv():
    """Mismo calculo que /cobertura, descargable. Una cifra citada en un informe debe poder
    reproducirse desde el archivo que la origino."""
    periodo, nivel = ambito()
    resumen = estado_resolucion(base(), periodo, nivel)
    buffer = io.StringIO()
    w = csv.writer(buffer)
    w.writerow(["# universo", resumen.universo])
    w.writerow(
        [
            "periodo",
            "nivel",
            "cod_mod",
            "anexo",
            "institucion",
            "nivel_recibido",
            "estado",
            "reportes",
            "reportes_revisados",
            "version_maxima",
            "filas",
            "sin_identidad",
            "alertas_pendientes",
            "alertas_criticas",
            "calendario_estado",
            "calendario_dias_sin_determinar",
        ]
    )
    for i in resumen.instituciones:
        w.writerow(
            [
                resumen.periodo.isoformat(),
                resumen.nivel,
                i.cod_mod,
                i.anexo,
                i.nombre_ie,
                i.nivel_modalidad,
                i.estado,
                i.reportes,
                i.reportes_revisados,
                i.version_maxima,
                i.filas,
                i.sin_identidad,
                i.alertas_pendientes,
                i.alertas_criticas,
                i.calendario_estado,
                i.calendario_dias_sin_determinar,
            ]
        )
    nombre = f"cobertura-{resumen.nivel}-{resumen.periodo:%Y-%m}.csv"
    return Response(
        buffer.getvalue(),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


@pages.get("/calendarios")
def calendarios():
    periodo, nivel = ambito()
    filas = (
        base()
        .execute(
            "SELECT cv.*,cl.anio,ie.nombre_ie,ie.cod_mod,ie.anexo FROM calendarizacion_version cv "
            "JOIN calendarizacion_local cl USING(calendarizacion_local_id) JOIN institucion_educativa ie USING(institucion_educativa_id) "
            "WHERE cl.anio=%s AND fn_nivel_canonico(ie.nivel_modalidad)=%s ORDER BY ie.nombre_ie,cv.version DESC",
            (periodo.year, nivel),
        )
        .fetchall()
    )
    return vista(
        "calendarios.html",
        calendarios=filas,
        codigos_desconocidos=codigos_desconocidos(base(), _FAMILIA_FORMATO_DEFECTO),
        tipos_dia=tipos_dia(base()),
        clasificaciones=clasificaciones_vigentes(base(), _FAMILIA_FORMATO_DEFECTO),
    )


@pages.post("/calendarios/codigos")
def clasificar_codigo_calendario():
    with escritura() as conn:
        clasificar_codigo(
            conn,
            _FAMILIA_FORMATO_DEFECTO,
            request.form.get("codigo_raw", ""),
            request.form.get("tipo_dia_id", type=int) or 0,
            motivo(),
            g.usuario["usuario_id"],
        )
    flash(
        "Código clasificado. Se aplicará a la próxima importación o reimportación del calendario; "
        "no cambia los días ya cargados con ese código."
    )
    return redirect(url_for("pages.calendarios"), code=303)


@pages.get("/calendarios/<uuid:cvid>")
def calendario(cvid):
    from .calendario_anual import calendario_anual
    from .calendario_revision import propuesta
    from .reporte_reglas import contexto_catalogo

    cv = lecturas.uno(
        base(),
        "SELECT cv.*,cl.anio,cl.institucion_educativa_id,ie.nombre_ie FROM calendarizacion_version cv JOIN calendarizacion_local cl USING(calendarizacion_local_id) JOIN institucion_educativa ie USING(institucion_educativa_id) WHERE calendarizacion_version_id=%s",
        (cvid,),
    )
    dias = (
        base()
        .execute(
            "SELECT d.*,c.nombre AS tipo,c.codigo_interno AS grupo_recibido FROM dia_calendarizacion d LEFT JOIN catalogo_tipo_dia c USING(tipo_dia_id) WHERE calendarizacion_version_id=%s ORDER BY fecha",
            (cvid,),
        )
        .fetchall()
    )
    sin_determinar = sum(1 for d in dias if d["estado_captura"] != "REGISTRADO")
    bloqueantes = dias_bloqueantes(base(), cvid)
    puede_aprobar = (
        cv["estado"] in ESTADOS_APROBABLES
        and not bloqueantes
        and cv["procedencia_extraccion"].get("tipo") != "DERIVADO_2025"
    )
    ultimo = lecturas.uno(
        base(),
        "SELECT calendarizacion_version_id,version FROM calendarizacion_version WHERE calendarizacion_local_id=%s ORDER BY version DESC LIMIT 1",
        (cv["calendarizacion_local_id"],),
    )
    actual = ultimo["calendarizacion_version_id"] == cvid
    periodo, _ = ambito()
    meses_con_datos = {d["fecha"].month for d in dias}
    mes_inicial = (
        periodo.month
        if periodo.month in meses_con_datos
        else min(meses_con_datos, default=1)
    )
    try:
        mes = int(request.args.get("mes", mes_inicial))
        if not 1 <= mes <= 12:
            raise ValueError
    except ValueError:
        raise ErrorDeTrabajo("Selecciona un mes entre enero y diciembre.") from None
    revision_calendario = propuesta(base(), cvid)
    efectivo = {**cv, "clasificacion_codigos": revision_calendario["categorias"]}
    anual = calendario_anual(cv, dias)
    celdas = {
        c["fecha"]: c
        for m in meses_calendario(efectivo, dias)
        for s in m["semanas"]
        for c in s
        if c
    }
    for m in anual["meses"]:
        m["celdas"] = [
            {**d, "edicion": celdas[d["fecha"]]} if d else None for d in m["celdas"]
        ]
    return vista(
        "calendario.html",
        calendario=cv,
        revision_calendario=revision_calendario,
        anual=anual,
        reglas_calendario=contexto_catalogo(base(), "calendario"),
        dias=[{**d, "clasificacion": categoria_dia(cv, d)} for d in dias],
        categorias=leyendas.codigos_documento(efectivo, dias),
        huella=leyendas.huella_calendario(base(), cvid),
        sin_determinar=sin_determinar,
        bloqueantes=bloqueantes,
        puede_aprobar=puede_aprobar,
        meses=meses_calendario(cv, dias),
        mes_visible=mes,
        opciones_dias=opciones_dias(efectivo),
        editable_dias=actual and cv["estado"] not in {"HISTORICA", "RECHAZADA"},
        ultima_version=ultimo,
        es_ultima=actual,
        hoy=datetime.now(ZoneInfo("America/Lima")).date(),
    )


@pages.post("/calendarios/<uuid:cvid>/dias")
def asignar_dias_calendario(cvid):
    mes = request.form.get("mes", type=int)
    if mes is not None and not 1 <= mes <= 12:
        raise ErrorDeTrabajo("Selecciona un mes entre enero y diciembre.")
    with escritura() as conn:
        nuevo = guardar_dias(
            conn,
            cvid,
            request.form,
            motivo("Asignación de días registrada; sin comentario adicional."),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Días guardados en una nueva versión en borrador. La versión anterior se conserva; revisa el calendario antes de aprobarlo."
    )
    return redirect(url_for("pages.calendario", cvid=nuevo, mes=mes), code=303)


@pages.post("/calendarios/<uuid:cvid>/leyenda")
def leyenda_calendario(cvid):
    with escritura() as conn:
        nuevo = leyendas.guardar_calendario(
            conn,
            cvid,
            request.form,
            motivo(
                "Clasificación de símbolo del calendario; sin comentario adicional."
            ),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Leyenda guardada en una nueva versión en borrador. Comprueba sus días antes de aprobar."
    )
    return redirect(url_for("pages.calendario", cvid=nuevo), code=303)


@pages.post("/reportes/<uuid:rid>/leyenda")
def leyenda_asistencia(rid):
    with escritura() as conn:
        leyendas.guardar_asistencia(
            conn,
            rid,
            request.form,
            motivo("Clasificación de leyenda registrada; sin comentario adicional."),
            request.form.get("huella"),
            operacion(),
            g.usuario["usuario_id"],
        )
    flash(
        "Significado guardado para este reporte. Si ya no quedan observaciones, pulsa «Guardar revisión» junto al estado."
    )
    return redirect(url_for("pages.reporte", rid=rid), code=303)


@pages.post("/calendarios/<uuid:cvid>/aprobar")
def aprobar_calendario(cvid):
    with escritura() as conn:
        marcar_vigente(conn, cvid, g.usuario["usuario_id"])
    flash(
        "Calendario aprobado: queda vigente y disponible para calcular días esperados."
    )
    return redirect(url_for("pages.calendario", cvid=cvid), code=303)


@pages.post("/calendarios/<uuid:cvid>/confirmar")
def confirmar_calendario(cvid):
    from .calendario_revision import confirmar

    cv = lecturas.uno(
        base(),
        "SELECT cl.* FROM calendarizacion_version cv JOIN calendarizacion_local cl USING(calendarizacion_local_id) WHERE calendarizacion_version_id=%s",
        (cvid,),
    )
    mes = request.form.get("mes", 1, type=int)
    if mes is None or not 1 <= mes <= 12:
        raise ErrorDeTrabajo("Selecciona un mes entre enero y diciembre.")
    destino = request.form.get("destino", "calendario")
    rid = request.form.get("reporte_id")
    if destino == "reporte":
        try:
            rid = uuid.UUID(rid or "")
        except ValueError:
            raise ErrorDeTrabajo("El reporte de retorno no es válido.") from None
        r = lecturas.reporte(base(), rid)
        if (
            r["institucion_educativa_id"] != cv["institucion_educativa_id"]
            or r["periodo"].year != cv["anio"]
        ):
            raise ErrorDeTrabajo("El reporte no corresponde a esta institución y año.")
    elif destino not in {"calendario", "institucion"}:
        raise ErrorDeTrabajo("El destino de retorno no es válido.")
    with escritura() as conn:
        nuevo = confirmar(
            conn, cvid, request.form.get("huella"), operacion(), g.usuario["usuario_id"]
        )
    flash(
        "Calendarización anual confirmada: categorías y días revisados, calendario vigente. Puedes volver a editarlo cuando lo necesites."
    )
    if destino == "reporte":
        return redirect(
            url_for("pages.reporte", rid=rid) + "#report-calendar-heading", code=303
        )
    if destino == "institucion":
        return redirect(
            url_for(
                "pages.institucion",
                iid=cv["institucion_educativa_id"],
                anio_calendario=cv["anio"],
                mes=mes,
            )
            + "#calendario-institucion-titulo",
            code=303,
        )
    return redirect(url_for("pages.calendario", cvid=nuevo, mes=mes), code=303)


@pages.get("/calendarios/<uuid:cvid>/original")
def original_calendario(cvid):
    o, path = ruta_original_calendario(cvid)
    return send_file(path, as_attachment=True, download_name=o["nombre_original"])


def ruta_original_calendario(cvid):
    o = lecturas.uno(
        base(),
        "SELECT o.*,d.nombre_original FROM calendarizacion_version cv JOIN documento_recibido d USING(documento_recibido_id) JOIN objeto_archivo o USING(objeto_archivo_id) WHERE calendarizacion_version_id=%s",
        (cvid,),
    )
    copia = (
        base()
        .execute(
            "SELECT ruta_copia FROM web_carga WHERE objeto_archivo_id=%s LIMIT 1",
            (o["objeto_archivo_id"],),
        )
        .fetchone()
    )
    return o, ruta_verificada(
        copia["ruta_copia"] if copia else o["ruta_objeto"], o["sha256"]
    )


@pages.get("/calendarios/<uuid:cvid>/fuente")
def fuente_calendario(cvid):
    ruta_original_calendario(cvid)
    return render_template(
        "visor_original.html",
        manifest_url=url_for("pages.vista_calendario", cvid=cvid),
        original_url=url_for("pages.original_calendario", cvid=cvid),
        textual_url=None,
    )


def contexto_visor_calendario(cvid):
    from .visor_original import EXCEL, hoja_inicial, hojas_excel, preparar

    o, path = ruta_original_calendario(cvid)
    hojas = hojas_excel(path) if path.suffix.lower() in EXCEL else []
    origen = (
        base()
        .execute(
            "SELECT hoja_origen FROM dia_calendarizacion WHERE calendarizacion_version_id=%s "
            "AND hoja_origen IS NOT NULL ORDER BY fecha LIMIT 1",
            (cvid,),
        )
        .fetchone()
    )
    hoja = (
        hoja_inicial(
            hojas, request.args.get("hoja"), origen["hoja_origen"] if origen else None
        )
        if hojas
        else None
    )
    if path.suffix.lower() in EXCEL and not hoja:
        raise ErrorDeTrabajo("El libro no tiene hojas visibles para consultar.", 422)
    carpeta, meta = preparar(path, o["sha256"], hoja)
    return path, hojas, hoja, carpeta, meta


@pages.get("/calendarios/<uuid:cvid>/fuente/vista")
def vista_calendario(cvid):
    _path, hojas, hoja, _carpeta, meta = contexto_visor_calendario(cvid)
    return jsonify(
        paginas=meta["paginas"],
        tipo=meta["tipo"],
        hojas=[h["nombre"] for h in hojas],
        hoja=hoja["nombre"] if hoja else None,
        imagen_url=url_for(
            "pages.pagina_calendario",
            cvid=cvid,
            pagina=1,
            **({"hoja": hoja["nombre"]} if hoja else {}),
        ),
    )


@pages.get("/calendarios/<uuid:cvid>/fuente/pagina/<int:pagina>")
def pagina_calendario(cvid, pagina):
    from .visor_original import pagina_png

    path, _hojas, _hoja, carpeta, meta = contexto_visor_calendario(cvid)
    if not 1 <= pagina <= meta["paginas"]:
        raise ErrorDeTrabajo("Esa página no existe en el documento.", 404)
    return send_file(
        path if meta["tipo"] == "imagen" else pagina_png(carpeta, meta, pagina),
        as_attachment=False,
    )


@pages.get("/resultados")
def resultados():
    from .trabajo import tareas

    salidas = (
        base()
        .execute(
            "SELECT w.web_salida_id,w.estado_revision,c.periodo,c.nivel_modalidad,c.version FROM web_salida w JOIN consolidado_dre c USING(consolidado_dre_id) ORDER BY c.generado_en DESC"
        )
        .fetchall()
    )
    return vista(
        "resultados.html",
        tareas=tareas(base(), g.usuario["usuario_id"]),
        salidas_trabajo=salidas,
    )


@pages.post("/trabajo")
def trabajo():
    from .trabajo import registrar_trabajo

    periodo, nivel = ambito()
    with escritura() as conn:
        registrar_trabajo(
            conn, g.usuario["usuario_id"], periodo, nivel, request.form, operacion()
        )
    flash(
        "Registro de trabajo guardado. La evaluación de calidad y el recorrido observado se registran por separado."
    )
    return redirect(url_for("pages.resultados"), code=303)


@pages.get("/resultados/registro.csv")
def registro_csv():
    from .trabajo import tareas

    filas = tareas(base(), g.usuario["usuario_id"])
    campos = [
        "web_tarea_id",
        "cohorte",
        "tarea_referencia",
        "lote_referencia",
        "version_referencia",
        "criterio_terminacion",
        "usuario_id",
        "periodo",
        "nivel",
        "uso",
        "estado",
        "ayuda_tecnica",
        "segundos_capturados",
        "intervalo_abierto",
        "salida_id",
        "creado_en",
        "terminado_en",
    ]
    buffer = io.StringIO()
    w = csv.DictWriter(buffer, fieldnames=campos, extrasaction="ignore")
    w.writeheader()
    for fila in filas:
        # Texto libre de cohorte/referencia nunca se ejecuta como fórmula al abrir CSV.
        w.writerow(
            {
                k: (
                    "'" + v
                    if isinstance(v, str) and v.startswith(("=", "+", "-", "@"))
                    else v
                )
                for k, v in fila.items()
            }
        )
    return Response(
        "\ufeff" + buffer.getvalue(),
        mimetype="text/csv",
        headers={
            "Content-Disposition": "attachment; filename=asistia_registro_trabajo.csv"
        },
    )


@pages.get("/resultados/evidencia.json")
def evidencia_trabajo():
    from .trabajo import exportar_captura

    return Response(
        json.dumps(
            exportar_captura(base(), g.usuario["usuario_id"]),
            ensure_ascii=False,
            default=str,
        ),
        mimetype="application/json",
        headers={
            "Content-Disposition": "attachment; filename=asistia_captura_trabajo.json"
        },
    )


from .cierre import registrar_rutas

registrar_rutas(pages)

from .experimental import registrar_rutas as registrar_experimental

registrar_experimental(pages)


from .catalogo import registrar_rutas as registrar_catalogo

registrar_catalogo(pages)

from .monitoreo import registrar_rutas as registrar_monitoreo

registrar_monitoreo(pages)

from .monitoreo_diario import registrar_rutas as registrar_monitoreo_diario

registrar_monitoreo_diario(pages)

from .calidad_leyendas import registrar_rutas as registrar_calidad_leyendas

registrar_calidad_leyendas(pages)

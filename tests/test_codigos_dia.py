"""Feedback RRHH 2026-09-12: códigos de calendario que las instituciones inventan (ej. "P" de
Planificación) quedaban CODIGO_DESCONOCIDO sin que nadie pudiera clasificarlos. Prueba
`asistia.web.codigos_dia` (agrupación, clasificación, no-duplicados) y que la clasificación se
aplique desde la próxima importación sin tocar los días ya cargados con ese código.

`catalogo_codigo_tipo_dia_fuente` es una tabla `catalogo_%`: `tests/conftest.py` la excluye a
propósito del truncado entre pruebas (es catálogo, no un hecho por período). Por eso cada prueba
usa un código propio (prefijo ZTEST) y lo limpia al inicio, para no depender de si esta suite ya
corrió antes contra la misma base efímera.
"""

from __future__ import annotations

from datetime import date

import openpyxl

from asistia.importar.calendario import _FAMILIA_FORMATO_DEFECTO, importar_calendario
from asistia.importar.nexus import importar_nexus
from asistia.web import ErrorDeTrabajo
from asistia.web.auth import crear_usuario
from asistia.web.codigos_dia import (
    clasificaciones_vigentes,
    clasificar_codigo,
    codigos_desconocidos,
)

from .fixtures_excel import construir_nexus_minimo


def _limpiar_codigo(conn, codigo_raw):
    conn.execute(
        "DELETE FROM catalogo_codigo_tipo_dia_fuente WHERE familia_formato=%s AND codigo_raw=%s",
        (_FAMILIA_FORMATO_DEFECTO, codigo_raw),
    )
    conn.commit()


def _preparar_institucion_calendario(conn, tmp_path, cod_mod="9999006"):
    nexus = construir_nexus_minimo(
        [
            {
                "CODMOD I.E.": cod_mod,
                "CODIGO DE PLAZA": "PZ-CAL-01",
                "TIPO DE TRABAJADOR": "DOCENTE",
                "SITUACION LABORAL": "NOMBRADO",
                "TIPO DE REGISTRO": "ORGANICA",
                "DOCUMENTO DE IDENTIDAD": "10000006",
                "APELLIDO PATERNO": "Calendario",
                "APELLIDO MATERNO": "Prueba",
                "NOMBRES": "Docente Uno",
                "FECHA DE INICIO": date(2026, 1, 1),
                "CODIGO LOCAL": cod_mod + "0",
                "DISTRITO": "LUYA",
                "NIVEL EDUCATIVO": "Primaria",
                "NOMBRE DE LA INSTITUCION EDUCATIVA": f"{cod_mod} IE CALENDARIO PRUEBA",
                "CARGO": "DOCENTE",
            }
        ],
        tmp_path / "nexus_calendario.xlsx",
    )
    importar_nexus(conn, nexus, date(2026, 6, 1))


def _calendario_con_codigo(ruta, cod_mod, codigo_dia1):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "CALENDARIZACIÓN DEL AÑO ESCOLAR 2026"
    ws["A2"] = "NOMBRE DE LA IE"
    ws["B2"] = f"{cod_mod} IE CALENDARIO PRUEBA"
    ws["A3"] = "NIVEL O CICLO"
    ws["B3"] = "PRIMARIA"
    ws["A12"] = "JULIO"
    ws.cell(12, 2, 1)  # 1 de julio
    ws.cell(13, 2, codigo_dia1)
    ws.cell(12, 3, 2)  # 2 de julio, siempre lectivo, para no depender de un solo dia
    ws.cell(13, 3, "L")
    wb.save(ruta)
    return ruta


def test_codigos_desconocidos_agrupa_por_codigo_y_cuenta_dias(conn, tmp_path):
    _limpiar_codigo(conn, "ZTEST1")
    _preparar_institucion_calendario(conn, tmp_path)
    archivo = _calendario_con_codigo(tmp_path / "cal1.xlsx", "9999006", "ZTEST1")
    resultado = importar_calendario(conn, archivo)
    assert resultado.error is None
    assert resultado.dias_codigo_desconocido == 1

    filas = codigos_desconocidos(conn, _FAMILIA_FORMATO_DEFECTO)
    assert any(f["codigo_raw"] == "ZTEST1" and f["dias"] == 1 for f in filas)


def test_clasificar_codigo_nuevo_y_rechaza_duplicado(conn, tmp_path):
    _limpiar_codigo(conn, "ZTEST2")
    _preparar_institucion_calendario(conn, tmp_path)
    archivo = _calendario_con_codigo(tmp_path / "cal2.xlsx", "9999006", "ZTEST2")
    importar_calendario(conn, archivo)
    uid = crear_usuario(
        conn,
        "Especialista de prueba",
        "especialista@example.invalid",
        "ClaveDePrueba_2026x",
    )
    tipo_gestion = conn.execute(
        "SELECT tipo_dia_id FROM catalogo_tipo_dia WHERE codigo_interno='GESTION'"
    ).fetchone()["tipo_dia_id"]

    clasificar_codigo(
        conn,
        _FAMILIA_FORMATO_DEFECTO,
        "ZTEST2",
        tipo_gestion,
        "Planificación: programado y pagado, no es dictado de clases.",
        uid,
    )
    conn.commit()

    vigentes = clasificaciones_vigentes(conn, _FAMILIA_FORMATO_DEFECTO)
    fila = next(f for f in vigentes if f["codigo_raw"] == "ZTEST2")
    assert fila["tipo_dia_id"] == tipo_gestion
    assert fila["autor"] == "Especialista de prueba"
    assert "Planificación" in fila["motivo"]

    # Ya clasificado: no aparece mas como pendiente.
    assert not any(
        f["codigo_raw"] == "ZTEST2"
        for f in codigos_desconocidos(conn, _FAMILIA_FORMATO_DEFECTO)
    )

    try:
        clasificar_codigo(
            conn, _FAMILIA_FORMATO_DEFECTO, "ZTEST2", tipo_gestion, "Otra vez.", uid
        )
        raise AssertionError(
            "debia rechazar una segunda clasificacion del mismo codigo"
        )
    except ErrorDeTrabajo as exc:
        assert exc.status == 409
        assert "ya está clasificado" in exc.mensaje


def test_clasificacion_no_reclasifica_dias_ya_importados(conn, tmp_path):
    """La clasificacion aplica a la proxima importacion; los dias CODIGO_DESCONOCIDO ya guardados
    de una version anterior no cambian en silencio."""
    _limpiar_codigo(conn, "ZTEST3")
    _preparar_institucion_calendario(conn, tmp_path)
    primero = _calendario_con_codigo(tmp_path / "cal_v1.xlsx", "9999006", "ZTEST3")
    r1 = importar_calendario(conn, primero)
    assert r1.dias_codigo_desconocido == 1
    dia_v1_antes = conn.execute(
        "SELECT estado_captura,tipo_dia_id FROM dia_calendarizacion "
        "WHERE calendarizacion_version_id=%s AND fecha='2026-07-01'",
        (r1.calendarizacion_version_id,),
    ).fetchone()
    assert dia_v1_antes["estado_captura"] == "CODIGO_DESCONOCIDO"

    uid = crear_usuario(
        conn, "Especialista dos", "especialista2@example.invalid", "ClaveDePrueba_2026y"
    )
    tipo_gestion = conn.execute(
        "SELECT tipo_dia_id FROM catalogo_tipo_dia WHERE codigo_interno='GESTION'"
    ).fetchone()["tipo_dia_id"]
    clasificar_codigo(
        conn, _FAMILIA_FORMATO_DEFECTO, "ZTEST3", tipo_gestion, "Motivo de prueba.", uid
    )
    conn.commit()

    # La version ya importada no cambia con la clasificacion nueva.
    dia_v1_despues = conn.execute(
        "SELECT estado_captura,tipo_dia_id FROM dia_calendarizacion "
        "WHERE calendarizacion_version_id=%s AND fecha='2026-07-01'",
        (r1.calendarizacion_version_id,),
    ).fetchone()
    assert dia_v1_despues == dia_v1_antes

    # Reimportar (nueva version) SI aplica la clasificacion nueva.
    segundo = _calendario_con_codigo(tmp_path / "cal_v2.xlsx", "9999006", "ZTEST3")
    r2 = importar_calendario(conn, segundo)
    assert r2.dias_codigo_desconocido == 0
    assert r2.dias_registrados == 2
    dia_v2 = conn.execute(
        "SELECT estado_captura,tipo_dia_id FROM dia_calendarizacion "
        "WHERE calendarizacion_version_id=%s AND fecha='2026-07-01'",
        (r2.calendarizacion_version_id,),
    ).fetchone()
    assert dia_v2["estado_captura"] == "REGISTRADO"
    assert dia_v2["tipo_dia_id"] == tipo_gestion

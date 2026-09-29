"""P05: revision local (fuente/autor/motivo de correcciones automaticas), clasificacion de faltas
congelada por version (corrige DT03) y recuento de alertas reconciliado entre generacion y
exportacion (corrige DT04). Usa el caso real de reubicacion automatica de Tactamal (P02,
docs/casos/TACTAMAL_JULIO_2026.md §5, §9) para la parte verificable contra datos reales.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl

import asistia.consolidado.generar as generar_mod
from asistia.consolidado.generar import exportar_consolidado_xlsx, generar_consolidado
from asistia.consolidado.revision import (
    evaluar_revision,
    listar_alertas_con_fuente,
    resolver_alerta,
)
from asistia.importar.asistencia import importar_asistencia
from asistia.importar.nexus import importar_nexus

from .fixtures_excel import construir_anexo3_minimo

REPO_ROOT = Path(__file__).resolve().parent.parent
BASE = REPO_ROOT / "data/raw/data_brindada_por_ugel"
RUTA_NEXUS_JUNIO = BASE / "nexus/nexus 2026-06-01.xlsx"
RUTA_NEXUS_ABRIL = BASE / "nexus/NEXUS LUYA AL 01-04-2026.xls"
RUTA_ASISTENCIA = (
    BASE / "ASISTENCIAS JULIO 2026/ASISTENCIA SECUNDARIA JULIO"
    "/ASISTENCIA JULIO IEPySM N  18091 TACTAMAL (ANEXOS 3y4) UGEL LUYA.xlsx"
)


def _generar_tactamal(conn):
    importar_nexus(conn, RUTA_NEXUS_JUNIO, date(2026, 6, 1))
    importar_nexus(conn, RUTA_NEXUS_ABRIL, date(2026, 4, 1))
    importar_asistencia(conn, RUTA_ASISTENCIA)
    return generar_consolidado(conn, date(2026, 7, 1), "Primaria")


def test_clasificacion_faltas_congelada_no_cambia_con_la_regla(
    conn, tmp_path, monkeypatch
):
    """DT03: cambiar la regla de clasificacion despues de generar no debe cambiar la salida de una
    version ya generada al re-exportarla."""
    resultado = _generar_tactamal(conn)
    assert resultado.error is None

    salida_original = tmp_path / "original.xlsx"
    exportar_consolidado_xlsx(conn, resultado.consolidado_dre_id, salida_original)

    # cambia la regla DESPUES de generar (simula que el codigo evoluciona la clasificacion)
    monkeypatch.setitem(generar_mod._CLASIFICACION_FALTA, "L", "INJUSTIFICADA")

    salida_repetida = tmp_path / "repetida.xlsx"
    exportar_consolidado_xlsx(conn, resultado.consolidado_dre_id, salida_repetida)

    valores1 = [
        [c.value for c in row]
        for row in openpyxl.load_workbook(salida_original).active.iter_rows()
    ]
    valores2 = [
        [c.value for c in row]
        for row in openpyxl.load_workbook(salida_repetida).active.iter_rows()
    ]
    assert valores1 == valores2, (
        "re-exportar con una regla de clasificacion distinta no debe cambiar la salida ya generada"
    )


def test_recuento_de_alertas_reconciliado_entre_generacion_y_exportacion(
    conn, tmp_path
):
    """P05/DT04: el numero de alertas pendientes del resultado de generar_consolidado debe
    coincidir exacto con el que aparece en el XLSX exportado."""
    resultado = _generar_tactamal(conn)
    assert resultado.error is None
    assert (
        resultado.validaciones_pendientes >= 1
    )  # P02: reubicacion automatica de Tactamal (§9)

    salida = tmp_path / "consolidado.xlsx"
    exportar_consolidado_xlsx(conn, resultado.consolidado_dre_id, salida)

    texto_cabecera = openpyxl.load_workbook(salida).active["A7"].value
    assert (
        f"{resultado.validaciones_pendientes} alerta(s) pendiente(s)" in texto_cabecera
    )


def test_personas_omitidas_sin_identidad_expuestas(conn, tmp_path):
    """Construido: una persona sin DNI no entra al detalle (identidad no resuelta) -- P05 exige
    que quede contada, no silenciada."""
    with conn.cursor() as cur:
        cur.execute(
            "INSERT INTO local_educativo (codlocal_escale, distrito) VALUES ('900005', 'LUYA') "
            "RETURNING local_educativo_id"
        )
        local_id = cur.fetchone()["local_educativo_id"]
        cur.execute(
            "INSERT INTO institucion_educativa "
            "(local_educativo_id, cod_mod, nombre_ie, nivel_modalidad, distrito) "
            "VALUES (%s, '9999005', '9999005', 'PRIMARIA', 'LUYA')",
            (local_id,),
        )
    conn.commit()

    archivo = construir_anexo3_minimo(
        tmp_path / "anexo3_omitida.xlsx",
        institucion_raw="9999005",
        nivel_raw="PRIMARIA",
        personas=[
            {
                "dni": None,
                "nombres": "SIN DNI, PRUEBA",
                "cargo": "DOCENTE",
                "marcas": {},
            },
            {
                "dni": "30000001",
                "nombres": "CON DNI, PRUEBA",
                "cargo": "DOCENTE",
                "marcas": {1: "A"},
            },
        ],
    )
    importar_asistencia(conn, archivo)
    resultado = generar_consolidado(conn, date(2026, 7, 1), "PRIMARIA")

    assert resultado.error is None
    assert resultado.personas == 1
    assert resultado.personas_omitidas_sin_identidad == 1


def test_evaluar_revision_y_resolver_alerta(conn):
    """Bandeja de revision con fuente/autor/motivo: la reubicacion automatica de P02 (Tactamal)
    debe verse con su vinculo_trabajador_ie_confirmacion, y quedar resoluble."""
    resultado = _generar_tactamal(conn)
    assert resultado.error is None

    revision = evaluar_revision(conn, resultado.consolidado_dre_id)
    assert revision.revisable is True  # solo hay ADVERTENCIA (reubicacion), no ERROR

    with conn.cursor() as cur:
        alertas = listar_alertas_con_fuente(cur, resultado.consolidado_dre_id)
    automaticas = [
        a
        for a in alertas
        if a.codigo_regla == "INSTITUCION_ACTUALIZADA_AUTOMATICAMENTE"
    ]
    assert len(automaticas) == 1
    alerta = automaticas[0]
    assert alerta.origen_vinculo == "AUTOMATICO_IMPORTACION"
    assert alerta.motivo_vinculo is not None and "NEXUS ubica" in alerta.motivo_vinculo
    assert alerta.confirmado_por is None

    with conn.cursor() as cur:
        cur.execute(
            "INSERT INTO usuario (nombre, email) VALUES ('Prueba RRHH', 'prueba@ugel.test') "
            "RETURNING usuario_id"
        )
        usuario_id = cur.fetchone()["usuario_id"]
    conn.commit()

    resolver_alerta(
        conn,
        alerta.validacion_reporte_id,
        usuario_id,
        motivo_resolucion="Confirmado con RRHH: la reubicación corresponde.",
    )

    with conn.cursor() as cur:
        alertas_despues = listar_alertas_con_fuente(cur, resultado.consolidado_dre_id)
    alerta_despues = next(
        a
        for a in alertas_despues
        if a.validacion_reporte_id == alerta.validacion_reporte_id
    )
    assert alerta_despues.estado == "RESUELTA"
    assert alerta_despues.resuelta_por == usuario_id
    assert alerta_despues.resuelta_en is not None
    assert (
        alerta_despues.motivo_resolucion
        == "Confirmado con RRHH: la reubicación corresponde."
    )


def test_evaluar_revision_no_revisable_con_error_pendiente(conn):
    """Construido: una discrepancia ERROR pendiente vuelve la salida no revisable, aunque el
    borrador ya exista con sus vinculos automaticos en ADVERTENCIA."""
    resultado = _generar_tactamal(conn)
    assert resultado.error is None

    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT reporte_asistencia_id, trabajador_en_reporte_id FROM consolidado_dre_detalle_fuente
            WHERE consolidado_dre_detalle_id = (
                SELECT consolidado_dre_detalle_id FROM consolidado_dre_detalle
                WHERE consolidado_dre_id = %s LIMIT 1
            )
            """,
            (resultado.consolidado_dre_id,),
        )
        fuente = cur.fetchone()
        cur.execute(
            """
            INSERT INTO validacion_reporte
                (reporte_asistencia_id, trabajador_en_reporte_id, codigo_regla, severidad, mensaje)
            VALUES (%s, %s, 'ASISTENCIA_CONTRADICTORIA', 'ERROR', 'construido para prueba P05')
            """,
            (fuente["reporte_asistencia_id"], fuente["trabajador_en_reporte_id"]),
        )
    conn.commit()

    revision = evaluar_revision(conn, resultado.consolidado_dre_id)
    assert revision.revisable is False
    assert len(revision.pendientes_criticos) == 1
    assert revision.pendientes_criticos[0].codigo_regla == "ASISTENCIA_CONTRADICTORIA"

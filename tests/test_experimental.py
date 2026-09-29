"""Reportes sintéticos: procedencia, períodos, aislamiento, reproducción y web."""

import copy
import hashlib
import io
import json
import zipfile
from datetime import date

import numpy as np
import openpyxl
import psycopg
import pytest

from asistia.cierre.ingesta import padron_dre
from asistia.db import sha256_de
from asistia.experimental.datos import (
    cabecera,
    conteo,
    leer_corpus,
    mapa_periodos,
    publicar,
    redefinir,
)
from asistia.experimental.modelo import (
    ajustar,
    configurar,
    ejecutar,
    inferir,
    particiones,
)
from asistia.experimental.salidas import paquete_zip, reportes_lote

from .test_web import app, client, post  # noqa: F401 -- fixtures compartidas


def corpus_ficticio(conn, carpeta):
    carpeta.mkdir(parents=True, exist_ok=True)
    lid = conn.execute(
        "INSERT INTO local_educativo(nombre_local) VALUES('Local ficticio ML') RETURNING local_educativo_id"
    ).fetchone()["local_educativo_id"]
    ids = []
    for i in range(12):
        ids.append(
            conn.execute(
                """INSERT INTO institucion_educativa(local_educativo_id,cod_mod,nombre_ie,nivel_modalidad,distrito)
            VALUES(%s,%s,%s,'Primaria','LUYA') RETURNING institucion_educativa_id""",
                (lid, f"99991{i:02}", f"Institución ficticia {i}"),
            ).fetchone()["institucion_educativa_id"]
        )
    conn.commit()
    for m, nombre in [(3, "MARZO"), (4, "ABRIL"), (5, "MAYO"), (6, "JUNIO")]:
        w = openpyxl.Workbook()
        s = w.active
        s.title = nombre
        s["A1"] = "CONSOLIDADO PRIMARIA FICTICIO"
        s["A3"] = "MES DE OCTUBRE - 2025"
        s["A7"] = f"MES DE {nombre} - 2026"
        for col, value in {
            1: "N°",
            2: "INSTITUCIÓN EDUCATIVA",
            3: "LUGAR",
            5: "APELLIDOS Y NOMBRES",
            6: "CARGO",
            8: "FALTAS",
            12: "HORAS NO JUSTIFICADAS",
            13: "HORAS JUSTIFICADAS",
        }.items():
            s.cell(9, col, value)
        for cell, value in {
            "H10": "JUSTIFICADAS",
            "J10": "INJUSTIFICADAS",
            "H11": "N° DIAS",
            "J11": "N° DIAS",
        }.items():
            s[cell] = value
        for i in range(13):
            row = 12 + i
            s.cell(row, 1, i + 1)
            s.cell(
                row, 2, f"Institución ficticia {i}"
            )  # última institución ausente del padrón
            s.cell(row, 3, "LUYA")
            s.cell(row, 5, f"Persona Ficticia {i}")
            s.cell(row, 6, "DOCENTE")
            for c, v in {8: i % 4, 10: "-" if i % 3 else 1, 12: "-", 13: "-"}.items():
                s.cell(row, c, v)
        if m == 6:
            s["H12"] = None
            s["J12"] = None
            s["H13"] = 98  # Fuera del rango mensual; el original no se corrige.
        ruta = carpeta / f"PRIMARIA_{nombre}.xlsx"
        w.save(ruta)
        item = conn.execute(
            "INSERT INTO cierre_documento(sha256,tipo,ruta) VALUES(%s,'DRE',%s) RETURNING *",
            (sha256_de(ruta), str(ruta)),
        ).fetchone()
        conn.commit()
        padron_dre(conn, item)
    return ids


def lote_ficticio(conn, tmp_path):
    corpus_ficticio(conn, tmp_path)
    filas, manifiesto = leer_corpus(conn, tmp_path)
    lid, _ = publicar(conn, filas, manifiesto, motivo="Prueba ficticia")
    conn.commit()
    return lid


@pytest.mark.parametrize(
    "valor,estado,usado",
    [
        (None, "VACIO", None),
        ("-", "SUPUESTO_GUION_CERO", 0),
        (0, "OBSERVADO", 0),
        ("1,5", "OBSERVADO", 1.5),
        (98, "FUERA_DE_RANGO", None),
        ("ilegible", "NO_INTERPRETABLE", None),
        ("nan", "NO_INTERPRETABLE", None),
        (True, "NO_INTERPRETABLE", None),
    ],
)
def test_no_confundir_conteo_y_supuesto(valor, estado, usado):
    r = conteo(valor, "dias_justificados", date(2026, 6, 1))
    assert r["recibido"] == valor or str(r["recibido"]) == str(valor)
    assert r["estado"] == estado
    assert r["usado"] == usado
    if estado == "SUPUESTO_GUION_CERO":
        assert r["observado"] is None


def test_periodos_cruzan_anio_y_preservan_distancias():
    assert mapa_periodos(["2026-03", "2026-04", "2026-06"], "2027-11") == {
        "2026-03": "2027-11",
        "2026-04": "2027-12",
        "2026-06": "2028-02",
    }
    for invalido in ["2026-13", "2026-1", "2026-01-01", ""]:
        with pytest.raises(ValueError):
            mapa_periodos(["2026-03"], invalido)


def test_cabecera_cercana_y_contrato_no_por_posicion():
    w = openpyxl.Workbook()
    s = w.active
    s["A1"] = "MES DE OCTUBRE - 2025"
    s["A7"] = "MES DE MARZO - 2026"
    s["E9"] = "APELLIDOS Y NOMBRES"
    for celda, texto in {
        "H10": "JUSTIFICADAS",
        "J10": "INJUSTIFICADAS",
        "H11": "N° DIAS",
        "J11": "N° DIAS",
        "L11": "HORAS",
        "M9": "HORAS JUSTIFICADAS",
    }.items():
        s[celda] = texto
    cols, periodo, etiquetas = cabecera(s)
    assert cols["horas_no_justificadas"] == 12
    assert periodo == date(2026, 3, 1)
    assert etiquetas[0]["periodo"] == "2025-10"
    s["L11"] = "OTRA COSA"
    with pytest.raises(ValueError, match="encabezados"):
        cabecera(s)


def test_ingesta_real_xlsx_idempotente_y_aislada(conn, tmp_path):
    corpus_ficticio(conn, tmp_path)
    originales = {p: sha256_de(p) for p in tmp_path.glob("*.xlsx")}
    filas, manifiesto = leer_corpus(conn, tmp_path)
    assert manifiesto["filas_leidas"] == 52
    assert manifiesto["filas_aceptadas"] == 48
    assert manifiesto["exclusiones_por_motivo"] == {"INSTITUCION_NO_RESUELTA": 4}
    assert len(filas) == 48
    assert sum(r["datos"]["apto_modelo"] for r in filas) == 47
    lid, nuevo = publicar(conn, filas, manifiesto, motivo="Prueba de repetición")
    assert nuevo
    repetido, nuevo = publicar(conn, *leer_corpus(conn, tmp_path), motivo="Reintento")
    assert repetido == lid and not nuevo
    assert (
        conn.execute("SELECT count(*) n FROM ml_reporte_mensual").fetchone()["n"] == 48
    )
    assert conn.execute("SELECT count(*) n FROM asistencia_dia").fetchone()["n"] == 0
    assert (
        conn.execute("SELECT count(*) n FROM reporte_asistencia").fetchone()["n"] == 0
    )
    assert {p: sha256_de(p) for p in originales} == originales
    # Una celda no observada conserva el original y su localizador.
    assert all(
        f["conteos"]["horas_no_justificadas"]["observado"] is None
        for r in filas
        for f in r["fuentes"]
    )
    assert all(
        f["conteos"]["dias_injustificados"]["celda"].startswith("J")
        for r in filas
        for f in r["fuentes"]
    )


def test_redefinir_preserva_fuente_padre_y_particiones(conn, tmp_path):
    lid = lote_ficticio(conn, tmp_path)
    antes = reportes_lote(conn, lid)
    hijo, nuevo = redefinir(
        conn, lid, "2027-11", autor="PRUEBA", motivo="Trasladar intervalo ficticio"
    )
    assert nuevo and hijo != lid
    assert (
        redefinir(conn, lid, "2027-11", autor="PRUEBA", motivo="Reintento")[0] == hijo
    )
    despues = reportes_lote(conn, hijo)
    assert [r["periodo_fuente"] for r in antes] == [
        r["periodo_fuente"] for r in despues
    ]
    assert [r["datos"] for r in antes] == [r["datos"] for r in despues]
    assert despues[0]["periodo"] == date(2027, 11, 1)
    assert despues[-1]["periodo"] == date(2028, 2, 1)
    assert particiones(antes)[0] == particiones(despues)[0]
    assert reportes_lote(conn, lid) == antes


def test_fallo_publicacion_no_deja_lote_parcial(conn, tmp_path):
    corpus_ficticio(conn, tmp_path)
    filas, m = leer_corpus(conn, tmp_path)
    filas[-1]["institucion_educativa_id"] = 99999999
    with pytest.raises(psycopg.errors.ForeignKeyViolation):
        publicar(conn, filas, m, motivo="Provocar fallo aislado")
    assert conn.execute("SELECT count(*) n FROM ml_lote").fetchone()["n"] == 0
    assert (
        conn.execute("SELECT count(*) n FROM ml_reporte_mensual").fetchone()["n"] == 0
    )


def test_modelo_sin_fuga_reproduce_inferencia_y_conserva_zip(conn, tmp_path):
    lid = lote_ficticio(conn, tmp_path)
    eid, nuevo = ejecutar(conn, lid)
    assert nuevo and ejecutar(conn, lid) == (eid, False)
    exp = conn.execute(
        "SELECT * FROM ml_experimento WHERE ml_experimento_id=%s", (eid,)
    ).fetchone()
    filas = reportes_lote(conn, lid)
    puntuables = [r for r in filas if r["datos"]["apto_modelo"]]
    a, b = inferir(bytes(exp["modelo"]), puntuables)
    resultados = {r["reporte_id"]: r for r in exp["resultados"]}
    assert np.allclose(
        a,
        [resultados[str(r["ml_reporte_mensual_id"])]["score_if"] for r in puntuables],
        atol=1e-12,
    )
    assert np.allclose(
        b,
        [
            resultados[str(r["ml_reporte_mensual_id"])]["score_referencia"]
            for r in puntuables
        ],
        atol=1e-12,
    )
    # Modificar solo prueba final nunca cambia las medianas aprendidas ni scores train.
    alteradas = copy.deepcopy(filas)
    for r in alteradas:
        if r["periodo_fuente"].month == 6:
            r["datos"]["features"]["dias_justificados_por_registro_conocido"] = 30
    modelo_b, manifiesto_b, _resultado_b = ajustar(alteradas, configurar(alteradas))
    assert manifiesto_b["preprocesamiento"] == exp["manifiesto"]["preprocesamiento"]
    train = [r for r in filas if r["periodo_fuente"].month < 5]
    assert np.array_equal(
        inferir(bytes(exp["modelo"]), train)[0], inferir(modelo_b, train)[0]
    )
    particion = exp["manifiesto"]["particiones"]
    assert not (
        set(particion["ENTRENAMIENTO"]["ids"]) & set(particion["PRUEBA"]["ids"])
    )
    assert exp["manifiesto"]["exactitud_real"] == "NO_MEDIDO"
    archivo = paquete_zip(conn, eid)
    # Editar el nombre de una IE después no debe alterar el expediente congelado.
    conn.execute(
        "UPDATE institucion_educativa SET nombre_ie='Nombre ficticio actualizado'"
    )
    assert paquete_zip(conn, eid) == archivo
    with zipfile.ZipFile(io.BytesIO(archivo)) as z:
        hashes = json.loads(z.read("SHA256SUMS.json"))
        assert all(
            hashlib.sha256(z.read(n)).hexdigest() == h for n, h in hashes.items()
        )
        wb = openpyxl.load_workbook(
            io.BytesIO(z.read("reportes_mensuales_sinteticos.xlsx"))
        )
        assert wb.sheetnames == [
            "Leer primero",
            "2026-03",
            "2026-04",
            "2026-05",
            "2026-06",
        ]
        assert sum(s.max_row - 1 for s in list(wb)[1:]) == 48


def test_recorrido_web_periodos_descargas_y_errores(client, app, conn, tmp_path):  # noqa: F811
    carpeta = tmp_path / "dre"
    corpus_ficticio(conn, carpeta)
    app.config["EXPERIMENTAL_DRE_ROOT"] = carpeta
    respuesta = post(
        client, "/experimental/importar", {"inicio": "2026-03", "motivo": "Prueba web"}
    )
    assert respuesta.status_code == 303
    url = respuesta.headers["Location"]
    assert "Reportes mensuales sintéticos" in client.get(url).text
    assert post(client, url + "/entrenar").status_code == 303
    assert "Experimento ejecutado" in client.get(url).text
    r = conn.execute(
        "SELECT * FROM ml_reporte_mensual ORDER BY periodo LIMIT 1"
    ).fetchone()
    detalle = client.get(f"/experimental/reportes/{r['ml_reporte_mensual_id']}")
    assert detalle.status_code == 200 and "Abrir documento de origen" in detalle.text
    did = r["fuentes"][0]["cierre_documento_id"]
    assert client.get(f"/cierre/documentos/{did}").status_code == 200
    eid = conn.execute("SELECT ml_experimento_id FROM ml_experimento").fetchone()[
        "ml_experimento_id"
    ]
    assert client.get(
        f"/experimental/experimentos/{eid}/paquete.zip"
    ).data == paquete_zip(conn, eid)
    assert "SINTETICO_DERIVADO_DRE" in client.get(url + "/reportes.csv").text
    assert client.get(url + "?mes=2030-01").status_code == 400
    assert client.get(url + "?pagina=mala").status_code == 400
    assert (
        post(client, url + "/periodos", {"inicio": "2027-01", "motivo": ""}).status_code
        == 400
    )
    re = post(
        client,
        url + "/periodos",
        {"inicio": "2027-01", "motivo": "Escenario futuro sintético"},
    )
    assert re.status_code == 303
    assert "2027-01" in client.get(re.headers["Location"]).text
    assert conn.execute("SELECT count(*) n FROM ml_lote").fetchone()["n"] == 2
    assert (
        client.post(url + "/entrenar", data={"csrf": "incorrecto"}).status_code == 400
    )
    assert app.test_client().get(url).status_code == 302
    assert (
        conn.execute("SELECT count(*) n FROM reporte_asistencia").fetchone()["n"] == 0
    )


def test_correccion_tardia_no_duplica_personas_ni_oculta_conflictos(conn, tmp_path):
    anterior = lote_ficticio(conn, tmp_path)
    original = tmp_path / "PRIMARIA_MARZO.xlsx"
    correccion = tmp_path / "CORRECCION_MARZO.xlsx"
    libro = openpyxl.load_workbook(original)
    libro.active["H12"] = 3
    libro.save(correccion)
    item = conn.execute(
        "INSERT INTO cierre_documento(sha256,tipo,ruta) VALUES(%s,'DRE',%s) RETURNING *",
        (sha256_de(correccion), str(correccion)),
    ).fetchone()
    conn.commit()
    padron_dre(conn, item)
    filas, manifiesto = leer_corpus(conn, tmp_path)
    assert manifiesto["filas_leidas"] == 65
    assert manifiesto["filas_aceptadas"] == 47
    assert manifiesto["exclusiones_por_motivo"] == {
        "INSTITUCION_NO_RESUELTA": 5,
        "CONFLICTO_MISMA_IDENTIDAD_Y_ROL": 2,
        "DUPLICADO_EXACTO_MISMA_IDENTIDAD_Y_ROL": 11,
    }
    nuevo, creado = publicar(conn, filas, manifiesto, motivo="Fuente ficticia tardía")
    assert creado and nuevo != anterior
    assert len(reportes_lote(conn, anterior)) == 48
    assert len(reportes_lote(conn, nuevo)) == 47


def test_fallo_artefacto_no_publica_modelo_incompleto(conn, tmp_path, monkeypatch):
    lid = lote_ficticio(conn, tmp_path)

    def fallo(*args):
        raise OSError("Fallo simulado al construir expediente")

    with monkeypatch.context() as parche:
        parche.setattr("asistia.experimental.salidas.paquete_zip", fallo)
        with pytest.raises(OSError, match="simulado"):
            ejecutar(conn, lid)
    assert conn.execute("SELECT count(*) n FROM ml_experimento").fetchone()["n"] == 0
    eid, nuevo = ejecutar(conn, lid)
    assert nuevo
    assert paquete_zip(conn, eid)

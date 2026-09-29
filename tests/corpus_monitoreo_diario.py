"""Corpus enteramente ficticio para comprobar la web con un modelo entrenado."""

from datetime import date
from uuid import uuid4

import openpyxl

from asistia import categorias
from asistia.importar.asistencia import importar_asistencia
from asistia.importar.calendario import importar_calendario
from asistia.importar.nexus import importar_nexus

from .fixtures_excel import construir_anexo3_minimo, construir_nexus_minimo
from .test_catalogo_general import asociar, crear, operador
from .test_monitoreo_diario import perfiles_ficticios


def cargar(conn, directorio, n=100):
    directorio.mkdir(parents=True, exist_ok=True)
    perfiles = perfiles_ficticios(n)
    filas = []
    for i, p in enumerate(perfiles):
        filas.append(
            {
                "CODMOD I.E.": p["identidad"]["cod_mod"],
                "CODIGO DE PLAZA": f"FICTICIA-{i}",
                "TIPO DE TRABAJADOR": "DOCENTE",
                "SITUACION LABORAL": "NOMBRADO",
                "TIPO DE REGISTRO": "ORGANICA",
                "DOCUMENTO DE IDENTIDAD": f"{99000000 + i:08}",
                "APELLIDO PATERNO": "Ejemplo",
                "APELLIDO MATERNO": "Ficticio",
                "NOMBRES": f"Persona {i}",
                "FECHA DE INICIO": date(2026, 1, 1),
                "CODIGO LOCAL": f"{990000 + i:06}",
                "DISTRITO": "LUYA",
                "NIVEL EDUCATIVO": "Primaria",
                "NOMBRE DE LA INSTITUCION EDUCATIVA": p["identidad"]["nombre_ie"],
                "CARGO": "DOCENTE",
            }
        )
    importar_nexus(
        conn,
        construir_nexus_minimo(filas, directorio / "nexus_ficticio.xlsx"),
        date(2026, 1, 1),
    )
    existentes = [
        r
        for r in categorias.catalogo(conn, "asistencia")
        if r["nombre"] == "Asistencia"
    ]
    if existentes:
        c = existentes[0]
        cid = categorias.guardar_categoria(
            conn,
            {
                "version": str(c["version"]),
                "remuneracion": "SI",
                "es_falta": "NO",
                "motivo": "Configuración ficticia del corpus de prueba",
            },
            operador(conn),
            uuid4(),
            c["categoria_id"],
        )
    else:
        cid = crear(conn, nombre="Asistencia", pago="SI")
    asociar(conn, cid)
    conn.commit()
    for i, p in enumerate(perfiles):
        nombre = p["identidad"]["nombre_ie"]
        r = construir_anexo3_minimo(
            directorio / f"asistencia_ficticia_{i:03}.xlsx",
            institucion_raw=nombre,
            personas=[
                {
                    "dni": f"{99000000 + i:08}",
                    "nombres": f"Ejemplo Ficticio Persona {i}",
                    "cargo": "DOCENTE",
                    "marcas": {
                        j + 1: d["asistencia"]["estado_asistencia_codigo"]
                        for j, d in enumerate(p["datos"]["dias"])
                    },
                }
            ],
        )
        w = openpyxl.load_workbook(r)
        s = w.active
        for j in range(1, 32):
            s.cell(10, 4 + j, j)
        s["A14"] = "LEYENDA"
        for j, (codigo, desc) in enumerate(
            (
                ("A", "Asistencia"),
                ("F", "Feriado no remunerado"),
                ("I", "Inasistencia injustificada"),
                ("L", "Licencia con goce de remuneraciones"),
            ),
            15,
        ):
            s.cell(j, 1, codigo)
            s.cell(j, 2, desc)
        w.save(r)
        if i < n - 5:
            c = directorio / f"calendario_ficticio_{i:03}.xlsx"
            w = openpyxl.Workbook()
            s = w.active
            s["A1"] = "CALENDARIZACIÓN DEL AÑO ESCOLAR 2026"
            s["A2"], s["B2"] = "NOMBRE DE LA IE", nombre
            s["A3"], s["B3"] = "NIVEL O CICLO", "PRIMARIA"
            s["A12"] = "JULIO"
            for j, d in enumerate(p["datos"]["dias"], 1):
                s.cell(12, j + 1, j)
                s.cell(13, j + 1, "L" if d["actividad_esperada"] else "F")
            s["A15"] = "LEYENDA"
            s["A16"], s["B16"] = "L", "Día lectivo remunerado"
            s["A17"], s["B17"] = "F", "Feriado no remunerado"
            w.save(c)
            resultado = importar_calendario(conn, c)
            assert resultado.calendarizacion_version_id, resultado.error
        resultado = importar_asistencia(conn, r)
        assert resultado.reportes_asistencia_id, resultado.error
    conn.commit()

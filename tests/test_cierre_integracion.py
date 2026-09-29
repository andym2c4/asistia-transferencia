"""Casos de pérdida de bloques, identidad y procedencia del cierre."""

from datetime import date

import pytest

from asistia.cierre.ingesta import aplicar_extraccion, resolver_ie
from asistia.cierre.revision import proyectar_calendarios, revisar_instituciones
from asistia.consolidado.universo_esperado import calendario_aplicable
from asistia.db import registrar_documento
from asistia.importar.nexus import importar_nexus
from asistia.web import ErrorDeTrabajo
from asistia.web.calendarios import marcar_vigente

from .test_web import (  # noqa: F401 -- fixtures compartidas ficticias
    app,
    client,
    fuentes,
    post,
)


def base_ficticia(conn, tmp_path):
    nexus, _, _ = fuentes(tmp_path)
    importar_nexus(conn, nexus, date(2026, 6, 1))
    iid = conn.execute(
        "SELECT institucion_educativa_id FROM institucion_educativa"
    ).fetchone()["institucion_educativa_id"]
    with conn.cursor() as cur:
        did = registrar_documento(
            cur,
            nexus,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    item = conn.execute(
        "INSERT INTO cierre_documento(sha256,tipo,ruta,documento_recibido_id) VALUES(%s,'ASISTENCIA',%s,%s) RETURNING *",
        ("f" * 64, str(nexus), did),
    ).fetchone()
    conn.commit()
    return iid, item


def test_ocr_multipagina_se_conserva_en_un_reporte(conn, tmp_path):
    _iid, item = base_ficticia(conn, tmp_path)
    bloque = {
        "tipo": "asistencia",
        "institucion": "9999001 IE WEB FICTICIA",
        "cod_mod": "9999001",
        "nivel": "PRIMARIA",
        "anio": 2026,
        "mes": 7,
        "turno": "MAÑANA",
        "supuestos": [],
    }
    uno = {
        **bloque,
        "pagina": 1,
        "personas": [
            {
                "fila": 12,
                "dni": "99990001",
                "nombre": "Ejemplo Ficticio Persona Uno",
                "cargo": "DOCENTE",
                "marcas": ["A"] * 31,
            }
        ],
    }
    dos = {
        **bloque,
        "pagina": 2,
        "personas": [
            {
                "fila": 12,
                "dni": "99990002",
                "nombre": "Ejemplo Ficticio Persona Dos",
                "cargo": "DOCENTE",
                "marcas": ["I"] * 31,
            }
        ],
    }
    r = aplicar_extraccion(
        conn,
        item,
        {
            "datos": {"bloques": [uno, dos]},
            "modelo": "FICTICIO",
            "version": "prueba",
            "sha256": "f" * 64,
        },
    )
    assert len(r["ids"]) == 1
    assert (
        conn.execute("SELECT count(*) AS n FROM trabajador_en_reporte").fetchone()["n"]
        == 2
    )
    assert (
        conn.execute("SELECT count(*) AS n FROM asistencia_dia").fetchone()["n"] == 62
    )
    assert (
        conn.execute(
            "SELECT count(*) AS n FROM asistencia_dia WHERE celda_origen LIKE 'pagina:2;fila:12;%%'"
        ).fetchone()["n"]
        == 31
    )
    assert resolver_ie(conn, "NO COINCIDE", "SECUNDARIA", "9999001")[0] is None


def test_calendario_proyectado_no_sustituye_2026_y_no_se_aprueba(
    conn, tmp_path, monkeypatch
):
    from asistia.cierre import revision

    monkeypatch.setattr(revision, "SALIDA", tmp_path)
    iid, _item = base_ficticia(conn, tmp_path)
    cl = conn.execute(
        "INSERT INTO calendarizacion_local(institucion_educativa_id,anio) VALUES(%s,2025) RETURNING calendarizacion_local_id",
        (iid,),
    ).fetchone()["calendarizacion_local_id"]
    cv = conn.execute(
        "INSERT INTO calendarizacion_version(calendarizacion_local_id,version,origen,motivo_version) VALUES(%s,1,'IMPORTACION_IE','Prueba ficticia') RETURNING calendarizacion_version_id",
        (cl,),
    ).fetchone()["calendarizacion_version_id"]
    conn.execute(
        "INSERT INTO dia_calendarizacion(calendarizacion_version_id,fecha,estado_captura) VALUES(%s,'2025-07-01','VACIO')",
        (cv,),
    )
    primero = proyectar_calendarios(conn)
    assert len(primero) == 1
    cvid = primero[0]["calendarizacion_version_id"]
    with pytest.raises(ErrorDeTrabajo, match="basado en 2025"):
        marcar_vigente(conn, cvid, 1)
    assert proyectar_calendarios(conn) == []
    from asistia.cierre.ajustes import ajustar_calendario

    ajustada = ajustar_calendario(conn, cvid, {}, "Cotejo ficticio de la proyección.")
    fuente = conn.execute(
        "SELECT fuente_derivada_id FROM calendarizacion_version WHERE calendarizacion_version_id=%s",
        (ajustada,),
    ).fetchone()
    assert fuente["fuente_derivada_id"] == cv
    with pytest.raises(ErrorDeTrabajo, match="basado en 2025"):
        marcar_vigente(conn, ajustada, 1)
    local = conn.execute(
        "SELECT calendarizacion_local_id FROM calendarizacion_local WHERE institucion_educativa_id=%s AND anio=2026",
        (iid,),
    ).fetchone()["calendarizacion_local_id"]
    nuevo = conn.execute(
        "INSERT INTO calendarizacion_version(calendarizacion_local_id,version,origen,motivo_version) VALUES(%s,3,'IMPORTACION_IE','Fuente 2026') RETURNING calendarizacion_version_id",
        (local,),
    ).fetchone()["calendarizacion_version_id"]
    with conn.cursor() as cur:
        assert calendario_aplicable(cur, iid, 2026)[0] == nuevo
    assert revisar_instituciones(conn)["instituciones"] == 1


def test_revision_institucion_sin_fuentes_conserva_pendiente(
    conn, tmp_path, monkeypatch
):
    from asistia.cierre import revision

    monkeypatch.setattr(revision, "SALIDA", tmp_path)
    base_ficticia(conn, tmp_path)
    r = revisar_instituciones(conn)
    assert r["instituciones"] == 1
    assert r["por_resultado"] == {"REQUIERE_DATOS": 1}
    revisar_instituciones(conn)
    assert (
        conn.execute(
            "SELECT count(*) AS n FROM cierre_revision_institucion"
        ).fetchone()["n"]
        == 1
    )
    assert (
        conn.execute("SELECT count(*) AS n FROM reporte_asistencia").fetchone()["n"]
        == 0
    )


def test_cierre_web_estados_y_documento(client, conn, tmp_path):  # noqa: F811 -- fixture compartida
    _, item = base_ficticia(conn, tmp_path)
    assert client.get("/cierre?periodo=2026-07&nivel=PRIMARIA").status_code == 200
    r = client.get(f"/cierre/documentos/{item['cierre_documento_id']}")
    assert r.status_code == 200 and "Reintentar extracción" in r.text
    assert (
        client.get(
            "/cierre/instituciones.csv?periodo=2026-07&nivel=PRIMARIA"
        ).status_code
        == 200
    )
    assert (
        client.post(
            f"/cierre/documentos/{item['cierre_documento_id']}/reintentar"
        ).status_code
        == 400
    )


def test_salida_tecnica_expone_conflicto_sin_sumar_dos_veces(
    conn,
    tmp_path,
    app,  # noqa: F811 -- fixture compartida
    monkeypatch,
):
    import uuid

    import openpyxl

    from asistia.cierre import revision
    from asistia.cierre.salidas import alcance_nivel
    from asistia.web.salidas import preparar

    monkeypatch.setattr(revision, "SALIDA", tmp_path)
    _, item = base_ficticia(conn, tmp_path)
    b = {
        "tipo": "asistencia",
        "institucion": "9999001 IE WEB FICTICIA",
        "cod_mod": "9999001",
        "nivel": "PRIMARIA",
        "anio": 2026,
        "mes": 7,
        "pagina": 1,
        "supuestos": [],
    }
    p = {
        "fila": 12,
        "dni": "99990001",
        "nombre": "Ejemplo Ficticio Persona Uno",
        "cargo": "DOCENTE",
    }
    bloques = [
        {**b, "turno": "MAÑANA", "personas": [{**p, "marcas": ["A"] * 31}]},
        {**b, "turno": "TARDE", "personas": [{**p, "marcas": ["I"] + [None] * 30}]},
    ]
    aplicar_extraccion(
        conn,
        item,
        {
            "datos": {"bloques": bloques},
            "modelo": "FICTICIO",
            "version": "prueba",
            "sha256": "f" * 64,
        },
    )
    revisar_instituciones(conn)
    alcance = alcance_nivel(conn, date(2026, 7, 1), "PRIMARIA")
    assert len(alcance["filas_elegidas"]) == 1
    assert alcance["decisiones"][0]["fechas_conflicto"] == ["2026-07-01"]
    uid = conn.execute("SELECT usuario_id FROM usuario LIMIT 1").fetchone()[
        "usuario_id"
    ]
    sid = uuid.uuid4()
    with app.app_context():
        preparar(
            conn, date(2026, 7, 1), "PRIMARIA", False, sid, uid, cierre_tecnico=alcance
        )
    salida = conn.execute(
        "SELECT w.*,o.ruta_objeto FROM web_salida w JOIN objeto_archivo o USING(objeto_archivo_id) WHERE web_salida_id=%s",
        (sid,),
    ).fetchone()
    assert salida["manifiesto"]["personas_distintas"] == 1
    assert salida["manifiesto"]["filas_detalle"] == 1
    assert (
        salida["manifiesto"]["filas"][0]["fuente_calculo"]["clasificacion_faltas"][
            "injustificadas"
        ]
        == []
    )
    wb = openpyxl.load_workbook(salida["ruta_objeto"])
    assert "Selección de fuentes" in wb.sheetnames
    assert wb["Selección de fuentes"].cell(2, 5).value == "2026-07-01"
    assert "Revisión institucional" in wb.sheetnames


def test_carga_pdf_continua_con_ocr_desde_web(client, conn, tmp_path, monkeypatch):  # noqa: F811
    from asistia.cierre.gemini import Gemini
    from asistia.db import sha256_de

    from .test_web import cargar

    base_ficticia(conn, tmp_path)
    pdf = tmp_path / "asistencia.pdf"
    from pypdf import PdfWriter

    escritor = PdfWriter()
    escritor.add_blank_page(width=100, height=100)
    escritor.write(pdf)
    cargar(client, pdf, "asistencia")
    carga = conn.execute("SELECT * FROM web_carga WHERE tipo='asistencia'").fetchone()
    assert carga["estado"] == "PARCIAL"
    did = carga["resultado"]["cierre_documento_id"]
    assert (
        "Continuar con lectura asistida"
        in client.get(f"/documentos/{carga['web_carga_id']}").text
    )
    bloque = {
        "tipo": "asistencia",
        "cod_mod": "9999001",
        "institucion": "IE WEB FICTICIA",
        "nivel": "PRIMARIA",
        "anio": 2026,
        "mes": 7,
        "pagina": 1,
        "personas": [
            {
                "dni": "99990001",
                "nombre": "Persona Ficticia Uno",
                "fila": 12,
                "cargo": "DOCENTE",
                "marcas": ["A"] * 31,
            }
        ],
    }
    monkeypatch.setattr(
        Gemini,
        "extraer",
        lambda *a: {
            "datos": {"bloques": [bloque]},
            "modelo": "FICTICIO",
            "version": "test",
            "sha256": sha256_de(pdf),
        },
    )
    assert post(client, f"/cierre/documentos/{did}/reintentar").status_code == 303
    assert (
        conn.execute(
            "SELECT estado FROM web_carga WHERE web_carga_id=%s",
            (carga["web_carga_id"],),
        ).fetchone()["estado"]
        == "PROCESADO"
    )
    assert conn.execute("SELECT count(*) n FROM asistencia_dia").fetchone()["n"] == 31
    # Nueva lectura explícita de un parcial: conserva la decisión y no duplica hechos.
    from psycopg.types.json import Jsonb

    conn.execute(
        "UPDATE cierre_documento SET estado=%s,resultado=%s WHERE cierre_documento_id=%s",
        (
            "PARCIAL",
            Jsonb(
                {
                    "pendientes": [
                        {"bloque_original": 2, "mes": 6, "motivo": "Mes incompleto"}
                    ]
                }
            ),
            did,
        ),
    )
    conn.commit()
    pantalla = client.get(f"/cierre/documentos/{did}").text
    assert "Volver a leer pendientes" in pantalla
    assert "Bloque 2 · mes 6" in pantalla
    assert (
        post(
            client, f"/cierre/documentos/{did}/reintentar", {"nueva_lectura": "1"}
        ).status_code
        == 303
    )
    contexto = conn.execute(
        "SELECT contexto_tecnico FROM cierre_documento WHERE cierre_documento_id=%s",
        (did,),
    ).fetchone()["contexto_tecnico"]
    assert contexto["relectura_ocr"]
    assert contexto["historial"][-1]["accion"] == "RELECTURA_OCR"
    assert conn.execute("SELECT count(*) n FROM asistencia_dia").fetchone()["n"] == 31


def test_condicion_documental_conservada_y_rol_cas(conn, tmp_path):
    _, item = base_ficticia(conn, tmp_path)
    extraido = {
        "datos": {
            "bloques": [
                {
                    "tipo": "asistencia",
                    "institucion": "9999001 IE WEB FICTICIA",
                    "cod_mod": "9999001",
                    "nivel": "PRIMARIA",
                    "anio": 2026,
                    "mes": 7,
                    "pagina": 1,
                    "personas": [
                        {
                            "fila": 3,
                            "dni": "99990991",
                            "nombre": "Persona Ficticia CAS",
                            "cargo": "PSICOLOGO",
                            "condicion": "CAS",
                            "marcas": ["A"] * 31,
                        }
                    ],
                }
            ]
        },
        "modelo": "FICTICIO",
        "version": "prueba",
        "sha256": item["sha256"],
    }
    r = aplicar_extraccion(conn, item, extraido)
    assert aplicar_extraccion(conn, item, extraido) == r
    fila = conn.execute(
        """SELECT t.condicion_reportada_raw, c.codigo
        FROM trabajador_en_reporte t JOIN catalogo_rol_laboral c USING(rol_laboral_id)
        WHERE reporte_asistencia_id=%s""",
        (r["ids"][0],),
    ).fetchone()
    assert fila == {"condicion_reportada_raw": "CAS", "codigo": "CAS"}
    assert conn.execute("SELECT count(*) n FROM asistencia_dia").fetchone()["n"] == 31

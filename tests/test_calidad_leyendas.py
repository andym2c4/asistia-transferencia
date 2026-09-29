"""Desfase ficticio, abstención, revisión e independencia de las letras locales."""
# ruff: noqa: F811

from datetime import date
from uuid import uuid4

import openpyxl
import pytest

from asistia.calidad_leyendas import (
    REGLA_ALERTA,
    cotejar_excel,
    detectar_desfases,
    pendiente,
)
from asistia.db import sha256_de
from asistia.importar.asistencia import importar_asistencia
from asistia.importar.nexus import importar_nexus
from asistia.leyendas import (
    aplicar_leyenda_asistencia,
    leer_leyenda_xlsx,
    normalizar_leyenda,
)
from asistia.monitoreo.datos import construir
from asistia.web import ErrorDeTrabajo
from asistia.web.calidad_leyendas import comprobar_reporte, listar
from asistia.web.lecturas import huella_reporte
from asistia.web.leyendas import guardar_asistencia
from asistia.web.revision import confirmar_reporte, corregir_marca, decidir_alerta

from .test_catalogo_general import asociar, crear, operador
from .test_monitoreo_diario import fuente_minima
from .test_web import app, client, fuentes, post  # noqa: F401

CODIGOS = ["A", "I", "3T", "J", "L", "P", "T", "H"]
TEXTOS = [
    "Día laborado",
    "Inasistencia injustificada",
    "Tercera tardanza considerada como inasistencia injustificada",
    "Inasistencia justificada",
    "Licencia con goce de remuneraciones",
    "Permiso sin goce de remuneraciones",
    "Tardanza",
    "Huelga o paro",
]


def poner_leyenda(ws, desplazamiento=-1, codigos=CODIGOS, textos=TEXTOS):
    ws["A24"] = "LEYENDA:"
    for fila, codigo, texto in zip(range(25, 33), codigos, textos):
        ws.cell(fila, 1, codigo)
        ws.cell(fila + desplazamiento, 2, texto)


@pytest.mark.parametrize("desplazamiento", [-1, 1])
def test_detecta_extremos_y_conserva_pares_sin_corregir(desplazamiento):
    ws = openpyxl.Workbook().active
    poner_leyenda(ws, desplazamiento)
    h = detectar_desfases(ws)
    assert len(h) == 1 and h[0]["desplazamiento_descripcion"] == desplazamiento
    c = normalizar_leyenda(leer_leyenda_xlsx(ws), "asistencia")
    assert set(c) == set(CODIGOS)
    assert all(
        pendiente(v) and v["es_falta"] is None and v["estado_asistencia_codigo"] is None
        for v in c.values()
    )
    assert c["A"]["evidencia"]["revision_leyenda"]["pares"][0]["celda_codigo"] == "A25"
    if desplazamiento == -1:
        assert c["A"]["tipo_dia"] == "Inasistencia injustificada"
        assert h[0]["pares"][0]["descripcion_desplazada"] == "Día laborado"


@pytest.mark.parametrize("codigos", [CODIGOS, ["X", "Y", "Z", "Q", "W", "R", "V", "K"]])
def test_alineacion_correcta_no_depende_de_letras(codigos):
    ws = openpyxl.Workbook().active
    poner_leyenda(ws, 0, codigos)
    assert detectar_desfases(ws) == []
    c = normalizar_leyenda(leer_leyenda_xlsx(ws), "asistencia")
    assert c[codigos[0]]["estado_asistencia_codigo"] == "A"


def test_catalogo_reconoce_textos_sin_imponer_codigos():
    ws = openpyxl.Workbook().active
    textos = [f"Situación especial {n}" for n in range(8)]
    poner_leyenda(ws, textos=textos)
    assert detectar_desfases(ws) == []
    assert len(detectar_desfases(ws, textos)) == 1


def test_no_declara_desfase_con_geometria_ambigua_o_sin_extremos():
    ws = openpyxl.Workbook().active
    poner_leyenda(ws)
    ws["B32"] = "Otra descripción"
    assert detectar_desfases(ws) == []
    ws["B32"] = None
    ws.merge_cells("B24:B25")
    assert detectar_desfases(ws) == []


def test_pagina_numerica_no_es_indice_de_hoja(tmp_path):
    w = openpyxl.Workbook()
    w.active.title = "PORTADA"
    s = w.create_sheet("ANEXO 3")
    poner_leyenda(s)
    ruta = tmp_path / "ficticio.xlsx"
    w.save(ruta)
    assert pendiente(cotejar_excel(ruta, "1", {})["A"])
    w.create_sheet("ANEXO-3")
    w.save(ruta)
    with pytest.raises(ValueError, match="forma única"):
        cotejar_excel(ruta, "1", {})


def importar_desfase(conn, tmp_path):
    nx, _, ruta = fuentes(tmp_path)
    libro = openpyxl.load_workbook(ruta)
    poner_leyenda(libro.active)
    libro.save(ruta)
    libro.close()
    importar_nexus(conn, nx, date(2026, 1, 1))
    r = importar_asistencia(conn, ruta)
    assert not r.error
    return next(iter(r.reportes_asistencia_id.values())), ruta


def test_importacion_suspende_hechos_y_catalogo_no_rehabilita(conn, tmp_path):
    cid = crear(conn, nombre="Inasistencia injustificada", pago="NO")
    asociar(conn, cid, "Inasistencia injustificada")
    rid, _ = importar_desfase(conn, tmp_path)
    c = conn.execute(
        "SELECT clasificacion_codigos FROM reporte_asistencia WHERE reporte_asistencia_id=%s",
        (rid,),
    ).fetchone()["clasificacion_codigos"]
    assert pendiente(c["A"]) and c["A"]["es_remunerado"] is None
    assert (
        conn.execute(
            "SELECT count(*) n FROM asistencia_dia WHERE estado_asistencia_id IS NOT NULL"
        ).fetchone()["n"]
        == 0
    )
    assert (
        conn.execute("SELECT count(*) n FROM hecho_asistencia_dia").fetchone()["n"] == 0
    )
    assert (
        conn.execute(
            "SELECT count(*) n FROM validacion_reporte WHERE codigo_regla=%s",
            (REGLA_ALERTA,),
        ).fetchone()["n"]
        == 1
    )
    aplicar_leyenda_asistencia(conn, rid, c)
    aplicar_leyenda_asistencia(conn, rid, c)
    assert (
        conn.execute(
            "SELECT count(*) n FROM validacion_reporte WHERE codigo_regla=%s",
            (REGLA_ALERTA,),
        ).fetchone()["n"]
        == 1
    )
    with pytest.raises(ErrorDeTrabajo, match="desalineación"):
        confirmar_reporte(
            conn,
            rid,
            "Revisión ficticia",
            huella_reporte(conn, rid),
            uuid4(),
            operador(conn),
        )


def test_monitoreo_abstiene_y_respeta_correccion_diaria():
    ws = openpyxl.Workbook().active
    poner_leyenda(ws)
    c = normalizar_leyenda(leer_leyenda_xlsx(ws), "asistencia")["A"]
    f = fuente_minima()
    f["reportes"][0]["clasificacion_codigos"]["X"] = c
    p = construir(f)[0]["datos"]
    assert p["componentes"]["fraccion_inasistencia"]["denominador"] == 0
    assert p["dias"][0]["declaracion"] is None
    assert p["dias"][0]["pendiente"] == "LEYENDA_POR_REVISAR"
    assert not p["apto_modelo"]
    f["dias"][0]["evidencia_interpretacion"] = {
        "clasificacion_aceptada": {
            "estado_asistencia_codigo": "A",
            "fuente": "CORRECCION_WEB",
        }
    }
    assert construir(f)[0]["datos"]["dias"][0]["declaracion"] == "PRESENCIA"


def test_extraccion_alternativa_coteja_excel_antes_de_materializar(conn, tmp_path):
    from asistia.cierre.ingesta import aplicar_extraccion
    from asistia.db import registrar_documento

    from .test_leyendas_remuneracion import (
        bloque_asistencia,
        extraida,
        persona_ficticia,
    )

    nx, _, ruta = fuentes(tmp_path)
    w = openpyxl.load_workbook(ruta)
    poner_leyenda(w.active)
    w.save(ruta)
    w.close()
    importar_nexus(conn, nx, date(2026, 1, 1))
    with conn.cursor() as cur:
        did = registrar_documento(
            cur,
            ruta,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    item = conn.execute(
        "INSERT INTO cierre_documento(sha256,tipo,ruta,documento_recibido_id) VALUES(%s,'ASISTENCIA',%s,%s) RETURNING *",
        (sha256_de(ruta), str(ruta), did),
    ).fetchone()
    b = bloque_asistencia(
        [persona_ficticia(1, ["A"] * 31)],
        [{"codigo": "A", "descripcion": "Inasistencia injustificada"}],
    )
    result = aplicar_extraccion(conn, item, extraida(b))
    c = conn.execute(
        "SELECT clasificacion_codigos FROM reporte_asistencia WHERE reporte_asistencia_id=%s",
        (result["ids"][0],),
    ).fetchone()["clasificacion_codigos"]
    assert pendiente(c["A"])
    assert (
        conn.execute("SELECT count(*) n FROM hecho_asistencia_dia").fetchone()["n"] == 0
    )


def test_revision_por_codigo_cierra_alerta_solo_al_terminar(conn, tmp_path):
    rid, _ = importar_desfase(conn, tmp_path)
    uid = operador(conn)
    a = conn.execute(
        "SELECT * FROM validacion_reporte WHERE codigo_regla=%s", (REGLA_ALERTA,)
    ).fetchone()
    with pytest.raises(ErrorDeTrabajo, match="revisando los códigos"):
        decidir_alerta(
            conn,
            rid,
            a["validacion_reporte_id"],
            "DESCARTAR_ALERTA",
            "Ficticio",
            huella_reporte(conn, rid),
            uuid4(),
            uid,
        )
    did = conn.execute(
        "SELECT asistencia_dia_id FROM asistencia_dia LIMIT 1"
    ).fetchone()["asistencia_dia_id"]
    with pytest.raises(ErrorDeTrabajo, match="Primero revisa"):
        corregir_marca(
            conn,
            rid,
            did,
            "A",
            "REGISTRADO",
            "Ficticio",
            huella_reporte(conn, rid),
            uuid4(),
            uid,
        )
    for i, (codigo, texto) in enumerate(zip(CODIGOS, TEXTOS)):
        guardar_asistencia(
            conn,
            rid,
            {"codigo": codigo, "tipo_dia": texto, "remuneracion": "", "es_falta": ""},
            "Cotejo ficticio de celdas desplazadas",
            huella_reporte(conn, rid),
            uuid4(),
            uid,
        )
        estado = conn.execute(
            "SELECT estado FROM validacion_reporte WHERE validacion_reporte_id=%s",
            (a["validacion_reporte_id"],),
        ).fetchone()["estado"]
        assert estado == ("RESUELTA" if i == 7 else "PENDIENTE")
    assert listar(conn) == []


def test_comprobar_anterior_idempotente_y_preserva_revision_y_original(
    app, conn, tmp_path
):
    rid, ruta = importar_desfase(conn, tmp_path)
    # Representa una extracción antigua: no modifica el Excel ni las marcas recibidas.
    w = openpyxl.load_workbook(ruta)
    antiguas = normalizar_leyenda(
        leer_leyenda_xlsx(w.active, control_alineacion=False), "asistencia"
    )
    w.close()
    antiguas["A"].update(
        fuente="REVISION_WEB",
        estado_asistencia_codigo="A",
        motivo="Cotejo previo ficticio",
    )
    aplicar_leyenda_asistencia(conn, rid, antiguas, usar_catalogo=False)
    original = sha256_de(ruta)
    with app.app_context():
        assert comprobar_reporte(conn, rid, operador(conn))["estado"] == "DETECTADO"
        assert comprobar_reporte(conn, rid, operador(conn))["estado"] == "SIN_CAMBIOS"
    actual = conn.execute(
        "SELECT clasificacion_codigos FROM reporte_asistencia WHERE reporte_asistencia_id=%s",
        (rid,),
    ).fetchone()["clasificacion_codigos"]
    assert actual["A"] == antiguas["A"]
    assert actual["I"]["clasificacion_previa"] == antiguas["I"]
    assert sha256_de(ruta) == original
    assert (
        conn.execute(
            "SELECT count(*) n FROM web_revision_evento WHERE accion='COMPROBAR_LEYENDA'"
        ).fetchone()["n"]
        == 1
    )


def test_web_bandeja_filtros_y_evidencia(client, conn, tmp_path):
    assert "No hay reportes" in client.get("/monitoreo/leyendas").text
    rid, _ = importar_desfase(conn, tmp_path)
    assert "Posible desfase de leyenda" in client.get("/monitoreo/leyendas").text
    assert "No hay reportes" in client.get("/monitoreo/leyendas?mes=2026-08").text
    html = client.get(f"/reportes/{rid}").text
    assert "A24:B32" in html and "fila vecina, por comprobar" in html
    assert post(client, "/monitoreo/leyendas/comprobar").status_code == 303
    assert client.post("/monitoreo/leyendas/comprobar").status_code == 400

"""Paridad entre vacíos cotejados, cruce diario y faltas de la nueva salida."""

# ruff: noqa: F811
import io

import openpyxl
import pytest

from asistia.consolidado.remuneracion import cruce_persona

from .test_coherencia_reporte import cruce  # noqa: F401
from .test_reporte_vacios import payload
from .test_web import app, client, generar, post  # noqa: F401


@pytest.mark.parametrize(
    "numero,confirmar,calendario,esperado,falta",
    [
        (4, True, "D", "NO_APLICA", False),
        (2, True, "L", "OBSERVACION", None),
        (4, False, "D", "PENDIENTE", None),
        (4, True, "?", "PENDIENTE", None),
    ],
)
def test_no_aplica_requiere_cotejo_y_calendario_sin_actividad(
    client, conn, cruce, numero, confirmar, calendario, esperado, falta
):
    rid, cv = cruce
    fecha = f"2026-07-{numero:02}"
    d = conn.execute(
        "UPDATE asistencia_dia SET codigo_reportado_raw=NULL,codigo_interpretado=NULL,"
        "estado_captura='VACIO',estado_asistencia_id=NULL,hecho_asistencia_dia_id=NULL,"
        "validado=false,evidencia_interpretacion='{}' WHERE fecha=%s RETURNING *",
        (fecha,),
    ).fetchone()
    conn.execute(
        "UPDATE dia_calendarizacion SET codigo_reportado_raw=%s WHERE calendarizacion_version_id=%s AND fecha=%s",
        (calendario, cv, fecha),
    )
    conn.commit()
    if confirmar:
        assert (
            post(
                client,
                f"/reportes/{rid}/vacios",
                payload(conn, rid, [d["asistencia_dia_id"]]),
            ).status_code
            == 303
        )
    result = next(
        x
        for x in cruce_persona(conn, d["trabajador_en_reporte_id"])["dias"]
        if x["fecha"] == fecha
    )
    assert result["resultado"] == esperado
    assert result["es_falta"] is falta and result["es_remunerado"] is None
    after = conn.execute(
        "SELECT codigo_reportado_raw FROM asistencia_dia WHERE asistencia_dia_id=%s",
        (d["asistencia_dia_id"],),
    ).fetchone()
    assert after["codigo_reportado_raw"] is None


def test_vacio_resuelto_cero_faltas_en_excel_y_salida_anterior_conservada(
    client, conn, cruce
):
    rid, _ = cruce
    conn.execute(
        "UPDATE asistencia_dia SET codigo_reportado_raw='A',codigo_interpretado=NULL,estado_captura='REGISTRADO',estado_asistencia_id=(SELECT estado_asistencia_id FROM catalogo_estado_asistencia WHERE codigo='A'),evidencia_interpretacion='{}' WHERE extract(day FROM fecha)>=29"
    )
    d = conn.execute(
        "UPDATE asistencia_dia SET codigo_reportado_raw=NULL,codigo_interpretado=NULL,estado_captura='VACIO',estado_asistencia_id=NULL,hecho_asistencia_dia_id=NULL,validado=false,evidencia_interpretacion='{}' WHERE fecha='2026-07-04' RETURNING *"
    ).fetchone()
    conn.commit()
    antes = generar(client, "BORRADOR")
    assert antes.status_code == 303
    anterior = client.get(antes.location + "/descargar").data
    assert (
        openpyxl.load_workbook(io.BytesIO(anterior)).active["H12"].value == "Pendiente"
    )
    assert (
        post(
            client,
            f"/reportes/{rid}/vacios",
            payload(conn, rid, [d["asistencia_dia_id"]]),
        ).status_code
        == 303
    )
    nuevo = generar(client, "BORRADOR")
    assert nuevo.status_code == 303
    wb = openpyxl.load_workbook(
        io.BytesIO(client.get(nuevo.location + "/descargar").data)
    )
    assert len(wb.sheetnames) == 1
    assert wb.active["H12"].value == wb.active["J12"].value == 0
    assert (
        "AsistIA" in wb.active["A7"].value
        and "Control calendario" not in wb.active["A7"].value
    )
    assert client.get(antes.location + "/descargar").data == anterior
    frozen = conn.execute(
        "SELECT fuente_calculo FROM consolidado_dre_detalle ORDER BY consolidado_dre_id"
    ).fetchall()
    assert {x["fuente_calculo"]["clasificacion_faltas"]["estado"] for x in frozen} == {
        "NO_CALCULABLE",
        "CALCULABLE",
    }

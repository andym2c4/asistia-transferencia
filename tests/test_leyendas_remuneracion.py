"""Casos ficticios: semántica local, pago, conservación y tablas partidas."""

import uuid
from datetime import date

import openpyxl
import pytest

from asistia.cierre.ingesta import aplicar_extraccion
from asistia.cierre.tablas import reconstruir_columnas
from asistia.consolidado.generar import generar_consolidado
from asistia.consolidado.remuneracion import cruce_persona, cruzar_dia
from asistia.leyendas import leer_leyenda_xlsx, normalizar_leyenda
from asistia.web import ErrorDeTrabajo
from asistia.web.auth import crear_usuario
from asistia.web.lecturas import huella_reporte
from asistia.web.leyendas import (
    guardar_asistencia,
    guardar_calendario,
    huella_calendario,
)

from .test_cierre_integracion import base_ficticia
from .test_web import app, client, post  # noqa: F401


@pytest.mark.parametrize(
    "cal,asis,falta,resultado,es_falta",
    [
        (True, True, False, "REMUNERADO", False),
        (True, False, True, "NO_REMUNERADO", True),
        (True, False, False, "NO_REMUNERADO", False),
        (False, True, False, "OBSERVACION", False),
        (False, False, True, "NO_REMUNERADO", False),
        (None, True, False, "PENDIENTE", None),
        (True, None, None, "PENDIENTE", None),
    ],
)
def test_cruce_de_categorias(cal, asis, falta, resultado, es_falta):
    r = cruzar_dia(cal, asis, falta=falta)
    assert r["resultado"] == resultado
    assert r["es_falta"] is es_falta
    assert (
        cruzar_dia(cal, asis, impedimento="Fuente insuficiente")["es_remunerado"]
        is None
    )


def test_leyenda_no_depende_de_la_letra_y_no_cree_pago_sugerido_por_ocr():
    con = normalizar_leyenda(
        [{"codigo": "X", "descripcion": "Licencia con goce de remuneraciones"}],
        "asistencia",
    )
    sin = normalizar_leyenda(
        [{"codigo": "X", "descripcion": "Licencia sin goce de remuneraciones"}],
        "asistencia",
    )
    assert con["X"]["es_remunerado"] is True
    assert sin["X"]["es_remunerado"] is False
    assert sin["X"]["estado_asistencia_codigo"] == "LSG"
    assert (
        normalizar_leyenda(
            [{"codigo": "F", "descripcion": "Feriado", "es_remunerado": True}],
            "calendario",
        )["F"]["es_remunerado"]
        is None
    )
    conflicto = normalizar_leyenda(
        [
            {"codigo": "X", "descripcion": "Lectivo"},
            {"codigo": "X", "descripcion": "Descanso"},
        ],
        "calendario",
    )
    assert len(conflicto["X"]["conflictos"]) == 2
    assert conflicto["X"]["es_remunerado"] is None


def test_leyenda_excel_preserva_celdas():
    ws = openpyxl.Workbook().active
    ws.title = "ANEXO 3"
    ws.append(["LEYENDA"])
    ws.append(["F", "Feriado"])
    ws.append(["X = Licencia sin goce de remuneraciones"])
    filas = leer_leyenda_xlsx(ws)
    assert [(f["codigo"], f["celda"]) for f in filas] == [("F", "A2:B2"), ("X", "A3")]


def test_revision_institucional_expone_categorias_sin_regla(conn, tmp_path):
    from psycopg.types.json import Jsonb

    from asistia.cierre.revision import revisar_instituciones

    preparar_cruce(conn, tmp_path)
    conn.execute(
        "UPDATE calendarizacion_version SET clasificacion_codigos=%s",
        (
            Jsonb(
                normalizar_leyenda(
                    [
                        {"codigo": "U1", "descripcion": "Unidad lectiva remunerada"},
                        {"codigo": "D", "descripcion": "Sábados, domingos y feriados"},
                    ],
                    "calendario",
                )
            ),
        ),
    )
    revisar_instituciones(conn, evidencia_dir=tmp_path)
    r = conn.execute(
        "SELECT resultado,detalle FROM cierre_revision_institucion"
    ).fetchone()
    assert r["resultado"] == "REQUIERE_DATOS"
    assert r["detalle"]["remuneracion_pendiente"] == {
        "calendario": ["D"],
        "asistencia": [],
    }
    assert any("remuneración" in p and "D" in p for p in r["detalle"]["problemas"])


def test_leyenda_sin_simbolo_conserva_descripcion_y_no_clasifica_vacios(conn, tmp_path):
    _, item = base_ficticia(conn, tmp_path)
    cabecera = bloque_asistencia([])
    b = {
        **cabecera,
        "tipo": "calendario",
        "meses": [{"mes": 7, "codigos": ["L"] * 30 + [None]}],
        "leyenda": [
            {"codigo": "L", "descripcion": "Lectivo"},
            {"codigo": "", "descripcion": "Sábados, domingos y feriados", "pagina": 1},
        ],
    }
    ext = extraida(b)
    rid = aplicar_extraccion(conn, {**item, "tipo": "CALENDARIO_2026"}, ext)["ids"][0]
    r = conn.execute(
        "SELECT * FROM calendarizacion_version WHERE calendarizacion_version_id=%s",
        (rid,),
    ).fetchone()
    assert set(r["clasificacion_codigos"]) == {"L"}
    assert r["procedencia_extraccion"]["leyenda_sin_codigo"] == [b["leyenda"][1]]
    assert ext["datos"]["bloques"][0]["leyenda"] == b["leyenda"]
    d = conn.execute(
        "SELECT * FROM dia_calendarizacion WHERE calendarizacion_version_id=%s AND fecha='2026-07-31'",
        (rid,),
    ).fetchone()
    assert d["codigo_reportado_raw"] is None
    assert d["codigo_interpretado"] is None
    assert d["estado_captura"] == "VACIO"


def fragmentada():
    b = {
        "tipo": "asistencia",
        "institucion": "FICTICIA",
        "anio": 2026,
        "mes": 7,
        "personas": [],
    }
    base = {"id_fila": "001", "evidencia_ancla": "Correlativo 001 impreso", "fila": 12}
    b["fragmentos_columnas"] = [
        {
            "pagina": 2,
            "dias": [],
            "personas": [
                {**base, "nombre": "Persona Ficticia", "dni": "99990001", "marcas": []}
            ],
        },
        {
            "pagina": 3,
            "dias": list(range(1, 16)),
            "personas": [{**base, "marcas": ["A"] * 15}],
        },
        {
            "pagina": 4,
            "dias": list(range(16, 32)),
            "personas": [{**base, "marcas": [None] + ["A"] * 15}],
        },
    ]
    return b


def test_columnas_complementarias_sin_rellenar_vacios():
    b = fragmentada()
    r = reconstruir_columnas(b)
    assert len(r["personas"]) == 1
    assert r["personas"][0]["marcas"] == ["A"] * 15 + [None] + ["A"] * 15
    assert r["personas"][0]["localizadores_marcas"]["16"]["pagina"] == 4
    assert b["personas"] == []


@pytest.mark.parametrize(
    "fallo", ["mes", "identidad", "contradiccion", "ancla", "hueco"]
)
def test_no_unir_fragmentos_ambiguos(fallo):
    b = fragmentada()
    f = b["fragmentos_columnas"][-1]
    if fallo == "mes":
        f["mes"] = 6
    if fallo == "identidad":
        f["personas"][0]["dni"] = "99990002"
    if fallo == "contradiccion":
        f["dias"][0] = 15
    if fallo == "ancla":
        f["personas"][0]["id_fila"] = "002"
    if fallo == "hueco":
        f["dias"].pop()
        f["personas"][0]["marcas"].pop()
    with pytest.raises(ValueError):
        reconstruir_columnas(b)


def bloque_asistencia(personas, leyenda=None):
    return {
        "tipo": "asistencia",
        "institucion": "9999001 IE WEB FICTICIA",
        "cod_mod": "9999001",
        "nivel": "PRIMARIA",
        "anio": 2026,
        "mes": 7,
        "pagina": 2,
        "turno": "MAÑANA",
        "leyenda": leyenda or [],
        "personas": personas,
    }


def extraida(bloque):
    return {
        "datos": {"bloques": [bloque]},
        "modelo": "FICTICIO",
        "version": "leyendas-prueba",
        "sha256": "f" * 64,
    }


def persona_ficticia(n, marcas):
    return {
        "fila": n,
        "dni": f"9999{n:04}",
        "nombre": f"Persona Ficticia {n}",
        "cargo": "DOCENTE",
        "marcas": marcas,
    }


def test_feriado_vertical_no_materializa_inasistencias(conn, tmp_path):
    _, item = base_ficticia(conn, tmp_path)
    b = bloque_asistencia(
        [
            persona_ficticia(i, [letra] + ["A"] * 30)
            for i, letra in enumerate("FERIADO", 1)
        ]
    )
    r = aplicar_extraccion(conn, item, extraida(b))
    dias = conn.execute(
        "SELECT codigo_reportado_raw,codigo_interpretado,estado_captura,evidencia_interpretacion FROM asistencia_dia WHERE fecha='2026-07-01' ORDER BY codigo_reportado_raw"
    ).fetchall()
    assert len(dias) == 7
    assert all(
        d["codigo_interpretado"] == "FERIADO"
        and d["evidencia_interpretacion"]["texto"] == "FERIADO"
        for d in dias
    )
    assert sorted(d["codigo_reportado_raw"] for d in dias) == sorted("FERIADO")
    assert (
        conn.execute(
            "SELECT count(*) n FROM validacion_reporte WHERE codigo_regla='ASISTENCIA_CONTRADICTORIA'"
        ).fetchone()["n"]
        == 0
    )
    assert aplicar_extraccion(conn, item, extraida(b))["ids"] == r["ids"]
    assert conn.execute("SELECT count(*) n FROM asistencia_dia").fetchone()["n"] == 217


def test_feriado_incompleto_no_reinterpreta_i_aislada(conn, tmp_path):
    _, item = base_ficticia(conn, tmp_path)
    aplicar_extraccion(
        conn, item, extraida(bloque_asistencia([persona_ficticia(1, ["I"] * 31)]))
    )
    assert (
        conn.execute(
            "SELECT count(*) n FROM asistencia_dia WHERE codigo_interpretado IS NOT NULL"
        ).fetchone()["n"]
        == 0
    )


def preparar_cruce(conn, tmp_path):
    _, item = base_ficticia(conn, tmp_path)
    b = bloque_asistencia(
        [persona_ficticia(1, ["L", "I", "P", "A"] + ["L"] * 27)],
        [
            {"codigo": "L", "descripcion": "Licencia con goce de remuneraciones"},
            {"codigo": "I", "descripcion": "Inasistencia injustificada"},
            {"codigo": "P", "descripcion": "Licencia sin goce de remuneraciones"},
            {"codigo": "A", "descripcion": "Asistencia remunerada"},
        ],
    )
    rid = aplicar_extraccion(conn, item, extraida(b))["ids"][0]
    cal = {
        "tipo": "calendario",
        "institucion": b["institucion"],
        "cod_mod": b["cod_mod"],
        "nivel": "PRIMARIA",
        "anio": 2026,
        "meses": [{"mes": 7, "codigos": ["U1"] * 3 + ["D"] + ["U1"] * 27}],
        "leyenda": [
            {"codigo": "U1", "descripcion": "Unidad de aprendizaje 1 remunerada"},
            {"codigo": "D", "descripcion": "Descanso no remunerado"},
        ],
    }
    cvid = aplicar_extraccion(conn, {**item, "tipo": "CALENDARIO_2026"}, extraida(cal))[
        "ids"
    ][0]
    tid = conn.execute(
        "SELECT trabajador_en_reporte_id FROM trabajador_en_reporte WHERE reporte_asistencia_id=%s",
        (rid,),
    ).fetchone()["trabajador_en_reporte_id"]
    return rid, cvid, tid


def test_cruce_integrado_y_salida_congelada(conn, tmp_path):
    rid, _cvid, tid = preparar_cruce(conn, tmp_path)
    cr = cruce_persona(conn, tid)
    assert [d["resultado"] for d in cr["dias"][:4]] == [
        "REMUNERADO",
        "NO_REMUNERADO",
        "NO_REMUNERADO",
        "OBSERVACION",
    ]
    assert cr["faltas_en_dia_remunerado"] == 1
    assert sum(cr["resumen"].values()) == 31
    cons = generar_consolidado(conn, date(2026, 7, 1), "PRIMARIA")
    congelado = conn.execute(
        "SELECT fuente_calculo FROM consolidado_dre_detalle WHERE consolidado_dre_id=%s",
        (cons.consolidado_dre_id,),
    ).fetchone()["fuente_calculo"]["cruce_remuneracion"]
    uid = crear_usuario(
        conn, "Prueba", "leyenda@example.invalid", "ClaveDePrueba_2026z"
    )
    guardar_asistencia(
        conn,
        rid,
        {
            "codigo": "L",
            "tipo_dia": "Licencia con goce de remuneraciones",
            "remuneracion": "",
            "es_falta": "NO",
        },
        "Regla en revisión",
        huella_reporte(conn, rid),
        uuid.uuid4(),
        uid,
    )
    assert cruce_persona(conn, tid)["dias"][0]["resultado"] == "PENDIENTE"
    assert (
        conn.execute(
            "SELECT fuente_calculo FROM consolidado_dre_detalle WHERE consolidado_dre_id=%s",
            (cons.consolidado_dre_id,),
        ).fetchone()["fuente_calculo"]["cruce_remuneracion"]
        == congelado
    )


def test_calendario_clona_leyenda_y_rechaza_edicion_obsoleta(conn, tmp_path):
    _, cvid, _ = preparar_cruce(conn, tmp_path)
    uid = crear_usuario(conn, "Prueba", "cal@example.invalid", "ClaveDePrueba_2026z")
    h, op = huella_calendario(conn, cvid), uuid.uuid4()
    f = {
        "codigo": "D",
        "tipo_dia": "Descanso",
        "remuneracion": "SI",
        "grupo_actividad": "NO_LECTIVO_NI_GESTION",
    }
    nuevo = guardar_calendario(conn, cvid, f, "Regla explícita de prueba", h, op, uid)
    assert str(
        guardar_calendario(conn, cvid, f, "Regla explícita de prueba", h, op, uid)
    ) == str(nuevo)
    assert nuevo != cvid
    versiones = conn.execute(
        "SELECT clasificacion_codigos,estado FROM calendarizacion_version ORDER BY version"
    ).fetchall()
    assert versiones[0]["clasificacion_codigos"]["D"]["es_remunerado"] is False
    assert versiones[1]["clasificacion_codigos"]["D"]["es_remunerado"] is True
    assert versiones[1]["estado"] == "BORRADOR"
    with pytest.raises(ErrorDeTrabajo):
        guardar_calendario(conn, cvid, f, "Otra edición", h, uuid.uuid4(), uid)


def test_web_muestra_clasificaciones_y_cruce(client, conn, tmp_path):  # noqa: F811
    rid, cvid, tid = preparar_cruce(conn, tmp_path)
    conn.commit()
    for ruta in (f"/reportes/{rid}", f"/calendarios/{cvid}"):
        r = client.get(ruta)
        assert r.status_code == 200
        assert (
            "Remuneración"
            if "/reportes/" in ruta
            else "Clasificación por tipo de remuneración"
        ) in r.text
        assert (
            "Significado de la leyenda"
            if "/reportes/" in ruta
            else "Clasificación por tipo de día"
        ) in r.text
    detalle = client.get(f"/reportes/{rid}/personas/{tid}")
    assert detalle.status_code == 200
    assert "Cruce de remuneración por día" in detalle.text
    assert "Licencia sin goce" in detalle.text


def pdf_partido_ficticio(ruta):
    """PDF vectorial pequeño con dos identidades y una página de días; sin dependencia extra."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    w = PdfWriter()
    fuente = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )

    def texto(x, y, t):
        return f"BT /F1 9 Tf {x} {500 - y} Td ({t}) Tj ET"

    def linea(x1, y1, x2, y2):
        return f"{x1} {500 - y1} m {x2} {500 - y2} l S"

    for identidad in (True, False):
        pagina = w.add_blank_page(width=750, height=500)
        pagina[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): fuente})}
        )
        cmds = ["0.5 w", texto(350, 50, "JULIO / 2026")]
        if identidad:
            cmds.extend(
                [
                    texto(55, 104, "DNI"),
                    texto(55, 151, "99990001"),
                    texto(55, 181, "99990002"),
                ]
            )
            cmds.extend(linea(40, y, 700, y) for y in (90, 130, 160, 190))
        else:
            cmds.extend(linea(40 + 21 * i, 90, 40 + 21 * i, 190) for i in range(32))
            cmds.extend(linea(40, y, 691, y) for y in (90, 110, 130, 160, 190))
            for d in range(1, 32):
                x = 40 + 21 * (d - 1) + 6
                cmds.extend(
                    [
                        texto(x, 104, str(d)),
                        texto(x, 120, "LMMJVSD"[date(2026, 7, d).weekday()]),
                    ]
                )
                if d != 4:
                    cmds.extend(
                        [texto(x, 151, "I" if d == 2 else "A"), texto(x, 181, "L")]
                    )
        stream = DecodedStreamObject()
        stream.set_data("\n".join(cmds).encode())
        pagina[NameObject("/Contents")] = w._add_object(stream)
    w.write(ruta)


def test_pdf_digital_reconstruye_dias_por_bordes_y_preserva_ocr(tmp_path):
    from asistia.cierre.pdf_digital import cotejar_pdf_digital

    ruta = tmp_path / "tabla_dividida.pdf"
    pdf_partido_ficticio(ruta)
    datos = extraida(
        bloque_asistencia(
            [persona_ficticia(1, ["A"] * 31), persona_ficticia(2, ["A"] * 31)]
        )
    )
    r = cotejar_pdf_digital(ruta, datos)
    assert r["cotejo_geometrico"][0]["pagina_identidad"] == 1
    personas = r["datos"]["bloques"][0]["personas"]
    assert personas[0]["marcas"][1] == "I"
    assert personas[0]["marcas"][3] is None
    assert personas[1]["marcas"] == ["L"] * 3 + [None] + ["L"] * 27
    assert personas[0]["marcas_ocr"] == ["A"] * 31
    assert cotejar_pdf_digital(ruta, r) == r
    assert datos["datos"]["bloques"][0]["personas"][0]["marcas"] == ["A"] * 31


def test_ampliacion_leyenda_crea_version_sin_perder_filas(conn, tmp_path):
    from asistia.cierre.leyendas import ampliar_leyenda_reporte

    rid, _, tid = preparar_cruce(conn, tmp_path)
    antes = conn.execute(
        "SELECT to_jsonb(a) dato FROM asistencia_dia a WHERE trabajador_en_reporte_id=%s ORDER BY fecha",
        (tid,),
    ).fetchall()
    categorias = normalizar_leyenda(
        [{"codigo": "L", "descripcion": "Licencia con goce de remuneraciones"}],
        "asistencia",
    )
    nuevo = ampliar_leyenda_reporte(
        conn, rid, categorias, "Relectura de leyenda del original"
    )
    assert nuevo != rid
    assert (
        ampliar_leyenda_reporte(
            conn, rid, categorias, "Relectura de leyenda del original"
        )
        == nuevo
    )
    assert (
        conn.execute(
            "SELECT to_jsonb(a) dato FROM asistencia_dia a WHERE trabajador_en_reporte_id=%s ORDER BY fecha",
            (tid,),
        ).fetchall()
        == antes
    )
    assert conn.execute("SELECT count(*) n FROM asistencia_dia").fetchone()["n"] == 62
    assert conn.execute(
        "SELECT version_padre_id FROM reporte_asistencia WHERE reporte_asistencia_id=%s",
        (nuevo,),
    ).fetchone()["version_padre_id"] == uuid.UUID(rid)


def test_equivalencias_sin_titulo_leyenda_y_dimension_desconocida(tmp_path):
    w = openpyxl.Workbook()
    w.active.append(["L", "Lectivo", "BIMESTRE", "INICIO"])
    w.active.append(["G", "Gestión"])
    w.active.append(["D", "Sábados, domingos y feriados"])
    w.active.append(["N", "Persona Ficticia"])
    p = tmp_path / "sin_rotulo.xlsx"
    w.save(p)
    libro = openpyxl.load_workbook(p, read_only=True)
    libro.active.reset_dimensions()
    leyenda = leer_leyenda_xlsx(libro.active)
    libro.close()
    assert [c["codigo"] for c in leyenda] == ["L", "G", "D"]


def test_feriado_celda_combinada_respeta_ambito_y_vacios(conn, tmp_path):
    from asistia.importar.asistencia import importar_asistencia

    from .fixtures_excel import construir_anexo3_minimo

    base_ficticia(conn, tmp_path)
    ruta = construir_anexo3_minimo(
        tmp_path / "feriado_combinado.xlsx",
        "9999001 IE WEB FICTICIA",
        personas=[
            {
                "dni": "99990001",
                "nombres": "Uno Ficticio",
                "cargo": "DOCENTE",
                "marcas": {1: "F\nE\nR\nI\nA\nD\nO", 2: "A"},
            },
            {
                "dni": "99990002",
                "nombres": "Dos Ficticio",
                "cargo": "DOCENTE",
                "marcas": {2: "A"},
            },
        ],
    )
    wb = openpyxl.load_workbook(ruta)
    wb.active.merge_cells("E12:E13")
    wb.save(ruta)
    wb.close()
    resultado = importar_asistencia(conn, ruta)
    assert resultado.error is None
    filas = conn.execute(
        "SELECT codigo_reportado_raw,codigo_interpretado FROM asistencia_dia WHERE fecha='2026-07-01' ORDER BY codigo_reportado_raw NULLS LAST"
    ).fetchall()
    assert [f["codigo_interpretado"] for f in filas] == ["FERIADO", "FERIADO"]
    assert filas[1]["codigo_reportado_raw"] is None
    assert (
        conn.execute(
            "SELECT count(*) n FROM asistencia_dia WHERE fecha<>'2026-07-01' AND codigo_interpretado IS NOT NULL"
        ).fetchone()["n"]
        == 0
    )

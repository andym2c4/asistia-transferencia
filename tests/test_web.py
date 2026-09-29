"""P06: recorrido HTTP real contra PostgreSQL aislado, con fuentes totalmente ficticias."""

from __future__ import annotations

import io
import uuid
from datetime import date

import openpyxl
import pytest

from asistia.consolidado.generar import exportar_consolidado_xlsx
from asistia.web import create_app
from asistia.web.auth import crear_usuario
from asistia.web.lecturas import huella_reporte

from .conftest import TEST_DSN
from .fixtures_excel import construir_anexo3_minimo, construir_nexus_minimo


@pytest.fixture
def app(conn, tmp_path):
    crear_usuario(
        conn, "Operador de prueba", "operador@example.invalid", "UnaClaveDePrueba_2026"
    )
    return create_app(
        {
            "TESTING": True,
            "DEVELOPMENT_TOOLS": True,
            "SECRET_KEY": "clave_de_prueba_sin_uso_real_" * 2,
            "DATABASE_URL": TEST_DSN,
            "STORAGE_ROOT": tmp_path / "web",
            "ORIGINAL_ROOTS": [str(tmp_path)],
            "TRUSTED_HOSTS": ["localhost"],
        }
    )


@pytest.fixture
def client(app):
    c = app.test_client()
    c.get("/ingresar")
    response = post(
        c,
        "/ingresar",
        {"email": "operador@example.invalid", "password": "UnaClaveDePrueba_2026"},
    )
    assert response.status_code == 303
    return c


def post(client, path, datos=None, **kwargs):
    with client.session_transaction() as s:
        token = s.get("csrf")
    return client.post(
        path,
        data={"csrf": token, "operacion": str(uuid.uuid4()), **(datos or {})},
        **kwargs,
    )


def test_async_conserva_confirmacion_y_formulario_al_vencer_sesion(client, conn):
    data = {
        "accion": "INICIAR",
        "uso": "PRACTICA",
        "cohorte": "Ficticia",
        "tarea_referencia": "T1",
        "lote_referencia": "L1",
        "version_referencia": "V1",
        "criterio_terminacion": "Comprobación",
        "periodo": "2026-07",
        "nivel": "PRIMARIA",
    }
    response = post(client, "/trabajo", data, headers={"X-ASISTIA-Async": "1"})
    assert response.status_code == 200
    assert response.json == {"redirect": "/resultados"}
    assert "Registro de trabajo guardado" in client.get(response.json["redirect"]).text
    conn.execute("DELETE FROM web_sesion")
    conn.commit()
    response = post(client, "/trabajo", data, headers={"X-ASISTIA-Async": "1"})
    assert response.json["redirect"].startswith("/ingresar")
    assert conn.execute("SELECT count(*) AS n FROM web_tarea").fetchone()["n"] == 1


def fuentes(tmp_path, *, sin_dni=False, marca="A"):
    nexus = construir_nexus_minimo(
        [
            {
                "CODMOD I.E.": "9999001",
                "CODIGO DE PLAZA": "PZ-WEB-01",
                "TIPO DE TRABAJADOR": "DOCENTE",
                "SITUACION LABORAL": "NOMBRADO",
                "TIPO DE REGISTRO": "ORGANICA",
                "DOCUMENTO DE IDENTIDAD": "99990001",
                "APELLIDO PATERNO": "Ejemplo",
                "APELLIDO MATERNO": "Ficticio",
                "NOMBRES": "Persona Uno",
                "FECHA DE INICIO": date(2026, 1, 1),
                "CODIGO LOCAL": "999901",
                "DISTRITO": "LUYA",
                "NIVEL EDUCATIVO": "Primaria",
                "NOMBRE DE LA INSTITUCION EDUCATIVA": "9999001 IE WEB FICTICIA",
                "CARGO": "DOCENTE",
            }
        ],
        tmp_path / "nexus.xlsx",
    )
    asistencia = construir_anexo3_minimo(
        tmp_path / "asistencia.xlsx",
        institucion_raw="9999001 IE WEB FICTICIA",
        personas=[
            {
                "dni": None if sin_dni else "99990001",
                "nombres": "Ejemplo Ficticio Persona Uno",
                "cargo": "DOCENTE",
                "marcas": {1: marca, 2: "A", 3: "I", 4: "J", 5: "A"},
            }
        ],
    )
    cal = tmp_path / "calendario.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "CALENDARIZACIÓN DEL AÑO ESCOLAR 2026"
    ws["A2"] = "NOMBRE DE LA IE"
    ws["B2"] = "9999001 IE WEB FICTICIA"
    ws["A3"] = "NIVEL O CICLO"
    ws["B3"] = "PRIMARIA"
    ws["A12"] = "JULIO"
    for d in range(1, 32):
        ws.cell(12, d + 1, d)
        ws.cell(13, d + 1, "L")
    wb.save(cal)
    return nexus, cal, asistencia


def cargar(client, path, tipo, *, op=None):
    return post(
        client,
        "/documentos",
        {
            "tipo": tipo,
            "fecha_corte": "2026-06-01",
            "periodo": "2026-07",
            "nivel": "PRIMARIA",
            "operacion": str(op or uuid.uuid4()),
            "archivos": (io.BytesIO(path.read_bytes()), path.name),
        },
        follow_redirects=False,
    )


def ciclo(client, conn, tmp_path, **kwargs):
    nx, cal, asistencia = fuentes(tmp_path, **kwargs)
    for path, tipo in ((nx, "nexus"), (cal, "calendario"), (asistencia, "asistencia")):
        r = cargar(client, path, tipo)
        assert r.status_code == 303
    estados = conn.execute(
        "SELECT tipo,estado,mensaje FROM web_carga ORDER BY tipo"
    ).fetchall()
    assert all(r["estado"] == "PROCESADO" for r in estados), estados
    rid = conn.execute(
        "SELECT reporte_asistencia_id FROM reporte_asistencia ORDER BY version DESC LIMIT 1"
    ).fetchone()["reporte_asistencia_id"]
    return rid, (nx, cal, asistencia)


def revisar(client, conn, rid):
    return post(
        client,
        f"/reportes/{rid}/revisar",
        {
            "huella": huella_reporte(conn, rid),
            "motivo": "Cotejado con la fuente ficticia de prueba.",
        },
    )


def generar(client, estado="BORRADOR", op=None):
    return post(
        client,
        "/consolidado",
        {
            "estado": estado,
            "periodo": "2026-07",
            "nivel": "PRIMARIA",
            "operacion": str(op or uuid.uuid4()),
        },
    )


def test_acceso_csrf_sesion_y_host(app, conn):
    c = app.test_client()
    assert c.get("/").status_code == 302
    assert c.get("/resultados/registro.csv").status_code == 302
    assert c.post("/consolidado", data={"estado": "BORRADOR"}).status_code == 400
    assert c.get("/salud", headers={"Host": "sitio-ajeno.invalid"}).status_code == 400
    assert c.get("/salud").json == {"estado": "disponible"}
    c.get("/ingresar")
    assert (
        post(
            c,
            "/ingresar",
            {"email": "operador@example.invalid", "password": "incorrecta"},
        ).status_code
        == 401
    )
    assert (
        post(
            c,
            "/ingresar",
            {
                "email": "operador@example.invalid",
                "password": "UnaClaveDePrueba_2026",
                "next": "//sitio-ajeno.invalid",
            },
        ).location
        == "/"
    )
    protected = c.get("/")
    assert protected.status_code == 200
    assert "no-store" in protected.headers["Cache-Control"]
    assert "form-action 'self'" in protected.headers["Content-Security-Policy"]
    with c.session_transaction() as s:
        stolen_session = dict(s)
    assert post(c, "/salir").status_code == 303
    with c.session_transaction() as s:
        s.update(stolen_session)
    assert (
        c.get("/").status_code == 302
    )  # La sesión revocada no vuelve con la cookie firmada.


def test_ciclo_web_fuente_revision_y_descarga(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    for path in (
        "/",
        "/documentos",
        "/revision",
        "/consolidado",
        "/instituciones",
        "/calendarios",
        "/resultados",
        f"/reportes/{rid}",
        f"/reportes/{rid}/fuente",
    ):
        response = client.get(path)
        assert response.status_code == 200, (path, response.get_data(as_text=True))
    for c in conn.execute("SELECT web_carga_id FROM web_carga").fetchall():
        assert client.get(f"/documentos/{c['web_carga_id']}").status_code == 200
        assert (
            client.get(f"/documentos/{c['web_carga_id']}/original").status_code == 200
        )
    cv = conn.execute(
        "SELECT calendarizacion_version_id FROM calendarizacion_version"
    ).fetchone()["calendarizacion_version_id"]
    assert client.get(f"/calendarios/{cv}").status_code == 200
    assert client.get(f"/calendarios/{cv}/original").status_code == 200
    tid = conn.execute(
        "SELECT trabajador_en_reporte_id FROM trabajador_en_reporte"
    ).fetchone()["trabajador_en_reporte_id"]
    assert client.get(f"/reportes/{rid}/personas/{tid}?dni=99990001").status_code == 200
    assert generar(client, "REVISADO").status_code == 409
    assert revisar(client, conn, rid).status_code == 303
    assert generar(client, "REVISADO").status_code == 409  # Cruce aún incompleto.
    r = generar(client, "BORRADOR")
    assert r.status_code == 303, r.get_data(as_text=True)
    s = conn.execute("SELECT * FROM web_salida").fetchone()
    assert s["estado_revision"] == "BORRADOR"
    assert s["manifiesto"]["personas_distintas"] == 1
    assert s["manifiesto"]["cobertura_estado"] == "NO_CALCULABLE"
    assert client.get(r.location).status_code == 200
    xlsx = client.get(r.location + "/descargar")
    assert xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx.data))
    ws = wb.active
    assert "BORRADOR" in ws["A7"].value
    # Sin reglas locales de remuneración no se infieren faltas a partir de letras.
    assert ws["H12"].value == "Pendiente" and ws["J12"].value == "Pendiente"
    assert len(wb.sheetnames) == 1
    assert "2026" in ws["A6"].value


def test_recepcion_real_desconocida_y_duplicados(client, conn, tmp_path):
    nx, _, _ = fuentes(tmp_path)
    op = uuid.uuid4()
    assert cargar(client, nx, "nexus", op=op).status_code == 303
    assert cargar(client, nx, "nexus", op=op).status_code == 303
    assert conn.execute("SELECT count(*) AS n FROM web_recepcion").fetchone()["n"] == 1
    assert cargar(client, nx, "nexus").status_code == 303
    assert conn.execute("SELECT count(*) AS n FROM web_recepcion").fetchone()["n"] == 2
    assert conn.execute("SELECT count(*) AS n FROM nexus_carga").fetchone()["n"] == 1
    assert (
        conn.execute("SELECT count(*) AS n FROM vinculo_trabajador_ie").fetchone()["n"]
        == 1
    )
    assert (
        conn.execute("SELECT recibido_real_en FROM web_recepcion LIMIT 1").fetchone()[
            "recibido_real_en"
        ]
        is None
    )
    assert (
        conn.execute("SELECT count(*) AS n FROM web_carga_intento").fetchone()["n"] == 1
    )


def test_precondicion_y_archivo_no_soportado_conservan_original(client, conn, tmp_path):
    nx, _, asistencia = fuentes(tmp_path)
    assert cargar(client, asistencia, "asistencia").status_code == 303
    c = conn.execute("SELECT * FROM web_carga").fetchone()
    assert c["estado"] == "RECIBIDO"
    assert (
        client.get(f"/documentos/{c['web_carga_id']}/original").data
        == asistencia.read_bytes()
    )
    cargar(client, nx, "nexus")
    assert (
        post(client, f"/documentos/{c['web_carga_id']}/reintentar").status_code == 303
    )
    assert (
        conn.execute(
            "SELECT estado FROM web_carga WHERE web_carga_id=%s", (c["web_carga_id"],)
        ).fetchone()["estado"]
        == "PROCESADO"
    )
    r = post(
        client,
        "/documentos",
        {
            "tipo": "asistencia",
            "periodo": "2026-07",
            "archivos": (io.BytesIO(b"contenido no soportado"), "foto.jpg"),
        },
    )
    assert r.status_code == 303
    c = conn.execute("SELECT * FROM web_carga WHERE estado='NO_SOPORTADO'").fetchone()
    assert c is not None
    assert (
        client.get(f"/documentos/{c['web_carga_id']}/original").data
        == b"contenido no soportado"
    )


def test_omision_y_alerta_de_reporte_sin_detalle_no_desaparecen(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path, sin_dni=True)
    conn.execute(
        "INSERT INTO validacion_reporte(reporte_asistencia_id,codigo_regla,severidad,mensaje) VALUES (%s,'IDENTIDAD_PRUEBA','ERROR','Pendiente ficticio')",
        (rid,),
    )
    conn.commit()
    assert revisar(client, conn, rid).status_code == 409
    assert generar(client, "REVISADO").status_code == 409
    r = generar(client)
    assert r.status_code == 303
    salida = conn.execute("SELECT * FROM web_salida").fetchone()
    assert salida["manifiesto"]["filas_omitidas_sin_identidad"] == 1
    assert salida["manifiesto"]["alertas_pendientes"] == 1
    assert salida["manifiesto"]["filas_detalle"] == 0
    assert len(salida["manifiesto"]["reportes"]) == 1


def test_identidad_manual_y_marca_conservan_valor_recibido(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path, sin_dni=True)
    tid = conn.execute(
        "SELECT trabajador_en_reporte_id FROM trabajador_en_reporte"
    ).fetchone()["trabajador_en_reporte_id"]
    vid = conn.execute(
        "SELECT vinculo_trabajador_ie_id FROM vinculo_trabajador_ie"
    ).fetchone()["vinculo_trabajador_ie_id"]
    r = post(
        client,
        f"/reportes/{rid}/personas/{tid}/identidad",
        {
            "vinculo": vid,
            "huella": huella_reporte(conn, rid),
            "motivo": "DNI cotejado con fuente ficticia.",
        },
    )
    assert r.status_code == 303, r.get_data(as_text=True)
    fila = conn.execute("SELECT * FROM trabajador_en_reporte").fetchone()
    assert fila["dni_reportado_raw"] == ""
    assert fila["trabajador_id"] is not None
    did = conn.execute(
        "SELECT asistencia_dia_id FROM asistencia_dia WHERE fecha='2026-07-01'"
    ).fetchone()["asistencia_dia_id"]
    old_hash = huella_reporte(conn, rid)
    op = uuid.uuid4()
    datos = {
        "operacion": str(op),
        "codigo": "J",
        "captura": "REGISTRADO",
        "motivo": "Cambio cotejado con el original ficticio.",
        "huella": old_hash,
    }
    assert post(client, f"/reportes/{rid}/dias/{did}", datos).status_code == 303
    assert post(client, f"/reportes/{rid}/dias/{did}", datos).status_code == 303
    dia = conn.execute(
        "SELECT * FROM asistencia_dia WHERE asistencia_dia_id=%s", (did,)
    ).fetchone()
    assert dia["codigo_reportado_raw"] == "A"
    assert (
        dia["estado_asistencia_id"]
        == conn.execute(
            "SELECT estado_asistencia_id FROM catalogo_estado_asistencia WHERE codigo='J'"
        ).fetchone()["estado_asistencia_id"]
    )
    assert (
        conn.execute(
            "SELECT count(*) AS n FROM web_revision_evento WHERE accion='CORREGIR_MARCA'"
        ).fetchone()["n"]
        == 1
    )
    assert (
        post(
            client,
            f"/reportes/{rid}/dias/{did}",
            {**datos, "operacion": str(uuid.uuid4()), "codigo": "A"},
        ).status_code
        == 409
    )
    assert (
        post(
            client, f"/reportes/{rid}/dias/{did}", {**datos, "codigo": "A"}
        ).status_code
        == 409
    )


def test_salida_inmutable_reexportacion_y_rectificacion(client, conn, tmp_path):
    rid, fuentes_originales = ciclo(client, conn, tmp_path)
    revisar(client, conn, rid)
    op = uuid.uuid4()
    r = generar(client, "BORRADOR", op)
    contenido = client.get(r.location + "/descargar").data
    assert generar(client, "BORRADOR", op).location == r.location
    assert generar(client, "REVISADO", op).status_code == 409
    assert conn.execute("SELECT count(*) AS n FROM web_salida").fetchone()["n"] == 1
    conn.execute("UPDATE trabajador SET nombres='Cambio ficticio posterior'")
    conn.commit()
    assert client.get(r.location + "/descargar").data == contenido
    cid = conn.execute("SELECT consolidado_dre_id FROM web_salida").fetchone()[
        "consolidado_dre_id"
    ]
    otra = tmp_path / "reexportada.xlsx"
    exportar_consolidado_xlsx(conn, str(cid), otra)
    # La consola conserva el libro completo; la web deriva una hoja del original.
    from asistia.web.exportacion import primera_hoja

    assert primera_hoja(otra).getvalue() == contenido
    nueva = construir_anexo3_minimo(
        tmp_path / "rectificacion.xlsx",
        institucion_raw="9999001 IE WEB FICTICIA",
        personas=[
            {
                "dni": "99990001",
                "nombres": "Ejemplo Ficticio Persona Uno",
                "cargo": "DOCENTE",
                "marcas": {1: "I", 2: "A", 3: "I", 4: "J", 5: "A"},
            }
        ],
    )
    cargar(client, nueva, "asistencia")
    assert "Hay fuentes o decisiones posteriores" in client.get(r.location).get_data(
        as_text=True
    )
    assert client.get(r.location + "/descargar").data == contenido
    assert fuentes_originales[2].is_file()


def test_fallo_al_preparar_salida_no_publica_parcial(
    client, conn, tmp_path, monkeypatch
):
    import asistia.web.salidas as modulo

    ciclo(client, conn, tmp_path)
    exportador = modulo.exportar_consolidado_xlsx

    def fallar(*args, **kwargs):
        exportador(*args, **kwargs)
        raise RuntimeError("fallo sintético después del archivo")

    monkeypatch.setattr(modulo, "exportar_consolidado_xlsx", fallar)
    with pytest.raises(RuntimeError, match="fallo sintético"):
        generar(client)
    assert (
        conn.execute("SELECT count(*) AS n FROM consolidado_dre").fetchone()["n"] == 0
    )
    assert conn.execute("SELECT count(*) AS n FROM web_salida").fetchone()["n"] == 0
    monkeypatch.setattr(modulo, "exportar_consolidado_xlsx", exportador)
    assert generar(client).status_code == 303


def test_mapeo_niveles_sin_reescribir_fuentes(client, conn, tmp_path):
    ciclo(client, conn, tmp_path)
    pares = [
        ("Primaria", "PRIMARIA"),
        (" PRIMARIA ", "PRIMARIA"),
        ("Inicial - Jardín", "INICIAL"),
        ("Básica Especial-Inicial", "CEBE"),
        ("Administración", None),
    ]
    for valor, esperado in pares:
        assert (
            conn.execute("SELECT fn_nivel_canonico(%s) AS n", (valor,)).fetchone()["n"]
            == esperado
        )
    assert (
        conn.execute("SELECT nivel_modalidad FROM institucion_educativa").fetchone()[
            "nivel_modalidad"
        ]
        == "Primaria"
    )
    assert "9999001 IE WEB FICTICIA" in client.get(
        "/?periodo=2026-07&nivel=PRIMARIA"
    ).get_data(as_text=True)
    assert client.get("/?nivel=ADMINISTRACION").status_code == 400


def test_captura_trabajo_pausa_reintento_y_calidad_pendiente(client, conn, tmp_path):
    ciclo(client, conn, tmp_path)
    s = generar(client)
    sid = s.location.rsplit("/", 1)[-1]
    op = uuid.uuid4()
    datos = {
        "operacion": str(op),
        "accion": "INICIAR",
        "uso": "PRACTICA",
        "cohorte": "prueba",
        "tarea_referencia": "tarea-1",
        "lote_referencia": "lote-ficticio",
        "version_referencia": "v1",
        "criterio_terminacion": "Preparar y descargar salida de prueba.",
    }
    assert post(client, "/trabajo", datos).status_code == 303
    assert post(client, "/trabajo", datos).status_code == 303
    assert conn.execute("SELECT count(*) AS n FROM web_tarea").fetchone()["n"] == 1
    assert (
        post(client, "/trabajo", {"accion": "PAUSAR", "tarea": str(op)}).status_code
        == 303
    )
    assert (
        conn.execute(
            "SELECT count(*) AS n FROM web_tarea_intervalo WHERE fin IS NULL"
        ).fetchone()["n"]
        == 0
    )
    assert (
        post(client, "/trabajo", {"accion": "RETOMAR", "tarea": str(op)}).status_code
        == 303
    )
    assert (
        post(
            client,
            "/trabajo",
            {
                "accion": "FINALIZAR",
                "tarea": str(op),
                "salida": sid,
                "ayuda_tecnica": "si",
            },
        ).status_code
        == 303
    )
    tarea = conn.execute("SELECT * FROM web_tarea").fetchone()
    assert tarea["estado"] == "FINALIZADA" and tarea["ayuda_tecnica"]
    assert (
        conn.execute("SELECT count(*) AS n FROM web_tarea_intervalo").fetchone()["n"]
        == 2
    )
    assert "aún no medidos" in client.get("/resultados").get_data(as_text=True)
    csv = client.get("/resultados/registro.csv").get_data(as_text=True)
    assert "PRACTICA" in csv and "tarea-1" in csv


def test_captura_detallada_atribuye_solo_tarea_activa_del_ambito(
    client, conn, tmp_path
):
    rid, archivos = ciclo(client, conn, tmp_path)
    tid = uuid.uuid4()
    datos = {
        "operacion": str(tid),
        "accion": "INICIAR",
        "uso": "PRACTICA",
        "cohorte": "Ficticia",
        "tarea_referencia": "T1",
        "lote_referencia": "L1",
        "version_referencia": "V1",
        "criterio_terminacion": "Cotejo ficticio",
        "periodo": "2026-07",
        "nivel": "PRIMARIA",
    }
    assert post(client, "/trabajo", datos).status_code == 303
    assert revisar(client, conn, rid).status_code == 303
    assert (
        post(client, "/trabajo", {"accion": "PAUSAR", "tarea": str(tid)}).status_code
        == 303
    )
    dia = conn.execute(
        "SELECT asistencia_dia_id FROM asistencia_dia ORDER BY fecha LIMIT 1"
    ).fetchone()["asistencia_dia_id"]
    payload = {
        "operacion": str(uuid.uuid4()),
        "codigo": "J",
        "captura": "REGISTRADO",
        "motivo": "Corrección ficticia en pausa",
        "huella": huella_reporte(conn, rid),
    }
    assert post(client, f"/reportes/{rid}/dias/{dia}", payload).status_code == 303
    assert post(client, f"/reportes/{rid}/dias/{dia}", payload).status_code == 303
    cargar(client, archivos[2], "asistencia")
    captured = client.get("/resultados/evidencia.json")
    assert (
        captured.status_code == 200 and "no-store" in captured.headers["Cache-Control"]
    )
    data = captured.json
    assert data["resultados_metricas"] == "NO_MEDIDO"
    assert len(data["tareas"]) == len(data["intervalos"]) == 1
    assert data["intervalos"][0]["fin"] is not None
    assert len(data["recepciones"]) == 4 and len(data["intentos"]) == 3
    assert len(data["intervenciones"]) == 2
    revision, correccion = data["intervenciones"]
    assert revision["web_tarea_id"] == str(tid) and not revision["es_correccion_datos"]
    assert correccion["web_tarea_id"] is None and correccion["es_correccion_datos"]
    assert revision["sha256"] == correccion["sha256"]
    assert "nombres_reportados_raw" not in captured.text
    assert "lote_referencia" in client.get("/resultados/registro.csv").text

    # Una tarea activa de otro nivel tampoco se atribuye al reporte.
    post(client, "/trabajo", {"accion": "ABANDONAR", "tarea": str(tid)})
    datos.update(operacion=str(uuid.uuid4()), nivel="SECUNDARIA")
    assert post(client, "/trabajo", datos).status_code == 303
    # La corrección generó una contradicción: para este caso se captura otra corrección,
    # sin dar por revisado ni resolver la alerta artificialmente.
    payload.update(
        operacion=str(uuid.uuid4()), codigo="I", huella=huella_reporte(conn, rid)
    )
    assert post(client, f"/reportes/{rid}/dias/{dia}", payload).status_code == 303
    data = client.get("/resultados/evidencia.json").json
    assert data["intervenciones"][-1]["web_tarea_id"] is None


def test_original_con_hash_incorrecto_no_se_descarga(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    c = conn.execute("SELECT * FROM web_carga WHERE tipo='asistencia'").fetchone()
    from pathlib import Path

    p = Path(c["ruta_copia"])
    p.chmod(0o600)
    p.write_bytes(b"cambio simulado solo sobre copia de prueba")
    assert client.get(f"/reportes/{rid}/original").status_code == 409
    assert generar(client).status_code == 409


@pytest.mark.parametrize("defecto", ["mes_ausente", "dia_ausente", "ilegible"])
def test_calendario_incompleto_no_publica_total_parcial(
    client, conn, tmp_path, defecto
):
    from asistia.consolidado.universo_esperado import dias_esperados_de_vinculo

    ciclo(client, conn, tmp_path)
    if defecto == "mes_ausente":
        conn.execute(
            "DELETE FROM dia_calendarizacion WHERE fecha >= '2026-07-01' AND fecha < '2026-08-01'"
        )
    elif defecto == "dia_ausente":
        conn.execute("DELETE FROM dia_calendarizacion WHERE fecha='2026-07-04'")
    else:
        conn.execute(
            "UPDATE dia_calendarizacion SET tipo_dia_id=NULL,estado_captura='ILEGIBLE' WHERE fecha='2026-07-04'"
        )
    conn.commit()
    vid = conn.execute(
        "SELECT vinculo_trabajador_ie_id FROM vinculo_trabajador_ie"
    ).fetchone()["vinculo_trabajador_ie_id"]
    with conn.cursor() as cur:
        dias = dias_esperados_de_vinculo(cur, vid, date(2026, 7, 1))
    assert dias.lectivos is None and dias.gestion is None
    assert dias.fuente_calculo["motivo"] == "calendario_incompleto"
    assert dias.fuente_calculo["fechas_calendarizacion_sin_resolver"]
    response = generar(client)
    assert response.status_code == 303
    assert "Desconocidos" in client.get(response.location).text


def test_texto_de_fuente_no_ejecuta_html_ni_formula(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    conn.execute("UPDATE trabajador SET nombres='=1+1'")
    conn.execute("UPDATE institucion_educativa SET nombre_ie='=1+1'")
    conn.execute(
        "UPDATE trabajador_en_reporte SET nombres_reportados_raw='<script>alert(1)</script>'"
    )
    conn.commit()
    html = client.get(f"/reportes/{rid}").get_data(as_text=True)
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html
    r = generar(client)
    wb = openpyxl.load_workbook(io.BytesIO(client.get(r.location + "/descargar").data))
    assert all(c.data_type != "f" for row in wb.active for c in row)
    assert wb.active["B12"].value == "=1+1"


def test_reintento_nexus_parcial_no_duplica_vinculos(
    client, conn, tmp_path, monkeypatch
):
    from asistia.importar import nexus as modulo

    nx, _, _ = fuentes(tmp_path)
    libro = openpyxl.load_workbook(nx)
    fila = [c.value for c in libro.active[2]]
    fila[1] = "PZ-WEB-02"
    fila[5] = "99990002"
    libro.active.append(fila)
    libro.save(nx)
    original = modulo._obtener_o_crear_trabajador

    def fallo_una_fila(cur, dni, *args):
        if dni == "99990002":
            raise RuntimeError("fallo de fila ficticio")
        return original(cur, dni, *args)

    monkeypatch.setattr(modulo, "_obtener_o_crear_trabajador", fallo_una_fila)
    cargar(client, nx, "nexus")
    c = conn.execute("SELECT * FROM web_carga").fetchone()
    assert c["estado"] == "PARCIAL"
    assert (
        conn.execute("SELECT count(*) AS n FROM vinculo_trabajador_ie").fetchone()["n"]
        == 1
    )
    monkeypatch.setattr(modulo, "_obtener_o_crear_trabajador", original)
    op = str(uuid.uuid4())
    assert (
        post(
            client, f"/documentos/{c['web_carga_id']}/reintentar", {"operacion": op}
        ).status_code
        == 303
    )
    assert (
        post(
            client, f"/documentos/{c['web_carga_id']}/reintentar", {"operacion": op}
        ).status_code
        == 303
    )
    assert (
        conn.execute("SELECT estado FROM web_carga").fetchone()["estado"] == "PROCESADO"
    )
    assert (
        conn.execute("SELECT count(*) AS n FROM vinculo_trabajador_ie").fetchone()["n"]
        == 2
    )
    assert (
        conn.execute("SELECT count(*) AS n FROM web_carga_intento").fetchone()["n"] == 2
    )


def test_reintento_no_sustituye_decisiones_manuales(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    assert revisar(client, conn, rid).status_code == 303
    c = conn.execute("SELECT * FROM web_carga WHERE tipo='asistencia'").fetchone()
    # Representa un intento parcial que ya recibió una revisión local.
    conn.execute(
        "UPDATE web_carga SET estado='PARCIAL' WHERE web_carga_id=%s",
        (c["web_carga_id"],),
    )
    conn.commit()
    assert (
        post(client, f"/documentos/{c['web_carga_id']}/reintentar").status_code == 409
    )
    assert (
        conn.execute("SELECT count(*) AS n FROM reporte_asistencia").fetchone()["n"]
        == 1
    )
    assert (
        conn.execute("SELECT estado FROM reporte_asistencia").fetchone()["estado"]
        == "VALIDADO"
    )


def test_lock_impide_escrituras_concurrentes(client, conn, tmp_path):
    from asistia.web.db import WRITE_LOCK

    ciclo(client, conn, tmp_path)
    conn.execute("SELECT pg_advisory_lock(%s)", (WRITE_LOCK,))
    conn.commit()
    try:
        response = generar(client)
        assert response.status_code == 409
        assert "otra operación" in response.get_data(as_text=True)
        assert (
            conn.execute("SELECT count(*) AS n FROM consolidado_dre").fetchone()["n"]
            == 0
        )
    finally:
        conn.execute("SELECT pg_advisory_unlock(%s)", (WRITE_LOCK,))
        conn.commit()


def test_contexto_entre_pestanas_y_contrato_explicito(client, conn, tmp_path):
    ciclo(client, conn, tmp_path)
    pagina = client.get("/consolidado?periodo=2026-07&nivel=PRIMARIA").get_data(
        as_text=True
    )
    assert 'name="periodo" value="2026-07"' in pagina
    client.get("/?periodo=2026-08&nivel=SECUNDARIA")
    assert generar(client).status_code == 303
    assert conn.execute(
        "SELECT periodo,nivel_modalidad FROM consolidado_dre"
    ).fetchone() == {"periodo": date(2026, 7, 1), "nivel_modalidad": "PRIMARIA"}
    assert post(client, "/consolidado", {"estado": "BORRADOR"}).status_code == 400


def test_localizador_de_hoja_no_elige_la_portada(client, conn, tmp_path):
    nx, cal, asistencia = fuentes(tmp_path)
    libro = openpyxl.load_workbook(asistencia)
    portada = libro.create_sheet("PORTADA", 0)
    portada["A1"] = "Esto es la portada ficticia"
    libro.save(asistencia)
    for path, tipo in ((nx, "nexus"), (cal, "calendario"), (asistencia, "asistencia")):
        cargar(client, path, tipo)
    r = conn.execute("SELECT * FROM reporte_asistencia").fetchone()
    assert r["hoja_pagina_origen"] == "ANEXO 3"
    html = client.get(
        f"/reportes/{r['reporte_asistencia_id']}/fuente?vista=celdas"
    ).get_data(as_text=True)
    assert "Ejemplo Ficticio Persona Uno" in html
    assert "Esto es la portada ficticia" not in html


def test_calendario_posterior_advierte_sin_cambiar_descarga(client, conn, tmp_path):
    rid, (_, cal, _) = ciclo(client, conn, tmp_path)
    revisar(client, conn, rid)
    r = generar(client, "BORRADOR")
    anterior = client.get(r.location + "/descargar").data
    assert "Hay fuentes o decisiones posteriores" not in client.get(
        r.location
    ).get_data(as_text=True)
    libro = openpyxl.load_workbook(cal)
    libro.active["B13"] = "G"
    nuevo = tmp_path / "calendario_rectificado.xlsx"
    libro.save(nuevo)
    cargar(client, nuevo, "calendario")
    assert "Hay fuentes o decisiones posteriores" in client.get(r.location).get_data(
        as_text=True
    )
    assert client.get(r.location + "/descargar").data == anterior


def test_archivo_fallido_no_impide_otros_del_lote(client, conn, tmp_path):
    nx, _, _ = fuentes(tmp_path)
    r = post(
        client,
        "/documentos",
        {
            "tipo": "nexus",
            "fecha_corte": "2026-06-01",
            "periodo": "2026-07",
            "archivos": [
                (io.BytesIO(b"no soy un xlsx"), "corrupto.xlsx"),
                (io.BytesIO(nx.read_bytes()), "correcto.xlsx"),
            ],
        },
    )
    assert r.status_code == 303
    estados = {
        r["estado"] for r in conn.execute("SELECT estado FROM web_carga").fetchall()
    }
    assert estados == {"PROCESADO", "NO_SOPORTADO"}
    assert conn.execute("SELECT count(*) AS n FROM web_recepcion").fetchone()["n"] == 2


def test_persona_busqueda_por_nombre_cruza_instituciones_sin_vincular(
    client, conn, tmp_path
):
    rid, _ = ciclo(client, conn, tmp_path)
    tid = conn.execute(
        "SELECT trabajador_en_reporte_id FROM trabajador_en_reporte"
    ).fetchone()["trabajador_en_reporte_id"]
    otra_nexus = construir_nexus_minimo(
        [
            {
                "CODMOD I.E.": "9999005",
                "CODIGO DE PLAZA": "PZ-WEB-05",
                "TIPO DE TRABAJADOR": "DOCENTE",
                "SITUACION LABORAL": "NOMBRADO",
                "TIPO DE REGISTRO": "ORGANICA",
                "DOCUMENTO DE IDENTIDAD": "99990005",
                "APELLIDO PATERNO": "Ejemplo",
                "APELLIDO MATERNO": "Otra",
                "NOMBRES": "Persona Cinco",
                "FECHA DE INICIO": date(2026, 1, 1),
                "CODIGO LOCAL": "999905",
                "DISTRITO": "LUYA",
                "NIVEL EDUCATIVO": "Primaria",
                "NOMBRE DE LA INSTITUCION EDUCATIVA": "9999005 IE WEB SEGUNDA",
                "CARGO": "DOCENTE",
            }
        ],
        tmp_path / "nexus_segunda.xlsx",
    )
    assert cargar(client, otra_nexus, "nexus").status_code == 303
    r = client.get(f"/reportes/{rid}/personas/{tid}?nombre=Ejemplo")
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert 'id="vinculo"' in html
    seleccionables = html.split('id="vinculo"', 1)[1].split("</select>", 1)[0]
    assert "Persona Uno" in seleccionables  # misma institución: seleccionable
    assert "Persona Cinco" not in seleccionables  # otra institución: nunca en el select
    assert "Persona Cinco" in html and "9999005 IE WEB SEGUNDA" in html
    assert "Encontrada(s) en otra institución" in html


def test_institucion_detalle_muestra_vinculos_y_404_en_id_inexistente(
    client, conn, tmp_path
):
    ciclo(client, conn, tmp_path)
    iid = conn.execute(
        "SELECT institucion_educativa_id FROM institucion_educativa LIMIT 1"
    ).fetchone()["institucion_educativa_id"]
    pagina = client.get(f"/instituciones/{iid}").get_data(as_text=True)
    assert "9999001 IE WEB FICTICIA" in pagina
    assert "Ejemplo" in pagina and "Persona Uno" in pagina
    respuesta = client.get("/instituciones/999999999")
    assert respuesta.status_code == 404
    assert "No encontramos ese registro" in respuesta.get_data(as_text=True)


def crear_alerta(conn, rid, *, severidad="ADVERTENCIA", codigo="COBERTURA_PRUEBA"):
    return conn.execute(
        "INSERT INTO validacion_reporte(reporte_asistencia_id,codigo_regla,severidad,mensaje) "
        "VALUES (%s,%s,%s,'Alerta ficticia para comprobar la ruta de decisión.') "
        "RETURNING validacion_reporte_id",
        (rid, codigo, severidad),
    ).fetchone()["validacion_reporte_id"]


def test_alerta_confirmar_actualiza_estado_y_no_desaparece(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    aid = crear_alerta(conn, rid)
    conn.commit()
    pagina = client.get(f"/reportes/{rid}").get_data(as_text=True)
    assert "Confirmar dato cotejado" in pagina and 'value="DESCARTAR_ALERTA"' in pagina
    r = post(
        client,
        f"/reportes/{rid}/alertas/{aid}",
        {
            "accion": "CONFIRMAR_ALERTA",
            "motivo": "Cotejado con la fuente ficticia; el dato aplicado es correcto.",
            "huella": huella_reporte(conn, rid),
        },
    )
    assert r.status_code == 303, r.get_data(as_text=True)
    fila = conn.execute(
        "SELECT * FROM validacion_reporte WHERE validacion_reporte_id=%s", (aid,)
    ).fetchone()
    assert fila["estado"] == "RESUELTA" and fila["resuelta_por"] is not None
    pagina = client.get(f"/reportes/{rid}").get_data(as_text=True)
    assert "Decisión registrada" in pagina
    assert (
        conn.execute(
            "SELECT count(*) AS n FROM web_revision_evento WHERE accion='CONFIRMAR_ALERTA' AND validacion_reporte_id=%s",
            (aid,),
        ).fetchone()["n"]
        == 1
    )


def test_alerta_descartar_actualiza_estado(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    aid = crear_alerta(conn, rid)
    conn.commit()
    r = post(
        client,
        f"/reportes/{rid}/alertas/{aid}",
        {
            "accion": "DESCARTAR_ALERTA",
            "motivo": "Se revisó el original y la observación no corresponde.",
            "huella": huella_reporte(conn, rid),
        },
    )
    assert r.status_code == 303
    assert (
        conn.execute(
            "SELECT estado FROM validacion_reporte WHERE validacion_reporte_id=%s",
            (aid,),
        ).fetchone()["estado"]
        == "DESCARTADA"
    )
    assert "Observación descartada" in client.get(f"/reportes/{rid}").get_data(
        as_text=True
    )


def test_alerta_dejar_pendiente_conserva_estado_y_registra_nota(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    aid = crear_alerta(conn, rid)
    conn.commit()
    r = post(
        client,
        f"/reportes/{rid}/alertas/{aid}",
        {
            "accion": "DEJAR_PENDIENTE",
            "motivo": "Falta confirmar con la institución; se retoma la próxima semana.",
            "huella": huella_reporte(conn, rid),
        },
    )
    assert r.status_code == 303
    fila = conn.execute(
        "SELECT * FROM validacion_reporte WHERE validacion_reporte_id=%s", (aid,)
    ).fetchone()
    assert fila["estado"] == "PENDIENTE" and fila["resuelta_por"] is None
    pagina = client.get(f"/reportes/{rid}").get_data(as_text=True)
    assert 'value="DEJAR_PENDIENTE"' in pagina  # la acción sigue disponible
    assert (
        "Falta confirmar con la institución" in pagina
    )  # la nota queda en el historial


def test_alerta_accion_invalida_no_cambia_nada(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    aid = crear_alerta(conn, rid)
    conn.commit()
    r = post(
        client,
        f"/reportes/{rid}/alertas/{aid}",
        {
            "accion": "NO_EXISTE",
            "motivo": "Motivo de prueba con longitud suficiente.",
            "huella": huella_reporte(conn, rid),
        },
    )
    assert r.status_code == 400
    assert (
        conn.execute(
            "SELECT estado FROM validacion_reporte WHERE validacion_reporte_id=%s",
            (aid,),
        ).fetchone()["estado"]
        == "PENDIENTE"
    )


def test_alerta_huella_desactualizada_no_aplica_decision(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    aid = crear_alerta(conn, rid)
    conn.commit()
    vieja = huella_reporte(conn, rid)
    assert revisar(client, conn, rid).status_code == 303  # cambia el estado y la huella
    r = post(
        client,
        f"/reportes/{rid}/alertas/{aid}",
        {
            "accion": "CONFIRMAR_ALERTA",
            "motivo": "Intento con huella desactualizada.",
            "huella": vieja,
        },
    )
    assert r.status_code == 409
    assert (
        conn.execute(
            "SELECT estado FROM validacion_reporte WHERE validacion_reporte_id=%s",
            (aid,),
        ).fetchone()["estado"]
        == "PENDIENTE"
    )


def test_alerta_ya_resuelta_rechaza_nueva_decision(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    aid = crear_alerta(conn, rid)
    conn.commit()
    huella = huella_reporte(conn, rid)
    assert (
        post(
            client,
            f"/reportes/{rid}/alertas/{aid}",
            {
                "accion": "CONFIRMAR_ALERTA",
                "motivo": "Primera decisión.",
                "huella": huella,
            },
        ).status_code
        == 303
    )
    r = post(
        client,
        f"/reportes/{rid}/alertas/{aid}",
        {
            "accion": "DESCARTAR_ALERTA",
            "motivo": "Segunda decisión distinta sobre lo ya resuelto.",
            "huella": huella_reporte(conn, rid),
        },
    )
    assert r.status_code == 409
    assert (
        conn.execute(
            "SELECT estado FROM validacion_reporte WHERE validacion_reporte_id=%s",
            (aid,),
        ).fetchone()["estado"]
        == "RESUELTA"
    )


def test_alerta_reenvio_con_misma_operacion_no_duplica(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    aid = crear_alerta(conn, rid)
    conn.commit()
    op = uuid.uuid4()
    datos = {
        "operacion": str(op),
        "accion": "CONFIRMAR_ALERTA",
        "motivo": "Decisión única reenviada dos veces por reintento del navegador.",
        "huella": huella_reporte(conn, rid),
    }
    assert post(client, f"/reportes/{rid}/alertas/{aid}", datos).status_code == 303
    assert post(client, f"/reportes/{rid}/alertas/{aid}", datos).status_code == 303
    assert (
        conn.execute(
            "SELECT count(*) AS n FROM web_revision_evento WHERE validacion_reporte_id=%s",
            (aid,),
        ).fetchone()["n"]
        == 1
    )


def test_calendario_aprobar_marca_vigente_y_degrada_version_previa(
    client, conn, tmp_path
):
    ciclo(client, conn, tmp_path)
    cv = conn.execute(
        "SELECT calendarizacion_version_id,calendarizacion_local_id FROM calendarizacion_version"
    ).fetchone()
    r = post(client, f"/calendarios/{cv['calendarizacion_version_id']}/aprobar")
    assert r.status_code == 303, r.get_data(as_text=True)
    actual = conn.execute(
        "SELECT * FROM calendarizacion_version WHERE calendarizacion_version_id=%s",
        (cv["calendarizacion_version_id"],),
    ).fetchone()
    assert actual["estado"] == "VIGENTE"
    assert actual["aprobado_por"] is not None and actual["aprobado_en"] is not None
    conn.execute(
        "INSERT INTO calendarizacion_version"
        "(calendarizacion_local_id,version,origen,estado,motivo_version) "
        "VALUES (%s,2,'CORRECCION_UGEL','BORRADOR','Segunda versión de prueba')",
        (cv["calendarizacion_local_id"],),
    )
    nueva = conn.execute(
        "SELECT calendarizacion_version_id FROM calendarizacion_version WHERE version=2"
    ).fetchone()["calendarizacion_version_id"]
    conn.commit()
    r2 = post(client, f"/calendarios/{nueva}/aprobar")
    assert r2.status_code == 303, r2.get_data(as_text=True)
    assert (
        conn.execute(
            "SELECT estado FROM calendarizacion_version WHERE calendarizacion_version_id=%s",
            (cv["calendarizacion_version_id"],),
        ).fetchone()["estado"]
        == "HISTORICA"
    )
    r3 = post(client, f"/calendarios/{cv['calendarizacion_version_id']}/aprobar")
    assert r3.status_code == 409


def test_calendario_aprobar_bloquea_por_pendiente_fuera_de_enero_febrero(
    client, conn, tmp_path
):
    ciclo(client, conn, tmp_path)
    cv = conn.execute(
        "SELECT calendarizacion_version_id FROM calendarizacion_version"
    ).fetchone()["calendarizacion_version_id"]
    conn.execute(
        "INSERT INTO dia_calendarizacion(calendarizacion_version_id,fecha,estado_captura,codigo_reportado_raw) "
        "VALUES (%s,%s,'CODIGO_DESCONOCIDO','Z')",
        (cv, date(2026, 8, 3)),
    )
    conn.commit()
    r = post(client, f"/calendarios/{cv}/aprobar")
    assert r.status_code == 409
    assert "fuera de enero/febrero" in r.get_data(as_text=True)
    assert (
        conn.execute(
            "SELECT estado FROM calendarizacion_version WHERE calendarizacion_version_id=%s",
            (cv,),
        ).fetchone()["estado"]
        == "BORRADOR"
    )


def test_calendario_aprobar_no_bloquea_vacio_en_enero_o_febrero(client, conn, tmp_path):
    ciclo(client, conn, tmp_path)
    cv = conn.execute(
        "SELECT calendarizacion_version_id FROM calendarizacion_version"
    ).fetchone()["calendarizacion_version_id"]
    conn.execute(
        "INSERT INTO dia_calendarizacion(calendarizacion_version_id,fecha,estado_captura) "
        "VALUES (%s,%s,'VACIO'),(%s,%s,'VACIO')",
        (cv, date(2026, 1, 15), cv, date(2026, 2, 10)),
    )
    conn.commit()
    r = post(client, f"/calendarios/{cv}/aprobar")
    assert r.status_code == 303, r.get_data(as_text=True)
    assert (
        conn.execute(
            "SELECT estado FROM calendarizacion_version WHERE calendarizacion_version_id=%s",
            (cv,),
        ).fetchone()["estado"]
        == "VIGENTE"
    )


# --- Cobertura del mes (R1) -------------------------------------------------------------------


def test_cobertura_muestra_desglose_conciliado_y_declara_su_universo(
    client, conn, tmp_path
):
    ciclo(client, conn, tmp_path)
    pagina = client.get("/cobertura?periodo=2026-07&nivel=PRIMARIA").text
    assert "Extraído, pendiente de revisión" in pagina
    # El universo declarado debe viajar con la cifra: no es la cobertura validada de M04.
    assert "No es el manifiesto de unidades esperadas validado" in pagina


def test_cobertura_pasa_a_revisado_cuando_rrhh_confirma_el_reporte(
    client, conn, tmp_path
):
    """La pantalla refleja el trabajo real de RRHH, no un estado inventado."""
    rid, _ = ciclo(client, conn, tmp_path)
    antes = client.get("/cobertura?periodo=2026-07&nivel=PRIMARIA&estado=REVISADO").text
    assert "Ninguna institución en este estado" in antes

    assert revisar(client, conn, rid).status_code == 303

    despues = client.get(
        "/cobertura?periodo=2026-07&nivel=PRIMARIA&estado=REVISADO"
    ).text
    assert "IE WEB FICTICIA" in despues


def test_cobertura_nunca_presenta_la_estimacion_como_medicion(client, conn, tmp_path):
    ciclo(client, conn, tmp_path)
    pagina = client.get("/cobertura?periodo=2026-07&nivel=PRIMARIA").text
    assert "ESTIMADO_DECLARADO" in pagina
    assert "Esto no es una medición" in pagina
    assert "Banda de sensibilidad" in pagina
    assert "no es un tiempo cronometrado" in pagina.lower()


def test_cobertura_no_promete_tiempo_absoluto_por_nivel(client, conn, tmp_path):
    """El tiempo declarado es por UGEL-mes. La pantalla debe decir por que el nivel no lo muestra,
    en vez de afirmar que falta una declaracion que si existe (se contradeciria con el bloque de
    la UGEL, que la cita justo debajo)."""
    ciclo(client, conn, tmp_path)
    pagina = client.get("/cobertura?periodo=2026-07&nivel=PRIMARIA").text
    assert (
        "es para <strong>todas</strong> las instituciones a la vez, no por nivel"
        in pagina
    )
    assert "Falta la declaración de tiempo manual" not in pagina
    # El absoluto si aparece, una sola vez, en el bloque de toda la UGEL.
    assert pagina.count("Tiempo manual declarado:") == 1


def test_cobertura_con_universo_vacio_no_muestra_cero_por_ciento(
    client, conn, tmp_path
):
    ciclo(client, conn, tmp_path)
    pagina = client.get("/cobertura?periodo=2026-07&nivel=PRITE").text
    assert "El denominador es cero" in pagina
    assert "%" not in pagina.split("Desglose del universo")[1].split("</section>")[0]


def test_cobertura_csv_se_descarga_con_el_mismo_calculo(client, conn, tmp_path):
    ciclo(client, conn, tmp_path)
    r = client.get("/cobertura.csv?periodo=2026-07&nivel=PRIMARIA")
    assert r.status_code == 200
    assert "attachment" in r.headers["Content-Disposition"]
    cuerpo = r.get_data(as_text=True)
    assert cuerpo.startswith("# universo,")
    assert "IE WEB FICTICIA" in cuerpo
    assert "EXTRAIDO_PENDIENTE" in cuerpo

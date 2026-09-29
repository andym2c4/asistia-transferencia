"""Separar leyendas de metadatos: disposición real con contenido nominal ficticio."""

from datetime import date

import openpyxl
import pytest
from psycopg.types.json import Jsonb

from asistia.cierre.leyendas import releer_leyenda_excel
from asistia.db import sha256_de
from asistia.importar.asistencia import importar_asistencia
from asistia.importar.nexus import importar_nexus
from asistia.leyendas import leer_leyenda_xlsx

from .test_catalogo_general import operador
from .test_web import fuentes

CODIGOS = ["A", "I", "3T", "J", "L", "P", "T", "F", "H"]
TEXTOS = [
    "Día laborado",
    "Inasistencia injustificada",
    "Tercera tardanza considerada como inasistencia injustificada",
    "Inasistencia justificada(Licencia, permiso, vacaciones)",
    "Licencia con goce de remuneraciones",
    "Permiso sin goce de remuneraciones",
    "Tardanza",
    "Feriados",
    "Huelga  o paro",
]


def hoja_julio():
    ws = openpyxl.Workbook().active
    ws.title = "ANEXO 3"
    ws["A24"] = "LEYENDA:"
    for fila, (codigo, texto) in enumerate(zip(CODIGOS, TEXTOS, strict=True), 25):
        ws.cell(fila, 1, codigo)
        ws.cell(fila, 2, texto)
    for fila in (25, 31, 33):
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
    ws["J28"] = "APELLIDOS Y NOMBRES DEL DIRECTOR"
    ws["L29"] = "DNI:99990001"
    return ws


def test_julio_nueve_categorias_sin_dni_de_firma():
    entradas = leer_leyenda_xlsx(hoja_julio())
    assert [(e["codigo"], e["descripcion"]) for e in entradas] == list(
        zip(CODIGOS, TEXTOS, strict=True)
    )
    assert [e["celda"] for e in entradas] == [f"A{f}:B{f}" for f in range(25, 34)]


def test_descripcion_ausente_no_se_completa_con_una_firma_lejana():
    ws = hoja_julio()
    ws["B33"] = None
    ws["L33"] = "Firma de persona ficticia"
    assert [e["codigo"] for e in leer_leyenda_xlsx(ws)] == CODIGOS[:-1]


@pytest.mark.parametrize(
    "metadato",
    [
        "DNI:99990001",
        "99990001: Persona ficticia",
        "DNI: Persona ficticia",
        "TEL: 999 000 001",
    ],
)
def test_firma_lateral_no_es_parte_del_bloque(metadato):
    ws = hoja_julio()
    ws["L29"] = metadato
    assert [e["codigo"] for e in leer_leyenda_xlsx(ws)] == CODIGOS


def test_bloque_termina_y_otro_rotulo_abre_otra_leyenda():
    ws = hoja_julio()
    ws["A37"] = "REF: Documento adjunto"
    ws["A41"] = "LEYENDA"
    ws["A42"] = "X"
    ws["B42"] = "Situación excepcional propia"
    assert [e["codigo"] for e in leer_leyenda_xlsx(ws)] == [*CODIGOS, "X"]


def test_leyenda_inline_y_codigos_numericos_legitimos():
    ws = openpyxl.Workbook().active
    ws["A1"] = "LEYENDA: 1=Situación excepcional propia"
    ws["A2"] = "2"
    ws["B2"] = "Otra situación propia"
    assert [(e["codigo"], e["descripcion"]) for e in leer_leyenda_xlsx(ws)] == [
        ("1", "Situación excepcional propia"),
        ("2", "Otra situación propia"),
    ]


def test_identificadores_no_se_convierten_en_categorias_en_columna_de_leyenda():
    ws = openpyxl.Workbook().active
    ws.append(["LEYENDA"])
    ws.append(["DNI:99990001"])
    ws.append(["99990001: Persona ficticia"])
    ws.append(["A", "Día laborado"])
    assert [e["codigo"] for e in leer_leyenda_xlsx(ws)] == ["A"]


def test_titulo_centrado_y_texto_local_sin_corregir_ortografia():
    ws = openpyxl.Workbook().active
    ws["D15"] = "LEYENDA:"
    ws["B17"] = "A"
    ws["C17"] = "DIAS LABORABLES"
    ws["B18"] = "I"
    ws["C18"] = "INACISTENCIA INJUSTIFICADA"
    entradas = leer_leyenda_xlsx(ws)
    assert [(e["codigo"], e["descripcion"]) for e in entradas] == [
        ("A", "DIAS LABORABLES"),
        ("I", "INACISTENCIA INJUSTIFICADA"),
    ]


def test_tercera_tardanza_y_plural_sin_rotulo():
    ws = openpyxl.Workbook().active
    ws["AQ49"] = "3T"
    ws["AR49"] = "Tercera tardanza, considerada como inasistencia injustificada"
    ws["AQ50"] = "T"
    ws["AR50"] = "Tardanzas"
    assert [e["codigo"] for e in leer_leyenda_xlsx(ws)] == ["3T", "T"]


def importar_julio_ficticio(conn, tmp_path):
    nx, _, ruta = fuentes(tmp_path)
    w = openpyxl.load_workbook(ruta)
    w.active["A14"] = None
    for row in hoja_julio():
        for celda in row:
            if celda.value is not None:
                w.active[celda.coordinate] = celda.value
    w.save(ruta)
    w.close()
    importar_nexus(conn, nx, date(2026, 1, 1))
    r = importar_asistencia(conn, ruta)
    assert not r.error and not r.errores_fila
    return str(next(iter(r.reportes_asistencia_id.values()))), ruta


def test_importacion_nativa_y_relectura_preservan_version_y_correcciones(
    conn, tmp_path
):
    rid, ruta = importar_julio_ficticio(conn, tmp_path)
    uid = operador(conn)
    categorias = conn.execute(
        "SELECT clasificacion_codigos FROM reporte_asistencia WHERE reporte_asistencia_id=%s",
        (rid,),
    ).fetchone()["clasificacion_codigos"]
    assert set(categorias) == set(CODIGOS)
    # Simular la contaminación del lector antiguo y una revisión local posterior.
    categorias["DNI"] = {
        "tipo_dia": "99990001",
        "fuente": "LEYENDA_DOCUMENTAL",
        "evidencia": {"hoja": "ANEXO 3", "celda": "L29"},
    }
    categorias["A"].update(fuente="REVISION_WEB", es_remunerado=True)
    conn.execute(
        "UPDATE reporte_asistencia SET clasificacion_codigos=%s WHERE reporte_asistencia_id=%s",
        (Jsonb(categorias), rid),
    )
    conn.execute(
        "UPDATE asistencia_dia SET codigo_reportado_raw='DNI' WHERE fecha='2026-07-02'"
    )
    conn.execute("UPDATE asistencia_dia SET validado=true WHERE fecha='2026-07-03'")
    antes = conn.execute(
        "SELECT to_jsonb(a) dato FROM asistencia_dia a ORDER BY fecha"
    ).fetchall()
    sha = sha256_de(ruta)
    nuevo = releer_leyenda_excel(
        conn, rid, "ANEXO 3", "Cotejo técnico ficticio de la firma", uid
    )
    assert nuevo != rid
    assert (
        releer_leyenda_excel(
            conn, rid, "ANEXO 3", "Cotejo técnico ficticio de la firma", uid
        )
        == nuevo
    )
    actual = conn.execute(
        "SELECT * FROM reporte_asistencia WHERE reporte_asistencia_id=%s", (nuevo,)
    ).fetchone()
    assert actual["estado"] == "EN_VALIDACION" and actual["validado_por"] is None
    assert actual["clasificacion_codigos"] == {
        k: v for k, v in categorias.items() if k != "DNI"
    }
    assert actual["procedencia_extraccion"]["relectura_leyenda"]["retirados"] == ["DNI"]
    assert sha256_de(ruta) == sha
    assert (
        conn.execute(
            "SELECT to_jsonb(a) dato FROM asistencia_dia a JOIN trabajador_en_reporte t USING(trabajador_en_reporte_id) WHERE t.reporte_asistencia_id=%s ORDER BY fecha",
            (rid,),
        ).fetchall()
        == antes
    )
    dias = conn.execute(
        "SELECT a.* FROM asistencia_dia a JOIN trabajador_en_reporte t USING(trabajador_en_reporte_id) WHERE t.reporte_asistencia_id=%s ORDER BY fecha",
        (nuevo,),
    ).fetchall()
    assert (
        dias[1]["codigo_reportado_raw"] == "DNI"
        and dias[1]["estado_captura"] == "PENDIENTE"
    )
    assert dias[1]["estado_asistencia_id"] is None
    assert (
        dias[2]["validado"]
        and dias[2]["evidencia_interpretacion"]["clasificacion_aceptada"]
    )
    assert (
        conn.execute(
            "SELECT count(*) n FROM web_revision_evento WHERE accion='RELEER_LEYENDA_EXCEL'"
        ).fetchone()["n"]
        == 1
    )
    with pytest.raises(ValueError, match="versión posterior"):
        releer_leyenda_excel(conn, rid, "ANEXO 3", "Otra operación", uid)


def test_relectura_rechaza_hoja_ambigua_y_original_alterado(conn, tmp_path):
    rid, ruta = importar_julio_ficticio(conn, tmp_path)
    uid = operador(conn)
    with pytest.raises(ValueError, match="nombre exacto"):
        releer_leyenda_excel(conn, rid, "1", "Ficticio", uid)
    with ruta.open("ab") as f:
        f.write(b"CAMBIO FICTICIO")
    with pytest.raises(ValueError, match="hash verificado"):
        releer_leyenda_excel(conn, rid, "ANEXO 3", "Ficticio", uid)
    assert (
        conn.execute("SELECT count(*) n FROM reporte_asistencia").fetchone()["n"] == 1
    )

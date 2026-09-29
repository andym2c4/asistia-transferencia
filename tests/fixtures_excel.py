"""Construye libros minimos que imitan la estructura real (NEXUS / ANEXO 3) para los dos casos de
P02 sin evidencia real disponible en esta sesion: fuente contradictoria e identidad no resuelta.
Identidades y codigos son ficticios (prefijo 9999...) y no deben confundirse con datos de la UGEL.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

_COLUMNAS_NEXUS = [
    "CODMOD I.E.",
    "CODIGO DE PLAZA",
    "TIPO DE TRABAJADOR",
    "SITUACION LABORAL",
    "TIPO DE REGISTRO",
    "DOCUMENTO DE IDENTIDAD",
    "APELLIDO PATERNO",
    "APELLIDO MATERNO",
    "NOMBRES",
    "FECHA DE INICIO",
    "FECHA DE TERMINO",
    "CODIGO LOCAL",
    "DISTRITO",
    "NIVEL EDUCATIVO",
    "NOMBRE DE LA INSTITUCION EDUCATIVA",
    "CARGO",
    "ESTADO",
    "MOTIVO DE VACANTE",
]


def construir_nexus_minimo(filas: list[dict[str, Any]], ruta: Path) -> Path:
    """Un corte NEXUS con las columnas reales (`asistia.importar.nexus._COLUMNAS_REQUERIDAS`) y
    las filas dadas (dict columna->valor; columnas omitidas quedan en blanco)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_COLUMNAS_NEXUS)
    for fila in filas:
        ws.append([fila.get(col) for col in _COLUMNAS_NEXUS])
    wb.save(ruta)
    return ruta


def construir_anexo3_minimo(
    ruta: Path,
    institucion_raw: str,
    nivel_raw: str = "PRIMARIA",
    periodo: date = date(2026, 7, 1),
    personas: list[dict[str, Any]] | None = None,
) -> Path:
    """ANEXO 3 minimo: encabezado I.E/periodo/nivel/turno, fila de columnas DNI/nombres/cargo, fila
    de dias (1..5) y las `personas` dadas (dict con dni, nombres, cargo, marcas por dia 1-5)."""
    personas = personas if personas is not None else []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ANEXO 3"

    ws["A1"] = "I.E:"
    ws["C1"] = institucion_raw
    ws["A2"] = "PERIODO(MES/AÑO):"
    ws["C2"] = periodo
    ws["A3"] = "NIVEL/MODALIDAD EDUCATIVA:"
    ws["C3"] = nivel_raw
    ws["A4"] = "TURNO:"
    ws["C4"] = "MAÑANA"

    fila_cols = 9
    ws.cell(row=fila_cols, column=1, value="N°")
    ws.cell(row=fila_cols, column=2, value="DNI")
    ws.cell(row=fila_cols, column=3, value="APELLIDOS Y NOMBRES")
    ws.cell(row=fila_cols, column=4, value="CARGO")

    fila_dias = fila_cols + 1
    for i in range(5):
        ws.cell(row=fila_dias, column=5 + i, value=i + 1)

    fila_datos = fila_cols + 3
    for idx, persona in enumerate(personas):
        fila = fila_datos + idx
        ws.cell(row=fila, column=1, value=idx + 1)
        ws.cell(row=fila, column=2, value=persona.get("dni"))
        ws.cell(row=fila, column=3, value=persona.get("nombres"))
        ws.cell(row=fila, column=4, value=persona.get("cargo"))
        marcas = persona.get("marcas", {})
        for dia, codigo in marcas.items():
            ws.cell(row=fila, column=5 + (dia - 1), value=codigo)

    ws.cell(row=fila_datos + len(personas) + 1, column=1, value="LEYENDA: A=asistencia")

    wb.save(ruta)
    return ruta

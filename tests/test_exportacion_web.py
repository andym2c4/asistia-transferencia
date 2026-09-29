"""RF-F30: descarga principal fiel y detalles disponibles solo en desarrollo."""

import io
from pathlib import Path

import openpyxl
import pytest

from asistia.db import sha256_de
from asistia.web.lecturas import salida

from .test_web import app, ciclo, client, generar  # noqa: F401

# ruff: noqa: F811


@pytest.mark.parametrize("desarrollo", [False, True])
def test_descarga_una_hoja_fiel_y_completo_segun_entorno(
    app, client, conn, tmp_path, desarrollo
):
    app.config["DEVELOPMENT_EXPORTS"] = desarrollo
    ciclo(client, conn, tmp_path)
    url = generar(client).location
    conservado = salida(conn, url.split("/")[-1])
    ruta = Path(conservado["ruta_objeto"])
    original = ruta.read_bytes()
    antes = conservado["manifiesto"]
    assert ("data-development-export" in client.get(url).text) is desarrollo
    normal = client.get(url + "/descargar")
    assert normal.status_code == 200
    assert "attachment" in normal.headers["Content-Disposition"]
    assert "completo_desarrollo" not in normal.headers["Content-Disposition"]
    libro = openpyxl.load_workbook(io.BytesIO(normal.data))
    fuente = openpyxl.load_workbook(io.BytesIO(original))
    assert len(fuente.sheetnames) > 1
    assert libro.sheetnames == [fuente.sheetnames[0]]
    a, b = fuente.worksheets[0], libro.worksheets[0]
    assert a.max_row == b.max_row and a.max_column == b.max_column
    for fila_a, fila_b in zip(a, b, strict=True):
        for celda_a, celda_b in zip(fila_a, fila_b, strict=True):
            assert celda_a.value == celda_b.value
            assert celda_a.data_type == celda_b.data_type
            assert celda_a._style == celda_b._style
    assert str(a.merged_cells) == str(b.merged_cells)
    assert a.freeze_panes == b.freeze_panes
    assert a.page_setup == b.page_setup
    assert a.page_margins == b.page_margins
    assert a.print_area == b.print_area
    for originales, derivados in (
        (a.row_dimensions, b.row_dimensions),
        (a.column_dimensions, b.column_dimensions),
    ):
        assert originales.keys() == derivados.keys()
        for key, dimension in originales.items():
            assert dict(dimension) == dict(derivados[key])
            assert dimension._style == derivados[key]._style
    assert "BORRADOR" in b["A7"].value
    assert client.get(url + "/descargar").data == normal.data
    completo = client.get(url + "/descargar?detalle=1")
    if desarrollo:
        assert completo.status_code == 200 and completo.data == original
        assert "completo_desarrollo" in completo.headers["Content-Disposition"]
    else:
        assert completo.status_code == 404
        assert (
            client.get(
                url + "/descargar?detalle=1&desarrollo=1",
                headers={
                    "X-Forwarded-Host": "localhost",
                    "X-Forwarded-For": "127.0.0.1",
                },
            ).status_code
            == 404
        )
    assert client.get(url + "/descargar?detalle=otro").status_code == 400
    assert sha256_de(ruta) == conservado["sha256"]
    assert ruta.read_bytes() == original
    assert salida(conn, url.split("/")[-1])["manifiesto"] == antes
    fuente.close()
    libro.close()


def test_descarga_verifica_integridad_y_autenticacion(app, client, conn, tmp_path):
    app.config["DEVELOPMENT_EXPORTS"] = True
    ciclo(client, conn, tmp_path)
    url = generar(client).location
    assert app.test_client().get(url + "/descargar").status_code == 302
    assert app.test_client().get(url + "/descargar?detalle=1").status_code == 302
    ruta = Path(salida(conn, url.split("/")[-1])["ruta_objeto"])
    assert ruta.is_relative_to(tmp_path)
    ruta.chmod(0o600)  # Simular corrupción solo sobre el archivo temporal ficticio.
    ruta.write_bytes(b"alteracion ficticia deliberada")
    for detalle in ("0", "1"):
        response = client.get(url + f"/descargar?detalle={detalle}")
        assert response.status_code == 409
        assert not response.headers.get("Content-Disposition")

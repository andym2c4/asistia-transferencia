"""P02: reproduccion de punta a punta (NEXUS -> calendario -> asistencia -> consolidado -> XLSX)
para Tactamal, julio 2026, comparando contra docs/casos/TACTAMAL_JULIO_2026.md. No sustituye una
evaluacion con referencia independiente (P03/P07); confirma que el codigo actual reproduce lo que
el caso ya documento a mano.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl

from asistia.consolidado.generar import exportar_consolidado_xlsx, generar_consolidado
from asistia.importar.asistencia import importar_asistencia
from asistia.importar.calendario import importar_calendario
from asistia.importar.nexus import importar_nexus

REPO_ROOT = Path(__file__).resolve().parent.parent
BASE = REPO_ROOT / "data/raw/data_brindada_por_ugel"
RUTA_NEXUS_JUNIO = BASE / "nexus/nexus 2026-06-01.xlsx"
RUTA_NEXUS_ABRIL = BASE / "nexus/NEXUS LUYA AL 01-04-2026.xls"
RUTA_CAL_PRIMARIA = (
    BASE
    / "calendarizaciones/CALENDARIZACIONES 2026 UGEL LUYA/CALENDARIZACIONES PRIMARIA 2026"
    "/18091-Tactamal.xlsx"
)
RUTA_CAL_SECUNDARIA = (
    BASE
    / "calendarizaciones/CALENDARIZACIONES 2026 UGEL LUYA/CALENDARIZACIONES SECUNDARIA 2026"
    "/CALENDARIZACIÓN 18091 -TACTAMAL.xlsx"
)
RUTA_ASISTENCIA = (
    BASE / "ASISTENCIAS JULIO 2026/ASISTENCIA SECUNDARIA JULIO"
    "/ASISTENCIA JULIO IEPySM N  18091 TACTAMAL (ANEXOS 3y4) UGEL LUYA.xlsx"
)


def _dias_julio_por_tipo(conn, calendarizacion_version_id) -> dict[str, int]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT ct.codigo_interno AS tipo, count(*) AS n
            FROM dia_calendarizacion dc
            JOIN catalogo_tipo_dia ct ON ct.tipo_dia_id = dc.tipo_dia_id
            WHERE dc.calendarizacion_version_id = %s
              AND dc.fecha BETWEEN '2026-07-01' AND '2026-07-31'
            GROUP BY ct.codigo_interno
            """,
            (calendarizacion_version_id,),
        )
        return {fila["tipo"]: fila["n"] for fila in cur.fetchall()}


def test_ciclo_completo_tactamal_julio(conn, tmp_path):
    importar_nexus(conn, RUTA_NEXUS_JUNIO, date(2026, 6, 1))
    importar_nexus(conn, RUTA_NEXUS_ABRIL, date(2026, 4, 1))

    cal_primaria = importar_calendario(conn, RUTA_CAL_PRIMARIA)
    cal_secundaria = importar_calendario(conn, RUTA_CAL_SECUNDARIA)
    assert cal_primaria.error is None
    assert cal_secundaria.error is None

    # docs/casos/TACTAMAL_JULIO_2026.md §5: "16 dias lectivos, 3 de gestion y 12 de tipo D"
    assert _dias_julio_por_tipo(conn, cal_primaria.calendarizacion_version_id) == {
        "LECTIVO": 16,
        "GESTION": 3,
        "NO_LECTIVO_NI_GESTION": 12,
    }
    # §6: "los totales son iguales en ambos calendarios" (la distribucion difiere en 2 dias,
    # el total no)
    assert _dias_julio_por_tipo(
        conn, cal_secundaria.calendarizacion_version_id
    ) == _dias_julio_por_tipo(conn, cal_primaria.calendarizacion_version_id)

    asistencia = importar_asistencia(conn, RUTA_ASISTENCIA)
    assert asistencia.error is None
    assert asistencia.total_trabajadores == 12
    assert asistencia.resueltos == 12

    consolidado_primaria = generar_consolidado(conn, date(2026, 7, 1), "Primaria")
    consolidado_secundaria = generar_consolidado(conn, date(2026, 7, 1), "Secundaria")

    # §5: 5 personas de primaria; P02 (Reubicada por Excedencia) queda pendiente de confirmar
    assert consolidado_primaria.personas == 5
    assert consolidado_primaria.reportes_incluidos == 1
    assert consolidado_primaria.validaciones_pendientes == 1
    # §10.1: 7 personas de secundaria, sin encargatura ni reemplazo pendiente
    assert consolidado_secundaria.personas == 7
    assert consolidado_secundaria.reportes_incluidos == 1
    assert consolidado_secundaria.validaciones_pendientes == 0

    salida = tmp_path / "consolidado_secundaria_2026-07.xlsx"
    ruta_exportada = exportar_consolidado_xlsx(
        conn, consolidado_secundaria.consolidado_dre_id, salida
    )
    assert ruta_exportada.exists()
    assert ruta_exportada.stat().st_size > 0

    wb = openpyxl.load_workbook(ruta_exportada)
    ws = wb.active
    assert "SECUNDARIA" in str(ws["A4"].value).upper()
    assert "JULIO" in str(ws["A6"].value).upper()

    with conn.cursor() as cur:
        cur.execute(
            "SELECT archivo_exportado_id FROM consolidado_dre WHERE consolidado_dre_id = %s",
            (consolidado_secundaria.consolidado_dre_id,),
        )
        assert cur.fetchone()["archivo_exportado_id"] is not None

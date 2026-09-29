"""RF-F14: representación fiel, privada y de consulta; originales ficticios."""
# ruff: noqa: F811

import hashlib
import io
import uuid

import openpyxl
import pytest
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image
from pypdf import PdfReader, PdfWriter

from asistia.web import ErrorDeTrabajo, routes
from asistia.web import visor_original as visor
from asistia.web.lecturas import huella_reporte

from .test_web import app, ciclo, client  # noqa: F401


def libro(path):
    wb = openpyxl.Workbook()
    wb.active.title = "Portada"
    wb.active["A1"] = "PORTADA FICTICIA"
    ws = wb.create_sheet("ANEXO 3")
    ws.merge_cells("A1:J2")
    ws["A1"] = "REPORTE FICTICIO DE JULIO"
    ws["A1"].fill = PatternFill("solid", fgColor="FFDA66")
    ws["A1"].font = Font(bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A4"] = "PERSONA FICTICIA"
    ws["B4"] = "A"
    ws["C4"] = "F"
    ws["B4"].fill = PatternFill("solid", fgColor="00B050")
    ws["C4"].fill = PatternFill("solid", fgColor="FF0000")
    ws.column_dimensions["A"].width = 24
    ws["J8"] = "FIN DE LA HOJA"
    wb.create_sheet("Oculta").sheet_state = "hidden"
    wb.save(path)
    return path


def original(monkeypatch, path, *, origen="1"):
    rid = uuid.uuid4()
    r = {
        "reporte_asistencia_id": rid,
        "nombre_original": path.name,
        "hoja_pagina_origen": origen,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }
    monkeypatch.setattr(routes, "ruta_original_reporte", lambda _: (r, path))
    return f"/reportes/{rid}/fuente", r


def test_excel_conserva_formato_hoja_original_y_cache(
    client, app, tmp_path, monkeypatch
):
    path = libro(tmp_path / "libro.xlsx")
    url, r = original(monkeypatch, path)
    antes = path.read_bytes()
    html = client.get(url)
    assert "data-original-viewer" in html.text and "data-viewer-image" in html.text
    assert "contenteditable" not in html.text and "data-report-day" not in html.text
    assert "blob:" in html.headers["Content-Security-Policy"]
    response = client.get(url + "/vista")
    assert response.status_code == 200, response.text
    meta = response.json
    assert meta["hoja"] == "ANEXO 3" and meta["hojas"] == ["Portada", "ANEXO 3"]
    assert meta["paginas"] == 1
    png = client.get(meta["imagen_url"])
    assert png.status_code == 200 and png.mimetype == "image/png"
    assert png.headers["Cache-Control"] == "no-store, private"
    with Image.open(io.BytesIO(png.data)) as image:
        colors = list(image.convert("RGB").get_flattened_data())
        assert any(
            red > 200 and green < 60 and blue < 60 for red, green, blue in colors
        )
        assert any(
            red > 200 and green > 150 and blue < 150 for red, green, blue in colors
        )
    cache = (
        app.config["STORAGE_ROOT"] / "vistas" / visor.RECETA / r["sha256"] / "hoja-1"
    )
    text = PdfReader(cache / "documento.pdf").pages[0].extract_text()
    assert (
        "REPORTE FICTICIO" in text
        and "FIN DE LA HOJA" in text
        and "PORTADA FICTICIA" not in text
    )
    monkeypatch.setattr(
        visor,
        "ejecutar_aislado",
        lambda *args, **kwargs: pytest.fail("No debe convertir una caché verificada"),
    )
    assert client.get(url + "/vista").json == meta
    assert client.get(meta["imagen_url"]).data == png.data
    assert path.read_bytes() == antes


def test_seleccion_hoja_visible_y_parametros_no_son_rutas(
    client, tmp_path, monkeypatch
):
    url, _ = original(monkeypatch, libro(tmp_path / "libro.xlsx"))
    response = client.get(url + "/vista?hoja=Portada")
    assert response.status_code == 200 and response.json["hoja"] == "Portada"
    for name in ["Oculta", "../archivo", "<script>", "No existe"]:
        response = client.get(url + "/vista", query_string={"hoja": name})
        assert response.status_code == 404 and "error" in response.json
    assert client.get(url + "/pagina/0?hoja=Portada").status_code == 404
    assert client.get(url + "/pagina/999?hoja=Portada").status_code == 404


def test_pdf_rasterizado_no_expone_campos_ni_acciones(client, tmp_path, monkeypatch):
    path = tmp_path / "original.pdf"
    writer = PdfWriter()
    writer.add_blank_page(400, 200)
    writer.add_blank_page(200, 400)
    writer.add_js("app.alert('Esta acción no debe ejecutarse en el visor');")
    writer.write(path)
    before = path.read_bytes()
    url, _ = original(monkeypatch, path)
    meta = client.get(url + "/vista").json
    assert meta["paginas"] == 2 and meta["tipo"] == "pdf" and not meta["hojas"]
    for n in [1, 2]:
        png = client.get(url + f"/pagina/{n}")
        assert png.status_code == 200 and png.mimetype == "image/png"
        assert png.data.startswith(b"\x89PNG")
    assert path.read_bytes() == before


def test_fotografia_conserva_bytes_y_autenticacion(client, app, tmp_path, monkeypatch):
    path = tmp_path / "foto.png"
    Image.new("RGB", (80, 40), "#3178aa").save(path)
    url, _ = original(monkeypatch, path)
    assert client.get(url + "/vista").json["tipo"] == "imagen"
    assert client.get(url + "/pagina/1").data == path.read_bytes()
    for endpoint in ["", "/vista", "/pagina/1"]:
        assert app.test_client().get(url + endpoint).status_code == 302
    assert (
        "blob:" not in client.get("/static/app.css").headers["Content-Security-Policy"]
    )


def test_cache_pdf_corrupto_se_regenera(client, app, tmp_path, monkeypatch):
    path = libro(tmp_path / "libro.xlsx")
    url, r = original(monkeypatch, path)
    assert client.get(url + "/vista").status_code == 200
    pdf = (
        app.config["STORAGE_ROOT"]
        / "vistas"
        / visor.RECETA
        / r["sha256"]
        / "hoja-1"
        / "documento.pdf"
    )
    pdf.write_bytes(b"corrupto")
    assert client.get(url + "/vista").status_code == 200
    assert len(PdfReader(pdf).pages) == 1


def test_convertidor_no_disponible_se_informa_sin_sustituir_original(
    client, tmp_path, monkeypatch
):
    url, _ = original(monkeypatch, libro(tmp_path / "libro.xlsx"))

    def falla(*args, **kwargs):
        raise ErrorDeTrabajo(
            "No se pudo representar este documento. Puedes descargar el original.", 503
        )

    monkeypatch.setattr(visor, "ejecutar_aislado", falla)
    assert client.get(url).status_code == 200
    response = client.get(url + "/vista")
    assert response.status_code == 503 and "descargar" in response.json["error"]


def test_limiter_paginas_y_pdf_roto(client, tmp_path, monkeypatch):
    path = tmp_path / "muchas.pdf"
    writer = PdfWriter()
    for _ in range(101):
        writer.add_blank_page(72, 72)
    writer.write(path)
    url, _ = original(monkeypatch, path)
    response = client.get(url + "/vista")
    assert response.status_code == 422 and "100 páginas" in response.json["error"]
    path.write_bytes(b"%PDF-roto")
    url, _ = original(monkeypatch, path)
    assert client.get(url + "/vista").status_code == 422


def test_visor_integrado_no_modifica_reporte_o_fuente(client, conn, tmp_path):
    rid, _ = ciclo(client, conn, tmp_path)
    before = huella_reporte(conn, rid)
    url = f"/reportes/{rid}/fuente"
    assert client.get(url).status_code == 200
    meta = client.get(url + "/vista").json
    assert client.get(meta["imagen_url"]).status_code == 200
    assert huella_reporte(conn, rid) == before
    assert (
        conn.execute("SELECT count(*) n FROM web_revision_evento").fetchone()["n"] == 0
    )
